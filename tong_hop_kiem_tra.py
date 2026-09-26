import sys
import os
import re
import json
import hashlib
import hmac
import subprocess
import urllib.request
import urllib.parse
import urllib.error
from datetime import datetime
from collections import Counter
import docx
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QLineEdit, QFileDialog, QTabWidget,
    QTableWidget, QTableWidgetItem, QHeaderView, QMessageBox,
    QProgressBar, QGroupBox, QGridLayout, QTextEdit, QSplitter,
    QComboBox, QSpinBox, QCheckBox
)
from PySide6.QtCore import Qt, QThread, Signal as pyqtSignal
from PySide6.QtGui import QFont, QColor

APP_VERSION_TH = "3.0.0"

# Nơi lưu các đường dẫn/cấu hình cho tab "Cập nhật dữ liệu & Build" - cùng thư mục với
# chương trình đang chạy (script hoặc .exe), để không mất khi đổi máy/đổi thư mục làm việc.
# Dùng sys.argv[0] thay vì __file__: khi build onefile (Nuitka/PyInstaller), __file__ trỏ vào
# thư mục giải nén tạm và bị xoá sau mỗi lần thoát -> cấu hình sẽ mất.
CONFIG_PATH = os.path.join(os.path.dirname(os.path.abspath(sys.argv[0])), "tong_hop_config.json")

def load_config():
    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}

def save_config(cfg):
    try:
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

_THANG_VIET_TAT = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def _danh_sach_thang_lui(thang_moc, so_thang):
    """Từ tháng mốc (vd '2026-Sep' hoặc '2026-09') tạo danh sách N tháng gần nhất tính lùi về trước,
    theo định dạng MSRC cần ('2026-Sep', '2026-Aug'...). Trả về [] nếu tháng mốc sai định dạng."""
    m = re.fullmatch(r"(\d{4})-([A-Za-z]{3,}|\d{1,2})", (thang_moc or "").strip())
    if not m:
        return []
    nam = int(m.group(1))
    khoa = m.group(2)
    if khoa.isdigit():
        thang = int(khoa)
    else:
        ten = khoa[:3].capitalize()
        if ten not in _THANG_VIET_TAT:
            return []
        thang = _THANG_VIET_TAT.index(ten) + 1
    if not (1 <= thang <= 12):
        return []
    ket_qua = []
    for _ in range(max(1, so_thang)):
        ket_qua.append(f"{nam}-{_THANG_VIET_TAT[thang - 1]}")
        thang -= 1
        if thang == 0:
            thang = 12
            nam -= 1
    return ket_qua


def _cap_nhat_data_version(path, ngay=None):
    """Ghi/cập nhật dòng '# DATA_VERSION: <ngày>' ở ĐẦU file dữ liệu (CVE hoặc IOC).
    Công cụ auto_fill sẽ đọc dòng này để in 'ngày dữ liệu' lên biên bản (phục vụ tính pháp lý)."""
    ngay = ngay or datetime.now().strftime("%d/%m/%Y")
    dong_moi = f"# DATA_VERSION: {ngay}\n"
    try:
        with open(path, "r", encoding="utf-8") as f:
            noi_dung = f.readlines()
    except Exception:
        noi_dung = []
    thay = False
    for i, dong in enumerate(noi_dung):
        if dong.strip().startswith("# DATA_VERSION:"):
            noi_dung[i] = dong_moi
            thay = True
            break
    if not thay:
        noi_dung.insert(0, dong_moi)
    with open(path, "w", encoding="utf-8") as f:
        f.writelines(noi_dung)


def doc_ngay_du_lieu_th(path):
    """Đọc dòng '# DATA_VERSION: <ngày>' ở đầu file dữ liệu. Trả về chuỗi ngày hoặc 'không rõ'."""
    if not path or not os.path.exists(path):
        return "không rõ"
    try:
        with open(path, encoding="utf-8") as f:
            for _ in range(30):
                line = f.readline()
                if not line:
                    break
                m = re.search(r"#\s*DATA_VERSION:\s*(.+)$", line)
                if m:
                    return m.group(1).strip()
    except Exception:
        pass
    return "không rõ"


def doc_phien_ban_script(duong_dan_script):
    """Đọc số phiên bản hiện tại (APP_VERSION) trong mã nguồn auto_fill_bien_ban.py.
    Trả về chuỗi kiểu '1.6', hoặc None nếu không đọc được."""
    try:
        with open(duong_dan_script, "r", encoding="utf-8") as f:
            noi_dung = f.read()
        m = re.search(r'^APP_VERSION\s*=\s*["\']([\d.]+)["\']', noi_dung, re.MULTILINE)
        return m.group(1) if m else None
    except Exception:
        return None

def tang_phien_ban_script(duong_dan_script):
    """Tự tăng số phụ (minor) của APP_VERSION trong mã nguồn lên 1 (vd '1.6' -> '1.7') và
    ghi thẳng lại vào file - để mỗi lần build, mã nguồn VÀ bản .exe xuất ra luôn đồng bộ,
    không phải tự tay sửa số phiên bản trước khi build. Chỉ giữ đúng 2 phần (major.minor)
    theo đúng ý dùng của người quản lý công cụ này.
    Trả về số phiên bản MỚI (str), hoặc None nếu không tìm/đọc/ghi được."""
    try:
        with open(duong_dan_script, "r", encoding="utf-8") as f:
            noi_dung = f.read()
        m = re.search(r'^(APP_VERSION\s*=\s*["\'])([\d.]+)(["\'])', noi_dung, re.MULTILINE)
        if not m:
            return None
        phan = m.group(2).split(".")
        major = int(phan[0]) if len(phan) > 0 and phan[0].isdigit() else 1
        minor = int(phan[1]) if len(phan) > 1 and phan[1].isdigit() else 0
        phien_ban_moi = f"{major}.{minor + 1}"
        noi_dung_moi = noi_dung[:m.start()] + m.group(1) + phien_ban_moi + m.group(3) + noi_dung[m.end():]
        with open(duong_dan_script, "w", encoding="utf-8") as f:
            f.write(noi_dung_moi)
        return phien_ban_moi
    except Exception:
        return None

# Danh bạ các lỗ hổng đặc biệt nguy hiểm cần gắn cờ cảnh báo
CRITICAL_VULNS_DICT = {
    "CVE-2017-0144": "EternalBlue (Khai thác SMB từ xa)",
    "CVE-2019-0708": "BlueKeep (Thực thi mã RDP)",
    "CVE-2020-0796": "SMBGhost (Thực thi mã SMBv3)",
    "CVE-2020-1472": "Zerologon (Chiếm quyền Netlogon Domain)",
    "CVE-2021-34527": "PrintNightmare (Thực thi mã Print Spooler)",
    "CVE-2022-30190": "Follina (Khai thác MSDT qua Office)"
}

def doc_muc_do_cve(duong_dan_cve):
    """Đọc file windows_vulnerabilities.txt -> dict {CVE: mức độ (CRITICAL/HIGH/...)}.
    Dùng để đánh giá rủi ro theo đúng mức độ của từng CVE thay vì chỉ đếm số lượng."""
    ket_qua = {}
    if not duong_dan_cve or not os.path.exists(duong_dan_cve):
        return ket_qua
    try:
        with open(duong_dan_cve, "r", encoding="utf-8") as f:
            for dong in f:
                dong = dong.strip()
                if not dong or dong.startswith("#"):
                    continue
                phan = dong.split("|")
                if len(phan) >= 3:
                    muc_do = phan[2].strip().upper()
                    # MSRC dùng thang Critical/Important/Moderate/Low -> quy về cùng thang với file gốc
                    muc_do = {"IMPORTANT": "HIGH", "MODERATE": "MEDIUM"}.get(muc_do, muc_do)
                    ket_qua[phan[0].strip().upper()] = muc_do
    except Exception:
        pass
    return ket_qua


# Ký tự Wingdings của ô đã tích (auto_fill_bien_ban dùng F0FE; F0FC/F0FD/F078/F0FB là các kiểu
# tích tay thường gặp khi cán bộ sửa lại biên bản trong Word). Ô trống là F0A8/F06F.
_KY_TU_O_DA_TICH = {"F0FE", "F0FD", "F0FC", "F0FB", "F078"}
_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"


def doc_o_tich(paragraph):
    """Trả về danh sách (nhãn, đã_tích) của các ô tích trong 1 đoạn văn, theo thứ tự.
    Nhãn là phần chữ đứng NGAY TRƯỚC ô tích (vd 'Đặt mật khẩu đăng nhập: Có [x]  không [ ]').
    Lưu ý: paragraph.text của python-docx KHÔNG chứa ký tự ô tích (w:sym), nên không thể dùng
    regex trên văn bản để biết ô nào được tích - phải đọc trực tiếp XML."""
    ket_qua = []
    nhan = ""
    for el in paragraph._p.iter():
        tag = el.tag
        if tag == f"{{{_W_NS}}}t":
            nhan += el.text or ""
        elif tag == f"{{{_W_NS}}}tab":
            nhan += " "
        elif tag == f"{{{_W_NS}}}sym":
            ky_tu = (el.get(f"{{{_W_NS}}}char") or "").upper()
            ket_qua.append((_lam_sach_nhan(nhan), ky_tu in _KY_TU_O_DA_TICH))
            nhan = ""
        elif tag == f"{{{_W14_NS}}}checked":  # ô tích dạng content control (Word 2010+)
            ket_qua.append((_lam_sach_nhan(nhan), el.get(f"{{{_W14_NS}}}val") in ("1", "true")))
            nhan = ""
    return ket_qua


def _lam_sach_nhan(nhan):
    nhan = nhan.split(":")[-1]
    return nhan.replace("|", " ").strip()


def _noi_dung_sau_tieu_de(paragraph_text):
    """Kết quả do auto_fill_bien_ban ghi thêm vào sau dòng tiêu đề mục (sau dấu xuống dòng)."""
    phan = paragraph_text.split("\n", 1)
    return [d.strip() for d in phan[1].split("\n") if d.strip()] if len(phan) > 1 else []


def _la_gia_tri_trong(gia_tri):
    return not gia_tri or re.fullmatch(r"[….\s]*", gia_tri) is not None


def parse_docx_content(file_path, muc_do_cve=None):
    muc_do_cve = muc_do_cve or {}
    doc = docx.Document(file_path)
    paragraphs = doc.paragraphs
    text = "\n".join(p.text.strip() for p in paragraphs if p.text.strip())

    filename = os.path.basename(file_path)
    data = {
        "File_Path": file_path,
        "File_Name": filename,
        "Thoi_Gian_KT": "Không rõ",
        "Dia_Diem": "Không rõ",
        "Can_Bo_KT": "Không rõ",
        "Chuc_Vu_KT": "",
        "Can_Bo_Quan_Ly": "Không rõ",
        "Mat_Khau": "Không rõ",
        "Cung_Cap_MK": "Không rõ",
        "Phan_Loai_May": "Không rõ",
        "Ten_May": "Không rõ",
        "He_Dieu_Hanh": "Không rõ",
        "Ngay_Cai_Dat": "Không rõ",
        "Dia_Chi_IP": "Không rõ",
        "Dia_Chi_MAC": "Không rõ",
        "CPU": "Không rõ",
        "RAM": "Không rõ",
        "Loai_OCung": "Không rõ",
        "DungLuong_OCung": "Không rõ",
        "Phan_Mem_Diet_Virus": "Không rõ",
        "Phan_Mem_Ung_Dung": "Không rõ",
        "Ket_Noi_Mang": "Không rõ",
        "Lich_Su_Internet": "Không rõ",
        "Tinh_Trang_Kiem_Tra_Lo_Hong": "Không rõ",
        "So_Luong_Lo_Hong": 0,
        "Danh_Sach_Lo_Hong": [],
        "Lo_Hong_Nguy_Hiem": [],
        "Ma_Doc": "Không rõ",
        "Ho_Ma_Doc": [],
        "Lich_Su_USB": [],
        "USB_Seri": [],
        "Muc_Do_Rui_Ro": "An toàn",
        "Ly_Do_Rui_Ro": [],
        "Toan_Ven": "Không có file kèm (biên bản cũ/thủ công)",
        "Nguon_Du_Lieu": "docx",
    }

    # 1. Thời gian & Địa điểm
    time_m = re.search(r"Vào hồi\s+([^,]+?),\s+ngày\s+(\d{1,2}\s+tháng\s+\d{1,2}\s+năm\s+\d{4})", text, re.I)
    if time_m:
        data["Thoi_Gian_KT"] = f"{time_m.group(1).strip()} - {time_m.group(2).strip()}"
    loc_m = re.search(r"Vào hồi[^\n]*?,\s*tại\s+([^\n]+?)\.?\s*$", text, re.I | re.M)
    if loc_m and not _la_gia_tri_trong(loc_m.group(1)):
        data["Dia_Diem"] = loc_m.group(1).strip()

    # 2. Cán bộ: "Đ/c <tên> – <chức vụ> tổ An ninh..." và "đồng chí: <tên> quản lý"
    cb_kt = re.search(r"Đ/c\s+(.+?)\s+[–\-]\s*(.*?)\s*tổ an ninh", text, re.I)
    if cb_kt:
        if not _la_gia_tri_trong(cb_kt.group(1)):
            data["Can_Bo_KT"] = cb_kt.group(1).strip()
        if not _la_gia_tri_trong(cb_kt.group(2)):
            data["Chuc_Vu_KT"] = cb_kt.group(2).strip()
    cb_ql = re.search(r"đồng chí[:\s]+(.+?)\s*quản lý", text, re.I)
    if cb_ql and not _la_gia_tri_trong(cb_ql.group(1)):
        data["Can_Bo_Quan_Ly"] = cb_ql.group(1).strip()
    elif len(doc.tables) > 1:
        # Dự phòng: ô ký tên "CÁN BỘ QUẢN LÝ, SỬ DỤNG THIẾT BỊ" ở bảng cuối biên bản
        try:
            ten = doc.tables[1].cell(1, 2).text.strip()
            if ten:
                data["Can_Bo_Quan_Ly"] = ten
        except Exception:
            pass

    # 3. Mật khẩu & Phân loại - đọc ô tích trực tiếp từ XML
    for p in paragraphs:
        dau = p.text.strip().lower()
        if dau.startswith("đặt mật khẩu đăng nhập"):
            o = doc_o_tich(p)
            if len(o) >= 2 and o[0][1] != o[1][1]:
                data["Mat_Khau"] = "Có đặt mật khẩu" if o[0][1] else "Không đặt MK"
        elif dau.startswith("có cung cấp mật khẩu"):
            o = doc_o_tich(p)
            if len(o) >= 2 and o[0][1] != o[1][1]:
                data["Cung_Cap_MK"] = "Có" if o[0][1] else "Không"
        elif dau.startswith("phân loại máy tính"):
            da_tich = [nhan for nhan, tich in doc_o_tich(p) if tich and nhan]
            if da_tich:
                data["Phan_Loai_May"] = ", ".join(da_tich)

    # 4. Thông số máy - đọc theo nhãn ở cột 1 của bảng thông tin (không phụ thuộc thứ tự dòng)
    if doc.tables:
        for row in doc.tables[0].rows:
            if len(row.cells) < 2:
                continue
            nhan = row.cells[0].text.strip().lower()
            gia_tri = row.cells[1].text.strip()
            if not gia_tri:
                continue
            if nhan.startswith("tên máy tính"):
                data["Ten_May"] = gia_tri
            elif nhan.startswith("hệ điều hành"):
                data["He_Dieu_Hanh"] = gia_tri
            elif nhan.startswith("thời gian cài đặt"):
                data["Ngay_Cai_Dat"] = gia_tri
            elif nhan == "mac":
                data["Dia_Chi_MAC"] = gia_tri
            elif nhan == "ip":
                data["Dia_Chi_IP"] = gia_tri
            elif nhan.startswith("cấu hình máy tính"):
                for key, pattern in (("CPU", r"^CPU:\s*(.+)$"), ("RAM", r"^RAM:\s*(.+)$"),
                                     ("Loai_OCung", r"^Ổ cứng loại:\s*(.+)$"),
                                     ("DungLuong_OCung", r"^Dung lượng:\s*(.+)$")):
                    m = re.search(pattern, gia_tri, re.I | re.M)
                    if m:
                        data[key] = m.group(1).strip()
            elif nhan.startswith("phần mềm diệt virus"):
                data["Phan_Mem_Diet_Virus"] = gia_tri
            elif nhan.startswith("phần mềm ứng dụng"):
                data["Phan_Mem_Ung_Dung"] = gia_tri
            elif nhan.startswith("tình trạng thiết bị"):
                m = re.search(r"Kết nối Internet:\s*(.+?)\.?\s*$", gia_tri, re.I | re.M)
                data["Ket_Noi_Mang"] = m.group(1).strip() if m else gia_tri

    # 5. Kết quả kiểm tra (mục II.1) - lấy các dòng auto_fill ghi thêm sau tiêu đề từng mục.
    # Chỉ lấy lần khớp ĐẦU TIÊN: mục II.2 cũng có dòng "- Mã độc hoặc phần mềm độc hại...".
    # Hỗ trợ CẢ 2 định dạng: cũ (kết quả nằm cùng đoạn tiêu đề, sau dấu xuống dòng) và mới (mỗi kết
    # quả là 1 đoạn văn riêng ngay dưới tiêu đề) - nên gom cả hai nguồn cho đến tiêu đề mục kế tiếp.
    muc = {}       # key -> danh sách dòng kết quả
    muc_head = {}  # key -> text đoạn tiêu đề (dùng cho mục ghi kết quả cùng dòng, vd Internet)
    n_par = len(paragraphs)
    tien_to_list = (("lo_hong", "- lỗ hổng bảo mật hệ điều hành"), ("ma_doc", "- mã độc"),
                    ("usb", "- lịch sử kết nối các thiết bị ngoại vi"),
                    ("internet", "- lịch sử kết nối internet"))
    for i, p in enumerate(paragraphs):
        dau = p.text.strip().lower()
        for key, tien_to in tien_to_list:
            if dau.startswith(tien_to) and key not in muc:
                lines = _noi_dung_sau_tieu_de(p.text)  # định dạng cũ (ngắt dòng mềm)
                for j in range(i + 1, n_par):          # định dạng mới (đoạn văn riêng)
                    t = paragraphs[j].text.strip()
                    if not t:
                        continue
                    tl = t.lower()
                    # Dừng ở tiêu đề mục kế tiếp. KHÔNG dùng số thứ tự chung vì dòng kết quả cũng
                    # đánh số ("1. WannaCry...", "1. Loại: USB..."). Chỉ dừng ở bullet "-", mục lớn
                    # (I./II./III.), mục "2. Đối với thiết bị khác..." và các đoạn kết.
                    if t.startswith("-") or re.match(r"^\d+\.\s*đối với", tl) \
                            or re.match(r"^(i{1,3})\.", tl) \
                            or tl.startswith(("sau khi tiến hành", "kết thúc quá trình",
                                              "biên bản kết thúc")):
                        break
                    lines.append(t)
                muc[key] = lines
                muc_head[key] = p.text

    # 5a. Lỗ hổng
    dong_lo_hong = muc.get("lo_hong", [])
    noi_dung_lo_hong = " ".join(dong_lo_hong)
    vuln_m = re.search(r"Phát hiện\s+(\d+)\s+lỗ hổng[^\(]*\((.*?)\)", noi_dung_lo_hong, re.I | re.DOTALL)
    if vuln_m:
        data["Tinh_Trang_Kiem_Tra_Lo_Hong"] = "Đã kiểm tra"
        data["So_Luong_Lo_Hong"] = int(vuln_m.group(1).strip())
        vuln_list = re.findall(r"CVE-\d{4}-\d+", vuln_m.group(2), re.I)
        data["Danh_Sach_Lo_Hong"] = [v.upper() for v in vuln_list]
    elif re.search(r"Không phát hiện lỗ hổng", noi_dung_lo_hong, re.I):
        data["Tinh_Trang_Kiem_Tra_Lo_Hong"] = "Đã kiểm tra"
    elif re.search(r"CẢNH BÁO", noi_dung_lo_hong):
        data["Tinh_Trang_Kiem_Tra_Lo_Hong"] = "Chưa kiểm tra (thiếu file CVE)"
    data["Lo_Hong_Nguy_Hiem"] = [v for v in data["Danh_Sach_Lo_Hong"]
                                 if v in CRITICAL_VULNS_DICT or muc_do_cve.get(v) == "CRITICAL"]

    # 5b. Mã độc: "1. <Họ mã độc>: <dấu hiệu>; ..." hoặc "Không phát hiện dấu hiệu (IOC)..."
    dong_ma_doc = muc.get("ma_doc", [])
    phat_hien = [re.sub(r"^\d+\.\s*", "", d) for d in dong_ma_doc if re.match(r"^\d+\.\s*\S", d)]
    if phat_hien:
        data["Ma_Doc"] = "PHÁT HIỆN: " + " | ".join(phat_hien)
        data["Ho_Ma_Doc"] = [d.split(":", 1)[0].strip() for d in phat_hien]
    elif any(re.search(r"Không phát hiện", d, re.I) for d in dong_ma_doc):
        data["Ma_Doc"] = "Không phát hiện"
    elif any("CẢNH BÁO" in d for d in dong_ma_doc):
        data["Ma_Doc"] = "Chưa kiểm tra (thiếu file IOC)"

    # 5c. Thiết bị ngoại vi / USB - mỗi dòng "N. Loại: .. | Seri: .. | Dung lượng: .. | Tên: .."
    for d in muc.get("usb", []):
        m = re.match(r"^\d+\.\s*Loại:\s*(.*?)\s*\|\s*Seri:\s*(.*?)\s*\|\s*Dung lượng:\s*(.*?)\s*\|\s*Tên:\s*(.+)$", d)
        if m:
            data["Lich_Su_USB"].append(f"{m.group(4).strip()} [{m.group(1).strip()}, {m.group(3).strip()}] "
                                       f"(Seri: {m.group(2).strip()})")
            data["USB_Seri"].append(m.group(2).strip())
        elif re.match(r"^\.\.\.\s*và\s+\d+", d):
            data["Lich_Su_USB"].append(d)

    # 5d. Lịch sử Internet (ghi cùng dòng với tiêu đề mục, hoặc đoạn văn ngay dưới)
    if muc_head.get("internet"):
        m = re.search(r"\(liệt kê chi tiết nếu có\):\s*(.+)$", muc_head["internet"], re.I | re.S)
        if m and m.group(1).strip():
            data["Lich_Su_Internet"] = m.group(1).strip()
        elif muc.get("internet"):
            data["Lich_Su_Internet"] = " ".join(muc["internet"]).strip()

    danh_gia_rui_ro(data, muc_do_cve)
    return data


MUC_RUI_RO = ["An toàn", "Trung bình (Medium)", "Cao (High)", "Nguy cấp (Critical)"]


def danh_gia_rui_ro(data, muc_do_cve):
    """Chấm mức rủi ro theo nhiều tiêu chí (không chỉ số lượng lỗ hổng) và ghi lại LÝ DO,
    để kiểm tra viên/lãnh đạo biết vì sao máy bị xếp mức đó và cần khắc phục gì."""
    muc = 0
    ly_do = []

    def nang(len_muc, ly):
        nonlocal muc
        muc = max(muc, len_muc)
        ly_do.append(ly)

    if data["Ho_Ma_Doc"]:
        nang(3, "Phát hiện dấu hiệu mã độc: " + ", ".join(data["Ho_Ma_Doc"]))
    if data["Lo_Hong_Nguy_Hiem"]:
        nang(3, f"{len(data['Lo_Hong_Nguy_Hiem'])} lỗ hổng mức NGUY CẤP chưa vá")
    so = data["So_Luong_Lo_Hong"]
    if so >= 20:
        nang(3, f"{so} lỗ hổng chưa vá (≥ 20)")
    elif so >= 10:
        nang(2, f"{so} lỗ hổng chưa vá (≥ 10)")
    elif so > 0:
        cao = [v for v in data["Danh_Sach_Lo_Hong"] if muc_do_cve.get(v) == "HIGH"]
        nang(2 if cao else 1, f"{so} lỗ hổng chưa vá" + (f" ({len(cao)} mức CAO)" if cao else ""))

    # Hệ điều hành hết hỗ trợ: không còn bản vá bảo mật -> lỗ hổng mới sẽ không bao giờ được vá
    hdh = data["He_Dieu_Hanh"].lower()
    if re.search(r"windows\s*(xp|vista|7|8(\.1)?)\b", hdh):
        nang(2, "Hệ điều hành đã hết hỗ trợ (không còn bản vá bảo mật)")
    elif re.search(r"windows\s*10\b", hdh) and "ltsc" not in hdh and "ltsb" not in hdh:
        nang(1, "Windows 10 hết hỗ trợ từ 14/10/2025 (trừ khi có đăng ký ESU) - cần nâng cấp")

    phan_loai = data["Phan_Loai_May"].lower()
    co_internet = "internet" in phan_loai or data["Ket_Noi_Mang"].lower().startswith("có")
    if "nội bộ" in phan_loai and co_internet:
        nang(2, "Máy nội bộ đồng thời kết nối Internet")
    if data["Mat_Khau"] == "Không đặt MK":
        nang(2 if co_internet else 1, "Không đặt mật khẩu đăng nhập")
    if data["Cung_Cap_MK"] == "Có":
        nang(1, "Cung cấp mật khẩu cho người khác")
    if re.search(r"không (có|xác định|phát hiện)", data["Phan_Mem_Diet_Virus"], re.I):
        nang(1, "Không xác định được phần mềm diệt virus")
    for key, ten in (("Ma_Doc", "mã độc"), ("Tinh_Trang_Kiem_Tra_Lo_Hong", "lỗ hổng")):
        if data[key].startswith("Chưa kiểm tra"):
            nang(1, f"Chưa đối chiếu {ten} (thiếu dữ liệu) - cần kiểm tra lại")
        elif data[key] == "Không rõ":
            # Biên bản không có kết quả mục này (điền tay/sửa mẫu) - không được coi là "An toàn"
            nang(1, f"Chưa đối chiếu {ten} (biên bản không ghi kết quả) - cần kiểm tra lại")

    data["Muc_Do_Rui_Ro"] = MUC_RUI_RO[muc]
    data["Ly_Do_Rui_Ro"] = ly_do


# ============================================================
# TOÀN VẸN BIÊN BẢN: xác minh chữ ký HMAC của file .attt.json do auto_fill ghi kèm
# ============================================================
DEFAULT_INTEGRITY_KEY = "ANATTT-CAX-TriPhu-2026-bien-ban-integrity-default-key"


def lay_integrity_key(cfg):
    """Khoá xác minh: ưu tiên khoá riêng của đơn vị trong cấu hình, nếu chưa có thì dùng khoá mặc định
    (khớp với auto_fill khi chưa nhúng khoá riêng)."""
    return (cfg.get("integrity_key") or "").strip() or DEFAULT_INTEGRITY_KEY


def _canonical_json(obj):
    return json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _tinh_hmac(key, obj):
    import hmac
    return hmac.new(key.encode("utf-8"), _canonical_json(obj).encode("utf-8"), hashlib.sha256).hexdigest()


def _sha256_file(path):
    try:
        h = hashlib.sha256()
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1024 * 1024), b""):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return None


def xac_minh_toan_ven(docx_path, key):
    """Tìm file kèm <biên_bản>.attt.json, kiểm chữ ký HMAC và so mã băm docx.
    Trả về (trang_thai, payload_hoặc_None):
      - 'Không có file kèm (biên bản cũ/thủ công)'
      - 'Chữ ký SAI - nghi bị giả mạo'         (HMAC không khớp -> KHÔNG tin dữ liệu)
      - 'Docx ĐÃ bị sửa sau khi tạo'           (HMAC đúng, mã băm docx khác -> dùng dữ liệu JSON gốc)
      - 'Hợp lệ'                               (HMAC đúng, docx khớp nguyên trạng)
    """
    sidecar = os.path.splitext(docx_path)[0] + ".attt.json"
    if not os.path.exists(sidecar):
        return "Không có file kèm (biên bản cũ/thủ công)", None
    try:
        with open(sidecar, "r", encoding="utf-8") as f:
            goi = json.load(f)
        payload = goi.get("payload")
        chu_ky = goi.get("hmac")
        if not isinstance(payload, dict) or not chu_ky:
            return "File kèm hỏng định dạng", None
        import hmac as _hmac
        if not _hmac.compare_digest(chu_ky, _tinh_hmac(key, payload)):
            return "Chữ ký SAI - nghi bị giả mạo", None
        docx_hash = _sha256_file(docx_path)
        if payload.get("docx_sha256") and docx_hash and payload["docx_sha256"] != docx_hash:
            return "Docx ĐÃ bị sửa sau khi tạo", payload
        return "Hợp lệ", payload
    except Exception:
        return "File kèm hỏng định dạng", None


def parse_sidecar(payload, file_path, muc_do_cve=None):
    """Dựng dữ liệu phân tích TRỰC TIẾP từ file kèm đã ký (đáng tin hơn regex trên docx)."""
    muc_do_cve = muc_do_cve or {}
    hw = payload.get("hardware") or {}
    manual = payload.get("manual") or {}
    ph = payload.get("password_has")
    nc = payload.get("network_class")
    online = bool(payload.get("online"))

    phan_loai = {"noi_bo": "Nội bộ", "doc_lap": "Độc lập"}.get(nc, "Không rõ")
    if online and phan_loai != "Không rõ":
        phan_loai += ", Internet"

    gio = manual.get("gio", ""); phut = manual.get("phut", "")
    ngay = manual.get("ngay", ""); thang = manual.get("thang", ""); nam = manual.get("nam", "")
    thoi_gian = f"{gio}:{phut} - {ngay}/{thang}/{nam}" if (gio or ngay) else "Không rõ"

    ho_ma_doc = payload.get("malware_families") or []
    cves = [c.upper() for c in (payload.get("vuln_cves") or [])]

    usb, seri = [], []
    for it in payload.get("peripherals") or []:
        ten = (it.get("ten") or "").strip()
        s = (it.get("serial") or "").strip()
        usb.append(f"{ten} [{it.get('loai','')}, {it.get('dung_luong','')}] (Seri: {s})")
        seri.append(s)

    data = {
        "File_Path": file_path,
        "File_Name": os.path.basename(file_path),
        "Thoi_Gian_KT": thoi_gian,
        "Dia_Diem": manual.get("dia_diem", "") or "Không rõ",
        "Can_Bo_KT": manual.get("ten_can_bo", "") or "Không rõ",
        "Chuc_Vu_KT": manual.get("chuc_vu", ""),
        "Can_Bo_Quan_Ly": manual.get("ten_doi_tuong", "") or "Không rõ",
        "Mat_Khau": {True: "Có đặt mật khẩu", False: "Không đặt MK"}.get(ph, "Không rõ"),
        "Cung_Cap_MK": "Không rõ",
        "Phan_Loai_May": phan_loai,
        "Ten_May": payload.get("computer_name", "") or "Không rõ",
        "He_Dieu_Hanh": payload.get("os_display") or payload.get("os_caption", "") or "Không rõ",
        "Ngay_Cai_Dat": payload.get("ngay_cai", "") or "Không rõ",
        "Dia_Chi_IP": payload.get("ip", "") or "Không rõ",
        "Dia_Chi_MAC": payload.get("mac", "") or "Không rõ",
        "CPU": hw.get("cpu", "") or "Không rõ",
        "RAM": str(hw.get("ram_gb", "")) or "Không rõ",
        "Loai_OCung": hw.get("loai_o_cung", "") or "Không rõ",
        "DungLuong_OCung": hw.get("dung_luong_o_cung", "") or "Không rõ",
        "Phan_Mem_Diet_Virus": payload.get("antivirus", "") or "Không rõ",
        "Phan_Mem_Ung_Dung": payload.get("top_apps", "") or "Không rõ",
        "Ket_Noi_Mang": "Có" if online else "Không",
        "Lich_Su_Internet": "Đang có kết nối Internet." if online else "Hiện không kết nối Internet.",
        "Tinh_Trang_Kiem_Tra_Lo_Hong": "Đã kiểm tra",
        "So_Luong_Lo_Hong": int(payload.get("vuln_count") or 0),
        "Danh_Sach_Lo_Hong": cves,
        "Lo_Hong_Nguy_Hiem": [v for v in cves if v in CRITICAL_VULNS_DICT or muc_do_cve.get(v) == "CRITICAL"],
        "Ma_Doc": ("PHÁT HIỆN: " + " | ".join(payload.get("malware_raw") or ho_ma_doc)) if ho_ma_doc else "Không phát hiện",
        "Ho_Ma_Doc": ho_ma_doc,
        "Lich_Su_USB": usb,
        "USB_Seri": seri,
        "Muc_Do_Rui_Ro": "An toàn",
        "Ly_Do_Rui_Ro": [],
        "Toan_Ven": "Hợp lệ",
        "Nguon_Du_Lieu": "file kèm đã ký (.attt.json)",
    }
    danh_gia_rui_ro(data, muc_do_cve)
    return data


def phan_tich_bien_ban(file_path, muc_do_cve=None, key=None):
    """Phân tích 1 biên bản: ưu tiên file kèm đã ký (.attt.json), dùng docx làm dự phòng.
    Luôn gắn trạng thái toàn vẹn để công cụ cảnh báo biên bản bị sửa/giả mạo."""
    key = key or DEFAULT_INTEGRITY_KEY
    trang_thai, payload = xac_minh_toan_ven(file_path, key)
    if payload is not None:  # HMAC hợp lệ -> tin dữ liệu trong file kèm
        data = parse_sidecar(payload, file_path, muc_do_cve)
        data["Toan_Ven"] = trang_thai
        return data
    # Không có file kèm, chữ ký sai, hoặc file kèm hỏng -> đọc docx và gắn cảnh báo tương ứng
    data = parse_docx_content(file_path, muc_do_cve)
    data["Toan_Ven"] = trang_thai
    if trang_thai.startswith("Chữ ký SAI"):
        if "Nghi biên bản bị giả mạo (chữ ký sai)" not in data["Ly_Do_Rui_Ro"]:
            data["Ly_Do_Rui_Ro"].insert(0, "Nghi biên bản bị giả mạo (chữ ký sai)")
        data["Muc_Do_Rui_Ro"] = "Nguy cấp (Critical)"
    return data


class WorkerThread(QThread):
    progress = pyqtSignal(int)
    file_processed = pyqtSignal(dict)
    finished = pyqtSignal(list)

    file_error = pyqtSignal(str, str)

    def __init__(self, folder_path, muc_do_cve=None, integrity_key=None):
        super().__init__()
        self.folder_path = folder_path
        self.muc_do_cve = muc_do_cve or {}
        self.integrity_key = integrity_key or DEFAULT_INTEGRITY_KEY

    def run(self):
        # Quét cả thư mục con (biên bản thường được xếp theo từng cơ quan/thôn/đợt kiểm tra)
        files = []
        for goc, _, ten_files in os.walk(self.folder_path):
            for f in sorted(ten_files):
                if f.lower().endswith(".docx") and not f.startswith("~$"):
                    files.append(os.path.join(goc, f))
        total = len(files)
        results = []
        for idx, file_path in enumerate(files):
            try:
                res = phan_tich_bien_ban(file_path, self.muc_do_cve, self.integrity_key)
                res["File_Name"] = os.path.relpath(file_path, self.folder_path)
                results.append(res)
                self.file_processed.emit(res)
            except Exception as e:
                self.file_error.emit(os.path.relpath(file_path, self.folder_path), str(e))
            if total > 0:
                self.progress.emit(int((idx + 1) / total * 100))
        self.finished.emit(results)


# ============================================================
#  CẬP NHẬT DỮ LIỆU (CVE Windows từ MSRC / IOC mã độc từ ThreatFox) & BUILD LẠI
#  auto_fill_bien_ban.exe - dùng cho tab 3.
# ============================================================
def fetch_cisa_kev():
    """Tải danh sách 'Known Exploited Vulnerabilities' (KEV) của CISA (công khai, miễn phí).
    Đây là các lỗ hổng ĐÃ BỊ KHAI THÁC THỰC TẾ - ưu tiên số 1 khi kiểm tra. Trả về set các CVE ID.
    Lỗi mạng chỉ trả về set rỗng (không chặn luồng CVE của MSRC)."""
    url = "https://www.cisa.gov/sites/default/files/feeds/known_exploited_vulnerabilities.json"
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=30) as resp:
            data = json.loads(resp.read().decode("utf-8"))
        return {(v.get("cveID") or "").strip().upper()
                for v in data.get("vulnerabilities", []) if v.get("cveID")}
    except Exception:
        return set()


def _bocs_fixed_build(chuoi):
    """Bóc '10.0.19045.4291' -> '19045.4291' (bản dựng gốc + UBR), bỏ tiền tố '10.0.'."""
    m = re.search(r"(\d{4,6})\.(\d{1,7})\s*$", (chuoi or "").strip())
    return f"{m.group(1)}.{m.group(2)}" if m else None


def fetch_msrc_cve(year_month, kev_set=None):
    """Tải + bóc tách CVE từ MSRC CVRF API (public, không cần API key) cho 1 tháng
    (vd '2026-Apr'), chỉ giữ lại CVE có liên quan Windows 10/11/Server.
    Trả về list dict: {cve, title, severity, kbs, fixed_builds, base_score, exploited, kev}."""
    kev_set = kev_set or set()
    url = f"https://api.msrc.microsoft.com/cvrf/v3.0/cvrf/{year_month}"
    req = urllib.request.Request(url, headers={"Accept": "application/json",
                                                "User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        data = json.loads(resp.read().decode("utf-8"))

    # Bản đồ ProductID -> tên sản phẩm, để lọc riêng các dòng liên quan Windows
    id_to_name = {}
    def _duyet(node):
        if isinstance(node, dict):
            if "ProductID" in node and "Value" in node:
                id_to_name[node["ProductID"]] = node["Value"]
            for v in node.values():
                _duyet(v)
        elif isinstance(node, list):
            for it in node:
                _duyet(it)
    _duyet((data.get("ProductTree") or {}))

    tu_khoa_windows = ("windows 10", "windows 11", "windows server")
    ket_qua = []
    for vuln in data.get("Vulnerability", []) or []:
        cve = (vuln.get("CVE") or "").strip()
        if not cve:
            continue
        title = ((vuln.get("Title") or {}).get("Value") or "").strip()

        pid_lien_quan = set()
        for ps in vuln.get("ProductStatuses", []) or []:
            pid_lien_quan.update(ps.get("ProductID", []) or [])
        ten_sp = {id_to_name.get(pid, "") for pid in pid_lien_quan}
        if not any(any(k in (ten or "").lower() for k in tu_khoa_windows) for ten in ten_sp):
            continue  # không liên quan Windows 10/11/Server -> bỏ qua (Edge/Office/Azure/...)

        kbs = set()
        fixed_builds = set()
        for rem in vuln.get("Remediations", []) or []:
            mo_ta = ((rem.get("Description") or {}).get("Value") or "").strip()
            so = re.sub(r"[^0-9]", "", mo_ta)
            if len(so) >= 6:
                kbs.add(f"KB{so}")
            # FixedBuild: bản dựng Windows đã vá lỗ hổng này (vd '10.0.19045.4291').
            # Chỉ giữ các dòng Windows máy ở xã có thể dùng (build gốc >= 10240 = Win10 trở lên).
            fb = _bocs_fixed_build(rem.get("FixedBuild") or "")
            if fb and int(fb.split(".")[0]) >= 10240:
                fixed_builds.add(fb)

        base_score = None
        for cvss in vuln.get("CVSSScoreSets", []) or []:
            if cvss.get("BaseScore") is not None:
                base_score = cvss["BaseScore"]
                break
        muc_do = ""
        exploited = cve.upper() in kev_set
        for th in vuln.get("Threats", []) or []:
            mo_ta = ((th.get("Description") or {}).get("Value") or "").strip()
            mo_ta_l = mo_ta.lower()
            if mo_ta_l in ("critical", "important", "moderate", "low"):
                muc_do = mo_ta_l.upper()
            if "exploited:yes" in mo_ta_l.replace(" ", ""):
                exploited = True  # MSRC ghi nhận đã bị khai thác thực tế
        if not muc_do and base_score is not None:
            muc_do = ("CRITICAL" if base_score >= 9.0 else "HIGH" if base_score >= 7.0
                      else "MEDIUM" if base_score >= 4.0 else "LOW")

        # Gộp fixed_build tốt nhất cho mỗi dòng Windows (UBR nhỏ nhất đủ để coi là đã vá)
        theo_dong = {}
        for fb in fixed_builds:
            goc, ubr = fb.split(".")
            goc, ubr = int(goc), int(ubr)
            if goc not in theo_dong or ubr < theo_dong[goc]:
                theo_dong[goc] = ubr
        fixed_list = [f"{g}.{u}" for g, u in sorted(theo_dong.items())]

        ket_qua.append({"cve": cve, "title": title, "severity": muc_do or "N/A",
                        "kbs": sorted(kbs), "fixed_builds": fixed_list,
                        "base_score": base_score, "exploited": exploited,
                        "kev": cve.upper() in kev_set})
    return ket_qua


def fetch_threatfox_recent(auth_key, days=7):
    """Tải các IOC mới nhất (N ngày gần đây) từ ThreatFox - cần Auth-Key miễn phí
    (đăng ký tại https://auth.abuse.ch/). Trả về list dict thô từ ThreatFox."""
    body = json.dumps({"query": "get_iocs", "days": days}).encode("utf-8")
    req = urllib.request.Request(
        "https://threatfox-api.abuse.ch/api/v1/", data=body,
        headers={"Auth-Key": auth_key, "Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        data = json.loads(resp.read().decode("utf-8"))
    if data.get("query_status") != "ok":
        raise RuntimeError(f"ThreatFox báo lỗi: {data.get('query_status')}")
    return data.get("data", []) or []


TEN_MIEN_HOP_PHAP = (
    "github.com", "githubusercontent.com", "google.com", "googleapis.com", "googleusercontent.com",
    "microsoft.com", "live.com", "sharepoint.com", "onedrive.com", "1drv.ms", "windows.net",
    "dropbox.com", "dropboxusercontent.com", "discord.com", "discordapp.com", "discordapp.net",
    "telegram.org", "t.me", "facebook.com", "fbcdn.net", "zalo.me", "amazonaws.com",
    "cloudfront.net", "bitbucket.org", "gitlab.com", "pastebin.com", "mediafire.com",
    "mega.nz", "4shared.com", "cloudflare.com", "workers.dev", "pages.dev", "azureedge.net",
)


# IP KHÔNG được đưa vào danh sách mã độc dù ThreatFox có gắn cờ: DNS công cộng và dịch vụ hợp pháp.
# Một dòng như 'ip:8.8.8.8' sẽ khiến MỌI máy có kết nối Internet bị báo "có mã độc".
IP_KHONG_DUNG = {
    "8.8.8.8", "8.8.4.4", "1.1.1.1", "1.0.0.1", "9.9.9.9", "149.112.112.112",
    "208.67.222.222", "208.67.220.220", "4.2.2.2", "4.2.2.1", "0.0.0.0",
}


def _ip_hop_le(ip):
    m = re.fullmatch(r"(\d{1,3})\.(\d{1,3})\.(\d{1,3})\.(\d{1,3})", ip or "")
    return bool(m) and all(0 <= int(x) <= 255 for x in m.groups())


def _ip_rieng_tu(ip):
    """IP nội bộ/không định tuyến (LAN, loopback, link-local) - không phải máy chủ mã độc trên Internet."""
    p = [int(x) for x in ip.split(".")]
    return (p[0] == 10 or p[0] == 127 or (p[0] == 192 and p[1] == 168)
            or (p[0] == 172 and 16 <= p[1] <= 31) or (p[0] == 169 and p[1] == 254)
            or p[0] == 0 or p[0] >= 224)


def _domain_hop_le(dom):
    dom = (dom or "").strip().lower().strip(".")
    if not dom or "." not in dom or len(dom) > 253 or re.fullmatch(r"[\d.]+", dom):
        return False
    return bool(re.fullmatch(r"[a-z0-9]([a-z0-9\-]*[a-z0-9])?(\.[a-z0-9]([a-z0-9\-]*[a-z0-9])?)+", dom))


def _domain_bi_cam(dom):
    dom = dom.lower().strip(".")
    return any(dom == d or dom.endswith("." + d) for d in TEN_MIEN_HOP_PHAP)


def _ioc_qua_han(item, max_age_days):
    """True nếu IOC (đặc biệt IP máy chủ điều khiển) đã quá cũ. IP C2 đổi rất nhanh nên IP cũ
    dễ đã được cấp lại cho dịch vụ hợp pháp -> gây báo nhầm."""
    if not max_age_days:
        return False
    raw = (item.get("last_seen") or item.get("first_seen") or "").strip()
    if not raw:
        return False
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M:%S UTC", "%Y-%m-%d"):
        try:
            return (datetime.now() - datetime.strptime(raw, fmt)).days > max_age_days
        except ValueError:
            continue
    return False


def quy_doi_ioc_threatfox(item, min_confidence=75, max_age_days=90):
    """Đổi 1 IOC thô từ ThreatFox sang đúng định dạng dòng của malware_signatures.txt và LỌC CHẤT LƯỢNG.
    Chỉ giữ loại công cụ kiểm tra dùng được (sha256/domain/ip). Trả về None nếu bị loại vì:
    độ tin cậy thấp, quá hạn, sai định dạng, hoặc là dữ liệu dễ gây báo nhầm (IP nội bộ, DNS công cộng,
    tên miền dịch vụ hợp pháp)."""
    try:
        if int(item.get("confidence_level") or 0) < min_confidence:
            return None
    except (TypeError, ValueError):
        pass
    if _ioc_qua_han(item, max_age_days):
        return None

    loai = (item.get("ioc_type") or "").lower()
    gia_tri = (item.get("ioc") or "").strip()
    if not gia_tri:
        return None

    if loai == "sha256_hash":
        return f"sha256:{gia_tri.lower()}" if re.fullmatch(r"[0-9a-fA-F]{64}", gia_tri) else None
    if loai == "domain":
        dom = gia_tri.lower().strip(".")
        return f"domain:{dom}" if _domain_hop_le(dom) and not _domain_bi_cam(dom) else None
    if loai in ("ip:port", "ip"):
        ip = gia_tri.split(":")[0].strip()
        return f"ip:{ip}" if _ip_hop_le(ip) and not _ip_rieng_tu(ip) and ip not in IP_KHONG_DUNG else None
    if loai == "url":
        try:
            ten_mien = (urllib.parse.urlparse(gia_tri).hostname or "").lower().strip(".")
            # URL độc hại đặt trên dịch vụ hợp pháp (github, google drive, discord...) KHÔNG được
            # đổi thành IOC tên miền: công cụ kiểm tra so khớp tên miền trong cache DNS, nên sẽ báo
            # "có mã độc" oan cho mọi máy chỉ vì từng mở Google Drive/GitHub.
            if _domain_hop_le(ten_mien) and not _domain_bi_cam(ten_mien):
                return f"domain:{ten_mien}"
            return None
        except Exception:
            return None
    return None  # md5/win_registry_key/... - chưa có chỗ tương ứng trong định dạng hiện tại


class CveFetchThread(QThread):
    ok = pyqtSignal(list)
    tien_trinh = pyqtSignal(str)
    error = pyqtSignal(str)

    def __init__(self, months, chi_da_khai_thac=False):
        super().__init__()
        self.months = months if isinstance(months, (list, tuple)) else [months]
        self.chi_da_khai_thac = chi_da_khai_thac

    def run(self):
        try:
            self.tien_trinh.emit("Đang tải danh sách lỗ hổng đã bị khai thác (CISA KEV)...")
            kev = fetch_cisa_kev()
            gom = {}
            loi = []
            for thang in self.months:
                self.tien_trinh.emit(f"Đang tải CVE tháng {thang} từ MSRC...")
                try:
                    for item in fetch_msrc_cve(thang, kev):
                        # Trùng CVE giữa các tháng: giữ bản có nhiều fixed_build/KB hơn
                        cu = gom.get(item["cve"])
                        if cu is None or len(item["fixed_builds"]) > len(cu["fixed_builds"]):
                            gom[item["cve"]] = item
                except Exception as e:
                    loi.append(f"{thang}: {e}")
            ket_qua = list(gom.values())
            if self.chi_da_khai_thac:
                ket_qua = [c for c in ket_qua if c["exploited"] or c["kev"]]
            # Sắp xếp: đã khai thác trước, rồi theo mức độ
            uu_tien = {"CRITICAL": 0, "HIGH": 1, "MEDIUM": 2, "LOW": 3, "N/A": 4}
            ket_qua.sort(key=lambda c: (not (c["exploited"] or c["kev"]),
                                        uu_tien.get(c["severity"], 5), c["cve"]))
            if loi and not ket_qua:
                self.error.emit("Không tải được tháng nào:\n" + "\n".join(loi))
                return
            self.ok.emit(ket_qua)
        except Exception as e:
            self.error.emit(str(e))


class IocFetchThread(QThread):
    ok = pyqtSignal(list)
    error = pyqtSignal(str)

    def __init__(self, auth_key, days):
        super().__init__()
        self.auth_key = auth_key
        self.days = days

    def run(self):
        try:
            self.ok.emit(fetch_threatfox_recent(self.auth_key, self.days))
        except Exception as e:
            self.error.emit(str(e))


class BuildThread(QThread):
    log_line = pyqtSignal(str)
    finished_build = pyqtSignal(bool, str)

    def __init__(self, cmd):
        super().__init__()
        self.cmd = cmd

    def run(self):
        try:
            proc = subprocess.Popen(self.cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                     text=True, encoding="utf-8", errors="replace")
            for line in proc.stdout:
                self.log_line.emit(line.rstrip())
            proc.wait()
            if proc.returncode == 0:
                self.finished_build.emit(True, "Build thành công.")
            else:
                self.finished_build.emit(False, f"Nuitka thoát với mã lỗi {proc.returncode}.")
        except Exception as e:
            self.finished_build.emit(False, str(e))


class ATTTAnalysisTool(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"HỆ THỐNG TỔNG HỢP & PHÂN TÍCH BIÊN BẢN KIỂM TRA ATTT - v{APP_VERSION_TH}")
        self.resize(1280, 780)
        self.data_list = []
        self.cfg = load_config()
        self.init_ui()

    def init_ui(self):
        main_widget = QWidget()
        self.setCentralWidget(main_widget)
        main_layout = QVBoxLayout(main_widget)

        # Top Bar: Chọn thư mục
        top_group = QGroupBox("Chọn Thư Mục Chứa Biên Bản Kiểm Tra (.docx)")
        top_layout = QHBoxLayout(top_group)

        self.txt_path = QLineEdit()
        self.txt_path.setPlaceholderText("Đường dẫn thư mục chứa các biên bản .docx...")
        self.txt_path.setReadOnly(True)
        top_layout.addWidget(self.txt_path)

        btn_browse = QPushButton("📁 Chọn Thư Mục")
        btn_browse.setStyleSheet("background-color: #0288D1; color: white; font-weight: bold; padding: 6px 15px;")
        btn_browse.clicked.connect(self.browse_folder)
        top_layout.addWidget(btn_browse)

        self.btn_run = QPushButton("🚀 Bắt Đầu Đọc & Tổng Hợp")
        self.btn_run.setStyleSheet("background-color: #2E7D32; color: white; font-weight: bold; padding: 6px 15px;")
        self.btn_run.clicked.connect(self.start_processing)
        top_layout.addWidget(self.btn_run)

        main_layout.addWidget(top_group)

        # Progress Bar
        self.progress_bar = QProgressBar()
        self.progress_bar.setValue(0)
        self.progress_bar.setTextVisible(True)
        main_layout.addWidget(self.progress_bar)

        # Tabs
        self.tabs = QTabWidget()
        main_layout.addWidget(self.tabs)

        self.init_tab1_data()
        self.init_tab2_analysis()
        self.init_tab3_update()

    def init_tab1_data(self):
        tab1 = QWidget()
        layout = QVBoxLayout(tab1)

        # Thanh Lọc dữ liệu nhanh & Tìm kiếm
        filter_layout = QHBoxLayout()
        filter_layout.addWidget(QLabel("🔍 Tìm kiếm nhanh:"))
        self.txt_search = QLineEdit()
        self.txt_search.setPlaceholderText("Gõ tên cán bộ, tên máy, IP, tên file...")
        self.txt_search.textChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.txt_search)

        filter_layout.addWidget(QLabel("Lọc mức nguy cơ:"))
        self.cbo_filter_risk = QComboBox()
        self.cbo_filter_risk.addItems(["Tất cả", "Nguy cấp (Critical)", "Cao (High)", "Trung bình (Medium)", "An toàn"])
        self.cbo_filter_risk.currentTextChanged.connect(self.apply_filter)
        filter_layout.addWidget(self.cbo_filter_risk)

        layout.addLayout(filter_layout)

        # Bảng hiển thị
        self.table = QTableWidget()
        self.table.setColumnCount(16)
        self.table.setHorizontalHeaderLabels([
            "STT", "Tên File", "Cán Bộ QL", "Tên Máy", "HĐH", "IP",
            "CPU", "RAM", "Ổ Cứng", "Số Lỗ Hổng", "Mức Nguy Cơ", "Mã Độc", "Lịch Sử USB",
            "Mật Khẩu", "Phân Loại", "Toàn Vẹn"
        ])
        self.table.setAlternatingRowColors(True)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.table.horizontalHeader().setStretchLastSection(True)
        self.table.doubleClicked.connect(self.open_selected_docx)
        layout.addWidget(self.table)

        # Bottom Bar: Xuất Excel
        bottom_layout = QHBoxLayout()
        self.lbl_status = QLabel("Nhấp đúp chuột vào bất kỳ hàng nào để mở trực tiếp file Word gốc.")
        self.lbl_status.setStyleSheet("color: #666; font-style: italic;")
        bottom_layout.addWidget(self.lbl_status)
        bottom_layout.addStretch()

        btn_export = QPushButton("📊 Xuất Báo Cáo Excel Tổng Hợp")
        btn_export.setStyleSheet("background-color: #1F4E78; color: white; font-weight: bold; padding: 8px 18px;")
        btn_export.clicked.connect(self.export_excel)
        bottom_layout.addWidget(btn_export)

        layout.addLayout(bottom_layout)
        self.tabs.addTab(tab1, "📋 1. Dữ Liệu Báo Cáo & Kết Xuất")

    def init_tab2_analysis(self):
        tab2 = QWidget()
        layout = QVBoxLayout(tab2)

        # Dashboard KPI
        kpi_group = QGroupBox("Tổng Quan Nguy Cơ Toàn Hệ Thống")
        kpi_layout = QGridLayout(kpi_group)

        self.lbl_kpi_total = QLabel("0\nMáy Kiểm Tra")
        self.lbl_kpi_critical = QLabel("0\nNguy Cấp")
        self.lbl_kpi_vuln = QLabel("0\nTổng Lỗ Hổng")
        self.lbl_kpi_usb = QLabel("0\nLượt Cắm USB")

        for lbl, color in zip([self.lbl_kpi_total, self.lbl_kpi_critical, self.lbl_kpi_vuln, self.lbl_kpi_usb],
                               ["#1976D2", "#D32F2F", "#F57C00", "#7B1FA2"]):
            lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
            lbl.setFont(QFont("Arial", 11, QFont.Weight.Bold))
            lbl.setStyleSheet(f"background-color: {color}; color: white; border-radius: 6px; padding: 10px;")

        kpi_layout.addWidget(self.lbl_kpi_total, 0, 0)
        kpi_layout.addWidget(self.lbl_kpi_critical, 0, 1)
        kpi_layout.addWidget(self.lbl_kpi_vuln, 0, 2)
        kpi_layout.addWidget(self.lbl_kpi_usb, 0, 3)

        layout.addWidget(kpi_group)

        # Splitter phân tích
        splitter = QSplitter(Qt.Orientation.Horizontal)

        vuln_group = QGroupBox("Top Lỗ Hổng Bảo Mật & Cảnh Báo Nguy Hiểm")
        vuln_layout = QVBoxLayout(vuln_group)
        self.txt_top_vuln = QTextEdit()
        self.txt_top_vuln.setReadOnly(True)
        self.txt_top_vuln.setFont(QFont("Consolas", 10))
        vuln_layout.addWidget(self.txt_top_vuln)
        splitter.addWidget(vuln_group)

        usb_group = QGroupBox("Danh Sách Thiết Bị Ngoại Vi / USB Đã Ghi Nhận")
        usb_layout = QVBoxLayout(usb_group)
        self.txt_usb_analysis = QTextEdit()
        self.txt_usb_analysis.setReadOnly(True)
        self.txt_usb_analysis.setFont(QFont("Consolas", 10))
        usb_layout.addWidget(self.txt_usb_analysis)
        splitter.addWidget(usb_group)

        layout.addWidget(splitter)
        self.tabs.addTab(tab2, "🔍 2. Phân Tích Lỗ Hổng & Cảnh Báo ATTT")

    def init_tab3_update(self):
        tab3 = QWidget()
        layout = QVBoxLayout(tab3)

        # ---- Khối đường dẫn (nhớ lại qua các lần mở, lưu trong tong_hop_config.json) ----
        path_group = QGroupBox("Đường dẫn (chỉ cần chọn 1 lần, tự nhớ cho lần sau)")
        path_grid = QGridLayout(path_group)
        self._duong_dan_edits = {}
        cac_duong_dan = [
            ("duong_dan_cve", "File CVE Windows (windows_vulnerabilities.txt)", False),
            ("duong_dan_ioc", "File IOC mã độc (malware_signatures.txt)", False),
            ("duong_dan_script", "Mã nguồn auto_fill_bien_ban.py", False),
            ("duong_dan_mau_docx", "File mẫu biên bản (mau_bien_ban.docx)", False),
            ("duong_dan_icon", "File icon .ico (không bắt buộc)", False),
            ("duong_dan_python", "Python có cài Nuitka (python.exe)", False),
            ("thu_muc_xuat", "Thư mục xuất bản build (có thể chọn thẳng USB)", True),
        ]
        for row, (key, nhan, la_thu_muc) in enumerate(cac_duong_dan):
            path_grid.addWidget(QLabel(nhan + ":"), row, 0)
            edit = QLineEdit(self.cfg.get(key, ""))
            edit.setReadOnly(True)
            path_grid.addWidget(edit, row, 1)
            btn = QPushButton("Chọn...")
            btn.clicked.connect(lambda _=False, k=key, e=edit, d=la_thu_muc: self._chon_duong_dan(k, e, d))
            path_grid.addWidget(btn, row, 2)
            self._duong_dan_edits[key] = edit
        layout.addWidget(path_group)

        splitter3 = QSplitter(Qt.Orientation.Horizontal)

        # ---- Khối CVE (MSRC) ----
        cve_group = QGroupBox("1. Cập nhật CVE Windows (nguồn: MSRC Security Update Guide)")
        cve_layout = QVBoxLayout(cve_group)
        cve_top = QHBoxLayout()
        cve_top.addWidget(QLabel("Tháng đến:"))
        self.txt_cve_thang = QLineEdit(datetime.now().strftime("%Y-%b"))
        self.txt_cve_thang.setMaximumWidth(100)
        self.txt_cve_thang.setToolTip("Tháng mốc (vd 2026-Sep). Công cụ sẽ tải lùi về trước theo số tháng bên cạnh.")
        cve_top.addWidget(self.txt_cve_thang)
        cve_top.addWidget(QLabel("Số tháng:"))
        self.spin_cve_so_thang = QSpinBox()
        self.spin_cve_so_thang.setRange(1, 24)
        self.spin_cve_so_thang.setValue(self.cfg.get("cve_so_thang", 12))
        self.spin_cve_so_thang.setMaximumWidth(56)
        self.spin_cve_so_thang.setToolTip("Tải bao nhiêu tháng gần nhất trong một lần (gộp và khử trùng).")
        cve_top.addWidget(self.spin_cve_so_thang)
        self.btn_cve_tai = QPushButton("🔄 Tải CVE")
        self.btn_cve_tai.clicked.connect(self.fetch_cve)
        cve_top.addWidget(self.btn_cve_tai)
        cve_top.addStretch()
        cve_layout.addLayout(cve_top)

        self.chk_cve_chi_khai_thac = QCheckBox("Chỉ lấy lỗ hổng ĐÃ bị khai thác thực tế (MSRC/CISA KEV) - danh sách gọn, đúng trọng tâm")
        self.chk_cve_chi_khai_thac.setChecked(self.cfg.get("cve_chi_khai_thac", False))
        cve_layout.addWidget(self.chk_cve_chi_khai_thac)

        self.table_cve = QTableWidget()
        self.table_cve.setColumnCount(5)
        self.table_cve.setHorizontalHeaderLabels(["Chọn", "CVE", "Khai thác", "Mức độ", "KB / Build đã vá"])
        self.table_cve.setAlternatingRowColors(True)
        self.table_cve.horizontalHeader().setStretchLastSection(True)
        cve_layout.addWidget(self.table_cve)

        cve_bottom = QHBoxLayout()
        self.lbl_cve_status = QLabel("")
        cve_bottom.addWidget(self.lbl_cve_status)
        cve_bottom.addStretch()
        btn_cve_save = QPushButton("💾 Lưu các dòng đã chọn vào file CVE")
        btn_cve_save.clicked.connect(self.save_cve_to_file)
        cve_bottom.addWidget(btn_cve_save)
        cve_layout.addLayout(cve_bottom)
        splitter3.addWidget(cve_group)

        # ---- Khối IOC (ThreatFox) ----
        ioc_group = QGroupBox("2. Cập nhật IOC mã độc (nguồn: ThreatFox - abuse.ch)")
        ioc_layout = QVBoxLayout(ioc_group)
        ioc_top = QHBoxLayout()
        ioc_top.addWidget(QLabel("Auth-Key:"))
        self.txt_threatfox_key = QLineEdit(self.cfg.get("threatfox_auth_key", ""))
        self.txt_threatfox_key.setEchoMode(QLineEdit.EchoMode.Password)
        ioc_top.addWidget(self.txt_threatfox_key)
        ioc_top.addWidget(QLabel("Số ngày:"))
        self.spin_ioc_days = QSpinBox()
        self.spin_ioc_days.setRange(1, 30)
        self.spin_ioc_days.setValue(7)
        self.spin_ioc_days.setMaximumWidth(56)
        ioc_top.addWidget(self.spin_ioc_days)
        self.btn_ioc_tai = QPushButton("🔄 Tải IOC")
        self.btn_ioc_tai.clicked.connect(self.fetch_ioc)
        ioc_top.addWidget(self.btn_ioc_tai)
        ioc_layout.addLayout(ioc_top)

        ioc_loc = QHBoxLayout()
        ioc_loc.addWidget(QLabel("Độ tin cậy ≥"))
        self.spin_ioc_confidence = QSpinBox()
        self.spin_ioc_confidence.setRange(0, 100)
        self.spin_ioc_confidence.setValue(self.cfg.get("ioc_min_confidence", 75))
        self.spin_ioc_confidence.setSuffix(" %")
        self.spin_ioc_confidence.setMaximumWidth(70)
        ioc_loc.addWidget(self.spin_ioc_confidence)
        ioc_loc.addWidget(QLabel("Loại bỏ IP/IOC cũ hơn"))
        self.spin_ioc_maxage = QSpinBox()
        self.spin_ioc_maxage.setRange(0, 365)
        self.spin_ioc_maxage.setValue(self.cfg.get("ioc_max_age", 90))
        self.spin_ioc_maxage.setSuffix(" ngày")
        self.spin_ioc_maxage.setMaximumWidth(90)
        self.spin_ioc_maxage.setToolTip("0 = không lọc theo thời gian. IP máy chủ điều khiển đổi nhanh nên IP cũ dễ gây báo nhầm.")
        ioc_loc.addWidget(self.spin_ioc_maxage)
        ioc_loc.addStretch()
        ioc_layout.addLayout(ioc_loc)

        self.table_ioc = QTableWidget()
        self.table_ioc.setColumnCount(4)
        self.table_ioc.setHorizontalHeaderLabels(["Chọn", "Họ mã độc", "Loại/Giá trị", "Ngày phát hiện"])
        self.table_ioc.setAlternatingRowColors(True)
        self.table_ioc.horizontalHeader().setStretchLastSection(True)
        ioc_layout.addWidget(self.table_ioc)

        ioc_bottom = QHBoxLayout()
        self.lbl_ioc_status = QLabel("")
        ioc_bottom.addWidget(self.lbl_ioc_status)
        ioc_bottom.addStretch()
        btn_ioc_save = QPushButton("💾 Lưu các dòng đã chọn vào file IOC")
        btn_ioc_save.clicked.connect(self.save_ioc_to_file)
        ioc_bottom.addWidget(btn_ioc_save)
        ioc_layout.addLayout(ioc_bottom)
        splitter3.addWidget(ioc_group)

        layout.addWidget(splitter3)

        # ---- Khối Build ----
        build_group = QGroupBox("3. Build lại auto_fill_bien_ban.exe (kèm dữ liệu mới nhất)")
        build_layout = QVBoxLayout(build_group)

        thongtin_grid = QGridLayout()
        thongtin_grid.addWidget(QLabel("Tên công ty/đơn vị (Company):"), 0, 0)
        self.txt_company_name = QLineEdit(self.cfg.get("company_name", "MDH"))
        self.txt_company_name.editingFinished.connect(
            lambda: (self.cfg.__setitem__("company_name", self.txt_company_name.text().strip()),
                     save_config(self.cfg)))
        thongtin_grid.addWidget(self.txt_company_name, 0, 1)

        thongtin_grid.addWidget(QLabel("Tên sản phẩm (Product):"), 0, 2)
        self.txt_product_name = QLineEdit(self.cfg.get("product_name", "Kiểm tra ANATTT"))
        self.txt_product_name.editingFinished.connect(
            lambda: (self.cfg.__setitem__("product_name", self.txt_product_name.text().strip()),
                     save_config(self.cfg)))
        thongtin_grid.addWidget(self.txt_product_name, 0, 3)

        thongtin_grid.addWidget(QLabel("Phiên bản hiện tại:"), 1, 0)
        self.lbl_phien_ban_hien_tai = QLabel("(chưa rõ - chọn mã nguồn trước)")
        self.lbl_phien_ban_hien_tai.setStyleSheet("font-weight: bold; color: #1976D2;")
        thongtin_grid.addWidget(self.lbl_phien_ban_hien_tai, 1, 1)

        self.chk_tu_tang_phien_ban = QCheckBox("Tự tăng phiên bản (vd 1.6 -> 1.7) mỗi lần build")
        self.chk_tu_tang_phien_ban.setChecked(self.cfg.get("tu_tang_phien_ban", True))
        self.chk_tu_tang_phien_ban.toggled.connect(
            lambda checked: (self.cfg.__setitem__("tu_tang_phien_ban", checked), save_config(self.cfg)))
        thongtin_grid.addWidget(self.chk_tu_tang_phien_ban, 1, 2, 1, 2)
        build_layout.addLayout(thongtin_grid)

        self.btn_build = QPushButton("🔨 Build lại auto_fill_bien_ban.exe")
        self.btn_build.setStyleSheet("background-color: #7B1FA2; color: white; font-weight: bold; padding: 8px;")
        self.btn_build.clicked.connect(self.run_build)
        build_layout.addWidget(self.btn_build)
        self.txt_build_log = QTextEdit()
        self.txt_build_log.setReadOnly(True)
        self.txt_build_log.setFont(QFont("Consolas", 9))
        self.txt_build_log.setMaximumHeight(160)
        build_layout.addWidget(self.txt_build_log)
        layout.addWidget(build_group)

        self.tabs.addTab(tab3, "🔧 3. Cập nhật dữ liệu & Build")
        self._cap_nhat_nhan_phien_ban()

    # ---------------- Đường dẫn ----------------
    def _chon_duong_dan(self, key, edit, la_thu_muc):
        if la_thu_muc:
            gia_tri = QFileDialog.getExistingDirectory(self, "Chọn thư mục")
        else:
            loc = {
                "duong_dan_cve": "Text Files (*.txt)",
                "duong_dan_ioc": "Text Files (*.txt)",
                "duong_dan_script": "Python Files (*.py)",
                "duong_dan_mau_docx": "Word Files (*.docx)",
                "duong_dan_icon": "Icon Files (*.ico)",
                "duong_dan_python": "Executable (*.exe)",
            }.get(key, "All Files (*)")
            gia_tri, _ = QFileDialog.getOpenFileName(self, "Chọn file", "", loc)
        if gia_tri:
            edit.setText(gia_tri)
            self.cfg[key] = gia_tri
            save_config(self.cfg)
            if key == "duong_dan_script":
                self._cap_nhat_nhan_phien_ban()

    def _cap_nhat_nhan_phien_ban(self):
        duong_dan_script = self._duong_dan_edits["duong_dan_script"].text().strip()
        if not duong_dan_script or not os.path.exists(duong_dan_script):
            self.lbl_phien_ban_hien_tai.setText("(chưa rõ - chọn mã nguồn trước)")
            return
        phien_ban = doc_phien_ban_script(duong_dan_script)
        self.lbl_phien_ban_hien_tai.setText(phien_ban or "(không đọc được APP_VERSION)")

    # ---------------- CVE (MSRC) ----------------
    def fetch_cve(self):
        thang = self.txt_cve_thang.text().strip()
        if not thang:
            QMessageBox.warning(self, "Thiếu thông tin", "Nhập tháng mốc cần tải (vd 2026-Sep).")
            return
        so_thang = self.spin_cve_so_thang.value()
        months = _danh_sach_thang_lui(thang, so_thang)
        if not months:
            QMessageBox.warning(self, "Sai định dạng tháng",
                                 "Tháng mốc phải dạng 'YYYY-Mon', vd 2026-Sep hoặc 2026-09.")
            return
        self.cfg["cve_so_thang"] = so_thang
        self.cfg["cve_chi_khai_thac"] = self.chk_cve_chi_khai_thac.isChecked()
        save_config(self.cfg)
        self.btn_cve_tai.setEnabled(False)
        self.lbl_cve_status.setText(f"Đang tải {len(months)} tháng từ MSRC...")
        self._cve_thread = CveFetchThread(months, self.chk_cve_chi_khai_thac.isChecked())
        self._cve_thread.tien_trinh.connect(self.lbl_cve_status.setText)
        self._cve_thread.ok.connect(self._tren_cve_tai_xong)
        self._cve_thread.error.connect(self._tren_cve_loi)
        self._cve_thread.start()

    def _tren_cve_loi(self, err):
        self.btn_cve_tai.setEnabled(True)
        self.lbl_cve_status.setText("Lỗi tải CVE.")
        QMessageBox.critical(self, "Lỗi tải CVE", f"Không tải được dữ liệu từ MSRC:\n{err}")

    def _tren_cve_tai_xong(self, ket_qua):
        self.btn_cve_tai.setEnabled(True)
        self._cve_data = ket_qua
        self.table_cve.setRowCount(0)
        da_co = self._doc_cve_id_da_co()
        so_khai_thac = 0
        for item in ket_qua:
            row = self.table_cve.rowCount()
            self.table_cve.insertRow(row)
            chk = QTableWidgetItem()
            chk.setFlags(Qt.ItemFlag.ItemIsUserCheckable | Qt.ItemFlag.ItemIsEnabled)
            # Bỏ tick sẵn các CVE đã có trong file - đỡ phải tự nhớ để tránh trùng
            chk.setCheckState(Qt.CheckState.Unchecked if item["cve"] in da_co else Qt.CheckState.Checked)
            self.table_cve.setItem(row, 0, chk)
            self.table_cve.setItem(row, 1, QTableWidgetItem(item["cve"] +
                                    (" (đã có)" if item["cve"] in da_co else "")))
            da_kt = item.get("exploited") or item.get("kev")
            if da_kt:
                so_khai_thac += 1
            kt_item = QTableWidgetItem("⚠ CÓ" if da_kt else "")
            if da_kt:
                kt_item.setForeground(QColor("#C62828"))
                kt_item.setFont(QFont("Arial", 9, QFont.Weight.Bold))
            self.table_cve.setItem(row, 2, kt_item)
            self.table_cve.setItem(row, 3, QTableWidgetItem(item["severity"]))
            build_txt = "; ".join(item.get("fixed_builds") or [])
            kb_txt = ", ".join(item["kbs"]) or "(chưa rõ KB)"
            self.table_cve.setItem(row, 4, QTableWidgetItem(
                kb_txt + (f"  |  build: {build_txt}" if build_txt else "")))
        self.lbl_cve_status.setText(
            f"Tìm thấy {len(ket_qua)} CVE Windows (trong đó {so_khai_thac} đã bị khai thác). "
            f"Xem lại rồi bấm Lưu.")

    def _doc_cve_id_da_co(self):
        duong_dan = self._duong_dan_edits["duong_dan_cve"].text().strip()
        da_co = set()
        if duong_dan and os.path.exists(duong_dan):
            try:
                with open(duong_dan, "r", encoding="utf-8") as f:
                    for dong in f:
                        dong = dong.strip()
                        if dong and not dong.startswith("#") and "|" in dong:
                            da_co.add(dong.split("|", 1)[0].strip())
            except Exception:
                pass
        return da_co

    def save_cve_to_file(self):
        duong_dan = self._duong_dan_edits["duong_dan_cve"].text().strip()
        if not duong_dan:
            QMessageBox.warning(self, "Thiếu đường dẫn", "Chọn file CVE (windows_vulnerabilities.txt) trước.")
            return
        if not hasattr(self, "_cve_data"):
            QMessageBox.warning(self, "Chưa có dữ liệu", "Bấm 'Tải CVE tháng này' trước.")
            return
        da_co = self._doc_cve_id_da_co()
        dong_moi = []
        for row in range(self.table_cve.rowCount()):
            if self.table_cve.item(row, 0).checkState() != Qt.CheckState.Checked:
                continue
            item = self._cve_data[row]
            if item["cve"] in da_co:
                continue
            kb_text = ";".join(item["kbs"])
            build_text = ";".join(item.get("fixed_builds") or [])
            ghi_chu = f"Tự động lấy từ MSRC ngày {datetime.now().strftime('%d/%m/%Y')}"
            if item.get("exploited") or item.get("kev"):
                ghi_chu = "[ĐÃ BỊ KHAI THÁC] " + ghi_chu
            # Cột 5 = 'build tối thiểu đã vá' (dạng 19045.4291;22631.3447) để so theo build.UBR
            dong_moi.append(f"{item['cve']}|{item['title']}|{item['severity']}|{kb_text}|{build_text}|{ghi_chu}")
        if not dong_moi:
            QMessageBox.information(self, "Không có gì để lưu",
                                     "Không có CVE mới nào được chọn (có thể đã có sẵn trong file).")
            return
        try:
            with open(duong_dan, "a", encoding="utf-8") as f:
                f.write(f"\n# --- Bổ sung tự động từ MSRC ngày {datetime.now().strftime('%d/%m/%Y')} ---\n")
                for dong in dong_moi:
                    f.write(dong + "\n")
            _cap_nhat_data_version(duong_dan)
        except Exception as e:
            QMessageBox.critical(self, "Lỗi ghi file", str(e))
            return
        QMessageBox.information(self, "Đã lưu", f"Đã thêm {len(dong_moi)} CVE mới vào file.")
        self.lbl_cve_status.setText(f"Đã lưu {len(dong_moi)} CVE mới.")

    # ---------------- IOC (ThreatFox) ----------------
    def fetch_ioc(self):
        key = self.txt_threatfox_key.text().strip()
        if not key:
            QMessageBox.warning(self, "Thiếu Auth-Key",
                                 "Nhập Auth-Key ThreatFox (đăng ký miễn phí tại https://auth.abuse.ch/).")
            return
        self.cfg["threatfox_auth_key"] = key
        self.cfg["ioc_min_confidence"] = self.spin_ioc_confidence.value()
        self.cfg["ioc_max_age"] = self.spin_ioc_maxage.value()
        save_config(self.cfg)
        self.btn_ioc_tai.setEnabled(False)
        self.lbl_ioc_status.setText("Đang tải từ ThreatFox...")
        self._ioc_thread = IocFetchThread(key, self.spin_ioc_days.value())
        self._ioc_thread.ok.connect(self._tren_ioc_tai_xong)
        self._ioc_thread.error.connect(self._tren_ioc_loi)
        self._ioc_thread.start()

    def _tren_ioc_loi(self, err):
        self.btn_ioc_tai.setEnabled(True)
        self.lbl_ioc_status.setText("Lỗi tải IOC.")
        QMessageBox.critical(self, "Lỗi tải IOC", f"Không tải được dữ liệu từ ThreatFox:\n{err}")

    def _tren_ioc_tai_xong(self, ket_qua_tho):
        self.btn_ioc_tai.setEnabled(True)
        da_co = self._doc_ioc_da_co()
        min_conf = self.spin_ioc_confidence.value()
        max_age = self.spin_ioc_maxage.value()
        self._ioc_data = []
        da_thay = set()
        for item in ket_qua_tho:
            dong = quy_doi_ioc_threatfox(item, min_conf, max_age)
            if dong and dong not in da_thay:  # khử trùng ngay khi tải
                da_thay.add(dong)
                self._ioc_data.append({"dong": dong, "ho": item.get("malware_printable") or
                                        item.get("malware") or "?", "ngay": item.get("first_seen", "")})
        self.table_ioc.setRowCount(0)
        for item in self._ioc_data:
            row = self.table_ioc.rowCount()
            self.table_ioc.insertRow(row)
            chk = QTableWidgetItem()
            chk.setFlags(Qt.ItemFlag.ItemIsUserCheckable | Qt.ItemFlag.ItemIsEnabled)
            chk.setCheckState(Qt.CheckState.Unchecked if item["dong"] in da_co else Qt.CheckState.Checked)
            self.table_ioc.setItem(row, 0, chk)
            self.table_ioc.setItem(row, 1, QTableWidgetItem(item["ho"]))
            self.table_ioc.setItem(row, 2, QTableWidgetItem(item["dong"] +
                                    (" (đã có)" if item["dong"] in da_co else "")))
            self.table_ioc.setItem(row, 3, QTableWidgetItem(item["ngay"]))
        self.lbl_ioc_status.setText(f"Tìm thấy {len(self._ioc_data)} IOC tương thích định dạng hiện tại "
                                     f"(trong tổng số {len(ket_qua_tho)} IOC ThreatFox trả về). Xem lại rồi bấm Lưu.")

    def _doc_ioc_da_co(self):
        duong_dan = self._duong_dan_edits["duong_dan_ioc"].text().strip()
        da_co = set()
        if duong_dan and os.path.exists(duong_dan):
            try:
                with open(duong_dan, "r", encoding="utf-8") as f:
                    for dong in f:
                        dong = dong.strip()
                        if dong and not dong.startswith("#"):
                            da_co.add(dong)
            except Exception:
                pass
        return da_co

    def save_ioc_to_file(self):
        duong_dan = self._duong_dan_edits["duong_dan_ioc"].text().strip()
        if not duong_dan:
            QMessageBox.warning(self, "Thiếu đường dẫn", "Chọn file IOC (malware_signatures.txt) trước.")
            return
        if not hasattr(self, "_ioc_data"):
            QMessageBox.warning(self, "Chưa có dữ liệu", "Bấm 'Tải IOC mới' trước.")
            return
        da_co = self._doc_ioc_da_co()
        theo_ho = {}
        for row in range(self.table_ioc.rowCount()):
            if self.table_ioc.item(row, 0).checkState() != Qt.CheckState.Checked:
                continue
            item = self._ioc_data[row]
            if item["dong"] in da_co:
                continue
            theo_ho.setdefault(item["ho"], []).append(item["dong"])
        if not theo_ho:
            QMessageBox.information(self, "Không có gì để lưu",
                                     "Không có IOC mới nào được chọn (có thể đã có sẵn trong file).")
            return
        try:
            with open(duong_dan, "a", encoding="utf-8") as f:
                f.write(f"\n# --- Bổ sung tự động từ ThreatFox ngày {datetime.now().strftime('%d/%m/%Y')} ---\n")
                for ho, dong_list in theo_ho.items():
                    f.write(f"\n# --- {ho} ---\n")
                    for dong in dong_list:
                        f.write(dong + "\n")
            _cap_nhat_data_version(duong_dan)
        except Exception as e:
            QMessageBox.critical(self, "Lỗi ghi file", str(e))
            return
        tong = sum(len(v) for v in theo_ho.values())
        QMessageBox.information(self, "Đã lưu", f"Đã thêm {tong} IOC mới ({len(theo_ho)} họ mã độc) vào file.")
        self.lbl_ioc_status.setText(f"Đã lưu {tong} IOC mới.")

    # ---------------- Build ----------------
    def run_build(self):
        can_thiet = ["duong_dan_script", "duong_dan_mau_docx", "duong_dan_python", "thu_muc_xuat"]
        thieu = [k for k in can_thiet if not self._duong_dan_edits[k].text().strip()]
        if thieu:
            QMessageBox.warning(self, "Thiếu đường dẫn",
                                 "Cần chọn đủ: mã nguồn, file mẫu docx, Python (có Nuitka), và thư mục xuất.")
            return
        duong_dan_script = self._duong_dan_edits["duong_dan_script"].text().strip()
        duong_dan_python = self._duong_dan_edits["duong_dan_python"].text().strip()
        thu_muc_xuat = self._duong_dan_edits["thu_muc_xuat"].text().strip()

        self.txt_build_log.clear()

        # Kiểm tra dữ liệu TRƯỚC khi build - tránh mang một bản .exe hỏng đi chạy trên hàng trăm máy
        loi, canh_bao = self._kiem_tra_truoc_build()
        for c in canh_bao:
            self.txt_build_log.append(f"⚠ {c}")
        if loi:
            self.txt_build_log.append("❌ DỪNG BUILD - lỗi nghiêm trọng:")
            for l in loi:
                self.txt_build_log.append(f"   • {l}")
            QMessageBox.critical(self, "Không thể build",
                                 "Dữ liệu chưa hợp lệ, cần sửa trước khi build:\n\n- " + "\n- ".join(loi))
            return

        # Tự tăng phiên bản (vd 1.6 -> 1.7) TRƯỚC khi build, nếu được tick - để mã nguồn
        # và bản .exe xuất ra luôn khớp nhau, không phải tự sửa tay APP_VERSION mỗi lần.
        if self.chk_tu_tang_phien_ban.isChecked():
            phien_ban_moi = tang_phien_ban_script(duong_dan_script)
            if phien_ban_moi:
                self.txt_build_log.append(f"Đã tự tăng phiên bản trong mã nguồn lên {phien_ban_moi}.")
                self._cap_nhat_nhan_phien_ban()
            else:
                self.txt_build_log.append("⚠ Không tự tăng được phiên bản (không tìm thấy dòng "
                                           "APP_VERSION trong mã nguồn) - vẫn tiếp tục build.")

        ten_cong_ty = self.txt_company_name.text().strip() or "MDH"
        ten_san_pham = self.txt_product_name.text().strip() or "Kiem tra ANATTT"

        cmd = [duong_dan_python, "-m", "nuitka", "--onefile",
               "--windows-console-mode=disable", "--assume-yes-for-downloads", "--no-progressbar",
               f"--output-dir={thu_muc_xuat}", "--output-filename=auto_fill_bien_ban.exe",
               f"--company-name={ten_cong_ty}", f"--product-name={ten_san_pham}"]

        # Nuitka bắt buộc phải có file-version/product-version khi đã cho company-name/product-name
        # (thiếu sẽ báo lỗi "FATAL: ... version information is given") - tự đọc đúng số phiên bản
        # (2 phần, vd "1.6" - đã tăng ở bước trên nếu có tick) rồi đổi sang dạng 4 số Windows cần
        # ("1.6" -> "1.6.0.0"); đây là yêu cầu bắt buộc của định dạng version resource trên Windows,
        # nên trong hộp thoại Properties của file .exe vẫn sẽ hiện đủ 4 số, số 2 phần (X.Y) mới là
        # số "chính thức" mà công cụ này theo dõi/tự tăng.
        phien_ban_4so = "1.0.0.0"
        phien_ban_hien = doc_phien_ban_script(duong_dan_script)
        if phien_ban_hien:
            phan = (phien_ban_hien.split(".") + ["0", "0", "0"])[:4]
            phien_ban_4so = ".".join(phan)
        cmd.append(f"--file-version={phien_ban_4so}")
        cmd.append(f"--product-version={phien_ban_4so}")

        icon = self._duong_dan_edits["duong_dan_icon"].text().strip()
        if icon:
            cmd.append(f"--windows-icon-from-ico={icon}")

        # Icon + file dữ liệu CVE/IOC/mẫu docx đều phải đóng gói cùng .exe (--include-data-files),
        # nếu không auto_fill_bien_ban.py sẽ không đọc lại được lúc chạy (xem ghi chú icon khay
        # từng gặp ở máy con - bài học tương tự áp dụng ở đây).
        for key in ("duong_dan_cve", "duong_dan_ioc", "duong_dan_mau_docx", "duong_dan_icon"):
            f = self._duong_dan_edits[key].text().strip()
            if f and os.path.exists(f):
                cmd.append(f"--include-data-files={f}={os.path.basename(f)}")

        # Nhúng KHOÁ TOÀN VẸN (bien_ban_key.dat): auto_fill dùng khoá này ký biên bản, công cụ tổng hợp
        # dùng chính khoá này xác minh. Sinh ngẫu nhiên 1 lần rồi lưu trong cấu hình để các đợt build sau
        # dùng cùng khoá (nếu không, biên bản đợt trước sẽ báo 'chữ ký sai').
        if not self.cfg.get("integrity_key"):
            import secrets
            self.cfg["integrity_key"] = secrets.token_hex(32)
            save_config(self.cfg)
            self.txt_build_log.append("Đã sinh khoá toàn vẹn mới cho đơn vị (lưu trong cấu hình).")
        try:
            key_file = os.path.join(os.path.dirname(CONFIG_PATH), "bien_ban_key.dat")
            with open(key_file, "w", encoding="utf-8") as kf:
                kf.write(self.cfg["integrity_key"])
            cmd.append(f"--include-data-files={key_file}=bien_ban_key.dat")
        except Exception as e:
            self.txt_build_log.append(f"⚠ Không ghi được khoá toàn vẹn: {e} - biên bản sẽ dùng khoá mặc định.")

        cmd.append(duong_dan_script)

        self.txt_build_log.append("Lệnh build: " + " ".join(cmd))
        # Lưu ngữ cảnh để ghi phiếu build (manifest) khi build xong
        self._build_context = {
            "thu_muc_xuat": thu_muc_xuat,
            "phien_ban": phien_ban_hien or "?",
            "cong_ty": ten_cong_ty,
            "san_pham": ten_san_pham,
        }
        self.btn_build.setEnabled(False)
        self._build_thread = BuildThread(cmd)
        self._build_thread.log_line.connect(self.txt_build_log.append)
        self._build_thread.finished_build.connect(self._tren_build_xong)
        self._build_thread.start()

    def _kiem_tra_truoc_build(self):
        """Kiểm tra file mẫu docx, file CVE, file IOC và số phiên bản trước khi build.
        Trả về (danh_sách_lỗi_nghiêm_trọng, danh_sách_cảnh_báo)."""
        loi, canh_bao = [], []
        E = self._duong_dan_edits

        # 1. Mã nguồn có APP_VERSION
        script = E["duong_dan_script"].text().strip()
        if not doc_phien_ban_script(script):
            loi.append("Không đọc được APP_VERSION trong mã nguồn auto_fill_bien_ban.py.")

        # 2. File mẫu biên bản đúng cấu trúc mà auto_fill cần (bảng ≥ 9 dòng, ≥ 14 đoạn văn)
        docx_path = E["duong_dan_mau_docx"].text().strip()
        try:
            d = docx.Document(docx_path)
            if not d.tables or len(d.tables[0].rows) < 9:
                loi.append("File mẫu biên bản thiếu bảng thông tin (cần ≥ 9 dòng).")
            if len(d.paragraphs) < 14:
                loi.append("File mẫu biên bản thiếu các đoạn văn cần thiết (cần ≥ 14 đoạn).")
        except Exception as e:
            loi.append(f"Không mở được file mẫu biên bản: {e}")

        # 3. File CVE đọc được và có ít nhất 1 dòng hợp lệ
        cve_path = E["duong_dan_cve"].text().strip()
        if cve_path and os.path.exists(cve_path):
            n = 0
            try:
                with open(cve_path, encoding="utf-8") as f:
                    for dong in f:
                        dong = dong.strip()
                        if dong and not dong.startswith("#") and len(dong.split("|")) >= 4:
                            n += 1
            except Exception as e:
                loi.append(f"Không đọc được file CVE: {e}")
            if n == 0:
                loi.append("File CVE không có dòng lỗ hổng hợp lệ nào.")
            else:
                canh_bao.append(f"File CVE: {n} lỗ hổng. Dữ liệu cập nhật: {doc_ngay_du_lieu_th(cve_path)}.")
        else:
            canh_bao.append("Chưa chọn file CVE - biên bản sẽ bỏ qua bước đối chiếu lỗ hổng.")

        # 4. File IOC đọc được và có ít nhất 1 dấu hiệu
        ioc_path = E["duong_dan_ioc"].text().strip()
        if ioc_path and os.path.exists(ioc_path):
            n = 0
            try:
                with open(ioc_path, encoding="utf-8") as f:
                    for dong in f:
                        dong = dong.strip()
                        if dong and not dong.startswith("#") and ":" in dong:
                            n += 1
            except Exception as e:
                loi.append(f"Không đọc được file IOC: {e}")
            if n == 0:
                canh_bao.append("File IOC không có dấu hiệu nào - biên bản sẽ luôn báo 'không phát hiện mã độc'.")
            else:
                canh_bao.append(f"File IOC: {n} dấu hiệu. Dữ liệu cập nhật: {doc_ngay_du_lieu_th(ioc_path)}.")
        else:
            canh_bao.append("Chưa chọn file IOC - biên bản sẽ bỏ qua bước quét mã độc.")

        return loi, canh_bao

    def _ghi_manifest(self):
        """Ghi phiếu build (manifest.json) cạnh file .exe: phiên bản, ngày build, ngày dữ liệu,
        số CVE/IOC và mã băm SHA256 của .exe và các file dữ liệu - phục vụ truy vết/pháp lý."""
        ctx = getattr(self, "_build_context", None)
        if not ctx:
            return
        thu_muc = ctx["thu_muc_xuat"]
        exe_path = os.path.join(thu_muc, "auto_fill_bien_ban.exe")
        E = self._duong_dan_edits

        def _sha256(path):
            try:
                import hashlib
                h = hashlib.sha256()
                with open(path, "rb") as f:
                    for chunk in iter(lambda: f.read(1024 * 1024), b""):
                        h.update(chunk)
                return h.hexdigest()
            except Exception:
                return None

        def _dem(path, la_cve):
            if not path or not os.path.exists(path):
                return 0
            n = 0
            with open(path, encoding="utf-8") as f:
                for dong in f:
                    dong = dong.strip()
                    if dong and not dong.startswith("#") and (("|" in dong) if la_cve else (":" in dong)):
                        n += 1
            return n

        cve_path = E["duong_dan_cve"].text().strip()
        ioc_path = E["duong_dan_ioc"].text().strip()
        manifest = {
            "san_pham": ctx["san_pham"],
            "cong_ty": ctx["cong_ty"],
            "phien_ban_cong_cu": ctx["phien_ban"],
            "ngay_build": datetime.now().strftime("%d/%m/%Y %H:%M:%S"),
            "du_lieu_cve": {"ngay": doc_ngay_du_lieu_th(cve_path), "so_luong": _dem(cve_path, True),
                            "sha256": _sha256(cve_path) if cve_path else None},
            "du_lieu_ioc": {"ngay": doc_ngay_du_lieu_th(ioc_path), "so_luong": _dem(ioc_path, False),
                            "sha256": _sha256(ioc_path) if ioc_path else None},
            "sha256_exe": _sha256(exe_path),
        }
        try:
            with open(os.path.join(thu_muc, "manifest.json"), "w", encoding="utf-8") as f:
                json.dump(manifest, f, ensure_ascii=False, indent=2)
            self.txt_build_log.append(f"Đã ghi phiếu build: {os.path.join(thu_muc, 'manifest.json')}")
        except Exception as e:
            self.txt_build_log.append(f"⚠ Không ghi được manifest.json: {e}")

    def _tren_build_xong(self, ok, msg):
        self.btn_build.setEnabled(True)
        if ok:
            self._ghi_manifest()
            QMessageBox.information(self, "Build xong", msg + "\n\nĐã ghi kèm phiếu build manifest.json.")
        else:
            QMessageBox.critical(self, "Build lỗi", msg)

    def browse_folder(self):
        dir_ = QFileDialog.getExistingDirectory(self, "Chọn thư mục chứa các biên bản .docx")
        if dir_:
            self.txt_path.setText(dir_)

    def start_processing(self):
        folder = self.txt_path.text().strip()
        if not folder or not os.path.exists(folder):
            QMessageBox.warning(self, "Cảnh báo", "Vui lòng chọn thư mục chứa file biên bản trước!")
            return

        self.table.setRowCount(0)
        self.data_list.clear()
        self.btn_run.setEnabled(False)
        self.lbl_status.setText("Đang đọc và phân tích dữ liệu...")
        self.progress_bar.setValue(0)

        self._file_loi = []
        muc_do_cve = doc_muc_do_cve(self.cfg.get("duong_dan_cve", ""))
        self.thread = WorkerThread(folder, muc_do_cve, lay_integrity_key(self.cfg))
        self.thread.file_processed.connect(self.add_row_to_table)
        self.thread.file_error.connect(lambda f, e: self._file_loi.append(f"{f}: {e}"))
        self.thread.progress.connect(self.progress_bar.setValue)
        self.thread.finished.connect(self.process_finished)
        self.thread.start()

    def add_row_to_table(self, item):
        row = self.table.rowCount()
        self.table.insertRow(row)

        self.table.setItem(row, 0, QTableWidgetItem(str(row + 1)))
        self.table.setItem(row, 1, QTableWidgetItem(item["File_Name"]))
        self.table.setItem(row, 2, QTableWidgetItem(item["Can_Bo_Quan_Ly"]))
        self.table.setItem(row, 3, QTableWidgetItem(item["Ten_May"]))
        self.table.setItem(row, 4, QTableWidgetItem(item["He_Dieu_Hanh"]))
        self.table.setItem(row, 5, QTableWidgetItem(item["Dia_Chi_IP"]))
        self.table.setItem(row, 6, QTableWidgetItem(item["CPU"]))
        self.table.setItem(row, 7, QTableWidgetItem(item["RAM"]))
        self.table.setItem(row, 8, QTableWidgetItem(f"{item['Loai_OCung']} ({item['DungLuong_OCung']})"))
        self.table.setItem(row, 9, QTableWidgetItem(str(item["So_Luong_Lo_Hong"])))

        risk_item = QTableWidgetItem(item["Muc_Do_Rui_Ro"])
        if "Nguy cấp" in item["Muc_Do_Rui_Ro"]:
            risk_item.setForeground(QColor("#D32F2F"))
        elif "Cao" in item["Muc_Do_Rui_Ro"]:
            risk_item.setForeground(QColor("#F57C00"))
        elif "Trung bình" in item["Muc_Do_Rui_Ro"]:
            risk_item.setForeground(QColor("#B8860B"))
        else:
            risk_item.setForeground(QColor("#2E7D32"))
        risk_item.setFont(QFont("Arial", 9, QFont.Weight.Bold))
        # Rê chuột lên ô mức nguy cơ để xem lý do bị xếp mức đó
        risk_item.setToolTip("\n".join(item["Ly_Do_Rui_Ro"]) or "Không ghi nhận vấn đề")
        self.table.setItem(row, 10, risk_item)

        malware_item = QTableWidgetItem(item["Ma_Doc"])
        if item["Ho_Ma_Doc"]:
            malware_item.setForeground(QColor("#D32F2F"))
            malware_item.setFont(QFont("Arial", 9, QFont.Weight.Bold))
        self.table.setItem(row, 11, malware_item)
        self.table.setItem(row, 12, QTableWidgetItem("; ".join(item["Lich_Su_USB"]) if item["Lich_Su_USB"] else "Không có"))
        self.table.setItem(row, 13, QTableWidgetItem(item["Mat_Khau"]))
        self.table.setItem(row, 14, QTableWidgetItem(item["Phan_Loai_May"]))

        tv = item.get("Toan_Ven", "")
        tv_item = QTableWidgetItem(tv)
        tv_item.setToolTip("Nguồn dữ liệu: " + item.get("Nguon_Du_Lieu", "docx"))
        if tv == "Hợp lệ":
            tv_item.setForeground(QColor("#2E7D32"))
        elif tv.startswith("Chữ ký SAI"):
            tv_item.setForeground(QColor("#D32F2F"))
            tv_item.setFont(QFont("Arial", 9, QFont.Weight.Bold))
        elif tv.startswith("Docx"):
            tv_item.setForeground(QColor("#B8860B"))
        else:
            tv_item.setForeground(QColor("#666"))
        self.table.setItem(row, 15, tv_item)

    def process_finished(self, results):
        self.data_list = results
        self.btn_run.setEnabled(True)
        self.lbl_status.setText(f"Đã xử lý xong {len(results)} biên bản! Nhấp đúp vào dòng để mở file Word.")
        self.update_analysis_tab()
        thong_bao = f"Đã quét và trích xuất thành công {len(results)} biên bản!"
        if self._file_loi:
            thong_bao += (f"\n\n{len(self._file_loi)} file KHÔNG đọc được (cần kiểm tra lại):\n"
                          + "\n".join(self._file_loi[:15]))
            QMessageBox.warning(self, "Thông báo", thong_bao)
        else:
            QMessageBox.information(self, "Thông báo", thong_bao)

    def open_selected_docx(self, index):
        row = index.row()
        if row < len(self.data_list):
            file_path = self.data_list[row]["File_Path"]
            if os.path.exists(file_path):
                os.startfile(file_path)

    def apply_filter(self):
        search_kw = self.txt_search.text().lower()
        risk_filter = self.cbo_filter_risk.currentText()

        for row in range(self.table.rowCount()):
            match_search = False
            for col in range(self.table.columnCount()):
                item = self.table.item(row, col)
                if item and search_kw in item.text().lower():
                    match_search = True
                    break

            match_risk = True
            if risk_filter != "Tất cả":
                risk_item = self.table.item(row, 10)
                if risk_item and risk_filter not in risk_item.text():
                    match_risk = False

            self.table.setRowHidden(row, not (match_search and match_risk))

    def update_analysis_tab(self):
        total_machines = len(self.data_list)
        critical_count = sum(1 for d in self.data_list if "Nguy cấp" in d["Muc_Do_Rui_Ro"])
        total_vulns = sum(d["So_Luong_Lo_Hong"] for d in self.data_list)
        total_usb = sum(len(d["Lich_Su_USB"]) for d in self.data_list)

        self.lbl_kpi_total.setText(f"{total_machines}\nMáy Kiểm Tra")
        self.lbl_kpi_critical.setText(f"{critical_count}\nNguy Cấp")
        self.lbl_kpi_vuln.setText(f"{total_vulns}\nTổng Lỗ Hổng")
        self.lbl_kpi_usb.setText(f"{total_usb}\nLượt Cắm USB")

        all_vulns = []
        for d in self.data_list:
            all_vulns.extend(d["Danh_Sach_Lo_Hong"])
        vuln_counts = Counter(all_vulns).most_common(10)

        vuln_report = "=== CẢNH BÁO TRỌNG ĐIỂM (cần xử lý ngay) ===\n\n"
        nhom_canh_bao = [
            ("Phát hiện dấu hiệu mã độc", lambda d: d["Ho_Ma_Doc"],
             lambda d: ", ".join(d["Ho_Ma_Doc"])),
            ("Có lỗ hổng mức NGUY CẤP chưa vá", lambda d: d["Lo_Hong_Nguy_Hiem"],
             lambda d: ", ".join(d["Lo_Hong_Nguy_Hiem"])),
            ("Máy nội bộ đồng thời kết nối Internet",
             lambda d: "Máy nội bộ đồng thời kết nối Internet" in d["Ly_Do_Rui_Ro"], lambda d: ""),
            ("Không đặt mật khẩu đăng nhập", lambda d: d["Mat_Khau"] == "Không đặt MK", lambda d: ""),
            ("Hệ điều hành hết/sắp hết hỗ trợ",
             lambda d: any("hết hỗ trợ" in x for x in d["Ly_Do_Rui_Ro"]), lambda d: d["He_Dieu_Hanh"]),
            ("Chưa đối chiếu được mã độc/lỗ hổng (thiếu dữ liệu)",
             lambda d: any(x.startswith("Chưa đối chiếu") for x in d["Ly_Do_Rui_Ro"]), lambda d: ""),
            ("BIÊN BẢN NGHI BỊ GIẢ MẠO (chữ ký sai)",
             lambda d: d.get("Toan_Ven", "").startswith("Chữ ký SAI"), lambda d: ""),
            ("Biên bản bị chỉnh sửa sau khi tạo (docx đổi so với bản gốc đã ký)",
             lambda d: d.get("Toan_Ven", "").startswith("Docx"), lambda d: ""),
            ("Biên bản không có file kèm đã ký (không kiểm chứng được nguồn gốc)",
             lambda d: d.get("Toan_Ven", "").startswith("Không có file kèm"), lambda d: ""),
        ]
        for tieu_de, dieu_kien, chi_tiet in nhom_canh_bao:
            may = [d for d in self.data_list if dieu_kien(d)]
            vuln_report += f"■ {tieu_de}: {len(may)} máy\n"
            for d in may:
                ct = chi_tiet(d)
                vuln_report += f"    - {d['Ten_May']} ({d['Can_Bo_Quan_Ly']})" + (f": {ct}" if ct else "") + "\n"
        vuln_report += "\n=== TOP 10 LỖ HỔNG BẢO MẬT XUẤT HIỆN NHIỀU NHẤT ===\n\n"
        for rank, (vuln, cnt) in enumerate(vuln_counts, 1):
            desc = f" ({CRITICAL_VULNS_DICT[vuln]})" if vuln in CRITICAL_VULNS_DICT else ""
            vuln_report += f"{rank:02d}. {vuln:<18}{desc} -> {cnt}/{total_machines} máy vi tính\n"

        self.txt_top_vuln.setText(vuln_report)

        usb_list = []
        for d in self.data_list:
            for usb in d["Lich_Su_USB"]:
                usb_list.append(f"[{d['Ten_May']} - {d['Can_Bo_Quan_Ly']}] -> {usb}")

        # USB (cùng số seri) đã cắm vào nhiều máy - đặc biệt là "cầu nối" giữa máy có Internet và
        # máy nội bộ/độc lập: con đường lây nhiễm mã độc và làm lộ lọt dữ liệu điển hình.
        seri_may = {}
        for d in self.data_list:
            for seri in set(d["USB_Seri"]):
                if seri and seri.lower() not in ("không xác định", "không rõ", "n/a", "?"):
                    seri_may.setdefault(seri, []).append(d)
        usb_report = "=== THIẾT BỊ LƯU TRỮ DÙNG CHUNG NHIỀU MÁY ===\n\n"
        dung_chung = {s: ds for s, ds in seri_may.items() if len(ds) > 1}
        if not dung_chung:
            usb_report += "Không phát hiện thiết bị cùng số seri trên nhiều máy.\n"
        for seri, ds in sorted(dung_chung.items(), key=lambda x: -len(x[1])):
            co_net = [d for d in ds if "internet" in d["Phan_Loai_May"].lower()
                      or d["Ket_Noi_Mang"].lower().startswith("có")]
            khong_net = [d for d in ds if d not in co_net]
            canh_bao = " ⚠ CẦU NỐI máy Internet <-> máy không Internet" if co_net and khong_net else ""
            usb_report += f"Seri {seri}: {len(ds)} máy{canh_bao}\n"
            for d in ds:
                usb_report += f"    - {d['Ten_May']} ({d['Can_Bo_Quan_Ly']}) - {d['Phan_Loai_May']}\n"

        usb_report += f"\n=== TỔNG HỢP {len(usb_list)} THIẾT BỊ NGOẠI VI GHI NHẬN ===\n\n"
        usb_report += "\n".join(usb_list) if usb_list else "Không ghi nhận thiết bị lưu trữ ngoài cắm vào hệ thống."
        self.txt_usb_analysis.setText(usb_report)

    def export_excel(self):
        if not self.data_list:
            QMessageBox.warning(self, "Cảnh báo", "Chưa có dữ liệu để xuất Excel!")
            return

        save_path, _ = QFileDialog.getSaveFileName(self, "Lưu file Excel tổng hợp", "Tong_Hop_Kiem_Tra_ATTT.xlsx", "Excel Files (*.xlsx)")
        if not save_path:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Tong_Hop_ATTT"
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "STT", "Tên Tệp", "Thời Gian KT", "Địa Điểm", "Cán Bộ KT", "Cán Bộ QL",
            "Mật Khẩu", "Phân Loại Máy", "Tên Máy", "HĐH", "Ngày Cài", "IP", "MAC",
            "CPU", "RAM", "Loại Ổ Cứng", "Dung Lượng Ổ", "Phần Mềm Diệt Virus", "Kết Nối Mạng",
            "Số Lượng Lỗ Hổng", "Mức Rủi Ro", "Danh Sách Lỗ Hổng Bảo Mật", "Tình Trạng Mã Độc", "Lịch Sử Cắm USB",
            "Chức Vụ Cán Bộ KT", "Cung Cấp MK Cho Người Khác", "Phần Mềm Ứng Dụng", "Lịch Sử Internet",
            "Lý Do Xếp Mức Rủi Ro", "Toàn Vẹn Biên Bản", "Nguồn Dữ Liệu"
        ]
        ws.append(headers)

        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9')
        )

        for col_num in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col_num)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")

        for idx, d in enumerate(self.data_list, 1):
            row_data = [
                idx, d["File_Name"], d["Thoi_Gian_KT"], d["Dia_Diem"], d["Can_Bo_KT"], d["Can_Bo_Quan_Ly"],
                d["Mat_Khau"], d["Phan_Loai_May"], d["Ten_May"], d["He_Dieu_Hanh"], d["Ngay_Cai_Dat"],
                d["Dia_Chi_IP"], d["Dia_Chi_MAC"], d["CPU"], d["RAM"], d["Loai_OCung"], d["DungLuong_OCung"],
                d["Phan_Mem_Diet_Virus"], d["Ket_Noi_Mang"], d["So_Luong_Lo_Hong"], d["Muc_Do_Rui_Ro"],
                ", ".join(d["Danh_Sach_Lo_Hong"]), d["Ma_Doc"], "; ".join(d["Lich_Su_USB"]),
                d["Chuc_Vu_KT"], d["Cung_Cap_MK"], d["Phan_Mem_Ung_Dung"], d["Lich_Su_Internet"],
                "; ".join(d["Ly_Do_Rui_Ro"]), d.get("Toan_Ven", ""), d.get("Nguon_Du_Lieu", "")
            ]
            ws.append(row_data)
            for c_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=idx + 1, column=c_idx)
                cell.font = Font(name="Arial", size=10)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            mau_rui_ro = {"Nguy cấp": "F8CBAD", "Cao": "FCE4D6", "Trung bình": "FFF2CC"}
            for tu_khoa, mau in mau_rui_ro.items():
                if tu_khoa in d["Muc_Do_Rui_Ro"]:
                    ws.cell(row=idx + 1, column=headers.index("Mức Rủi Ro") + 1).fill = PatternFill(
                        start_color=mau, end_color=mau, fill_type="solid")
            if d["Ho_Ma_Doc"]:
                ws.cell(row=idx + 1, column=headers.index("Tình Trạng Mã Độc") + 1).font = Font(
                    name="Arial", size=10, bold=True, color="C00000")

        for col in ws.columns:
            max_len = max(len(str(c.value or '')) for c in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        ws.row_dimensions[1].height = 26
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        try:
            wb.save(save_path)
        except PermissionError:
            QMessageBox.critical(self, "Lỗi ghi file",
                                 "Không ghi được file Excel - có thể file đang được mở trong Excel. "
                                 "Hãy đóng file rồi xuất lại.")
            return
        QMessageBox.information(self, "Thành công", f"Đã xuất báo cáo Excel thành công tại:\n{save_path}")


APP_STYLESHEET = """
QMainWindow, QWidget { background-color: #f4f6fb; color: #1f2937; font-family: 'Segoe UI'; font-size: 10pt; }
QGroupBox { background-color: #ffffff; border: 1px solid #dbe2ef; border-radius: 10px;
            margin-top: 14px; padding: 12px 10px 10px 10px; font-weight: bold; }
QGroupBox::title { subcontrol-origin: margin; left: 14px; padding: 2px 8px;
                   color: #1e3a8a; background-color: #e8efff; border-radius: 6px; }
QTabWidget::pane { border: 1px solid #dbe2ef; border-radius: 8px; top: -1px; background: #ffffff; }
QTabBar::tab { background: #e2e8f0; color: #334155; padding: 9px 18px; margin-right: 3px;
               border-top-left-radius: 8px; border-top-right-radius: 8px; font-weight: bold; }
QTabBar::tab:selected { background: #1e3a8a; color: #ffffff; }
QTabBar::tab:hover:!selected { background: #cbd5e1; }
QLineEdit, QSpinBox, QComboBox { background: #ffffff; border: 1px solid #cbd5e1; border-radius: 6px; padding: 5px 8px; }
QLineEdit:focus, QSpinBox:focus, QComboBox:focus { border: 1px solid #2563eb; }
QLineEdit:read-only { background: #f1f5f9; color: #475569; }
QPushButton { background-color: #2563eb; color: #ffffff; border: none; border-radius: 6px;
              padding: 7px 14px; font-weight: bold; }
QPushButton:hover { background-color: #1d4ed8; }
QPushButton:pressed { background-color: #1e40af; }
QPushButton:disabled { background-color: #94a3b8; }
QTableWidget { background: #ffffff; alternate-background-color: #f8fafc; gridline-color: #e5e7eb;
               border: 1px solid #dbe2ef; border-radius: 8px; selection-background-color: #dbeafe;
               selection-color: #0f172a; }
QHeaderView::section { background-color: #1e3a8a; color: #ffffff; padding: 6px; border: none; font-weight: bold; }
QProgressBar { border: 1px solid #cbd5e1; border-radius: 8px; text-align: center; background: #e2e8f0; height: 20px; }
QProgressBar::chunk { background-color: #16a34a; border-radius: 7px; }
QTextEdit { background: #ffffff; border: 1px solid #dbe2ef; border-radius: 8px; }
QCheckBox { spacing: 6px; }
QLabel { background: transparent; }
"""

if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    app.setStyleSheet(APP_STYLESHEET)
    win = ATTTAnalysisTool()
    win.show()
    sys.exit(app.exec())