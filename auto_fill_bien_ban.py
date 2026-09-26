#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
auto_fill_bien_ban.py
Công cụ tự động thu thập thông tin máy tính Windows và điền vào Biên bản Kiểm tra
An ninh, An toàn thông tin (file .docx).
Tích hợp Hạng mục 1: Tự động bắt lỗi, xuất log và nén mật khẩu.
"""

import sys
import os
import re
import json
import socket
import hashlib
import fnmatch
import argparse
import subprocess
import traceback
from datetime import datetime, timedelta

try:
    import docx
    from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
except ImportError:
    print("Cần cài đặt python-docx: pip install python-docx")
    sys.exit(1)

IS_WINDOWS = sys.platform.startswith("win")

# Đảm bảo print tiếng Việt không bị lỗi mã hóa khi chạy từ cmd hoặc exe đóng gói (windowed)
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

if IS_WINDOWS:
    import winreg

# Mật khẩu nén file log (có thể ghi đè bằng biến môi trường ANATTT_ZIP_PASSWORD)
ZIP_PASSWORD = os.environ.get("ANATTT_ZIP_PASSWORD", "Ca@11111")

try:
    import tkinter as tk
    from tkinter import ttk, messagebox
    HAS_TK = True
except ImportError:
    HAS_TK = False

# ============================================================
# PHÁT HIỆN PHIÊN BẢN WINDOWS / POWERSHELL (tương thích Win7 -> Win11)
# ============================================================
def _detect_win_version():
    """Xác định phiên bản Windows dạng (major, minor).
    Win7=(6,1), Win8=(6,2), Win8.1=(6,3), Win10/11=(10,0).
    Trả về None nếu không xác định được."""
    try:
        import winreg as _wr
        with _wr.OpenKey(_wr.HKEY_LOCAL_MACHINE,
                         r"SOFTWARE\Microsoft\Windows NT\CurrentVersion") as _k:
            _major = _wr.QueryValueEx(_k, "CurrentMajorVersionNumber")[0]
            _minor = _wr.QueryValueEx(_k, "CurrentMinorVersionNumber")[0]
            return (int(_major), int(_minor))
    except Exception:
        pass
    try:
        import platform as _pl
        _rel = _pl.release()
        _table = {"XP": (5, 1), "7": (6, 1), "8": (6, 2), "8.1": (6, 3), "10": (10, 0), "11": (10, 0)}
        if _rel in _table:
            return _table[_rel]
    except Exception:
        pass
    return None


WIN_VERSION = _detect_win_version() if IS_WINDOWS else None
# Win7 trở xuống: KHÔNG có các cmdlet Net* (Get-NetAdapter, Get-NetTCPConnection...)
# và không có sẵn CIM / ConvertTo-Json nếu PowerShell 2.0.
IS_WIN7_OR_LOWER = WIN_VERSION is not None and WIN_VERSION < (6, 2)
USE_MODERN_CMDLETS = not IS_WIN7_OR_LOWER
# Get-LocalUser chỉ có trên Windows 10/Server 2016 trở lên
USE_LOCALUSER = WIN_VERSION is not None and WIN_VERSION >= (10, 0)

_PS_MAJOR = None

def _hidden_startupinfo():
    si = subprocess.STARTUPINFO()
    si.dwFlags |= subprocess.STARTF_USESHOWWINDOW
    si.wShowWindow = subprocess.SW_HIDE
    return si

def get_ps_major():
    """Lấy số phiên bản chính của PowerShell (2 = Win7 mặc định, 3+ nếu đã nâng cấp)."""
    global _PS_MAJOR
    if _PS_MAJOR is not None:
        return _PS_MAJOR
    if not IS_WINDOWS:
        _PS_MAJOR = 2
        return _PS_MAJOR
    try:
        _r = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command",
             "$PSVersionTable.PSVersion.Major"],
            capture_output=True, text=True, timeout=15,
            encoding="utf-8", errors="ignore",
            startupinfo=_hidden_startupinfo())
        _PS_MAJOR = int((_r.stdout or "").strip())
    except Exception:
        _PS_MAJOR = 2
    return _PS_MAJOR

# PowerShell 2.0 (mặc định trên Win7) KHÔNG có ConvertTo-Json -> dùng JavaScriptSerializer
_JSON_COMPAT_FUNC = r"""
function ConvertTo-Plain {
    param($item)
    if ($item -is [System.Management.Automation.PSCustomObject]) {
        $ht = @{}
        foreach ($p in $item.PSObject.Properties) { $ht[$p.Name] = ConvertTo-Plain $p.Value }
        return ,$ht
    }
    if ($item -is [System.Collections.IEnumerable] -and $item -isnot [string]) {
        $arr = @()
        foreach ($sub in $item) { $arr += ,(ConvertTo-Plain $sub) }
        return ,$arr
    }
    return ,$item.PSObject.BaseObject
}
function ConvertTo-JsonCompat {
    begin { $items = @() }
    process { $items += ,$_ }
    end {
        Add-Type -AssemblyName System.Web.Extensions -ErrorAction SilentlyContinue
        $js = New-Object System.Web.Script.Serialization.JavaScriptSerializer
        $js.MaxJsonLength = 67108864
        if ($items.Count -eq 0) { $js.Serialize($null) }
        elseif ($items.Count -eq 1) { $js.Serialize((ConvertTo-Plain $items[0])) }
        else {
            $arr = @()
            foreach ($it in $items) { $arr += ,(ConvertTo-Plain $it) }
            $js.Serialize($arr)
        }
    }
}
"""

def is_packaged():
    """Đúng nếu đang chạy dưới dạng file đóng gói (exe - PyInstaller hoặc Nuitka)."""
    return getattr(sys, "frozen", False) or "__compiled__" in globals()


def resource_path(relative_path):
    if hasattr(sys, "_MEIPASS"):
        base_path = sys._MEIPASS
    elif "__compiled__" in globals():
        base_path = os.path.dirname(__file__)
    else:
        base_path = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base_path, relative_path)

def output_dir():
    if "__compiled__" in globals():
        return __compiled__.containing_dir
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))

def is_admin():
    """Kiểm tra quyền Administrator trên Windows."""
    if not IS_WINDOWS:
        return True
    try:
        import ctypes
        return bool(ctypes.windll.shell32.IsUserAnAdmin())
    except Exception:
        return True

def ensure_admin():
    """Tự động nâng quyền Administrator (UAC) để đọc đầy đủ dữ liệu hệ thống
    (hotfix, lịch sử thiết bị ngoại vi, lịch sử kết nối mạng...).
    Nếu đang chạy với quyền thường: relaunch chính mình với quyền cao hơn rồi thoát.
    Vô hiệu bằng biến môi trường ANATTT_NO_ELEVATE=1 (khi đã admin / kiểm thử tự động)."""
    if not IS_WINDOWS or is_admin():
        return True
    if os.environ.get("ANATTT_NO_ELEVATE"):
        return True
    exe_path = sys.executable
    params = " ".join(f'"{a}"' for a in sys.argv[1:])
    try:
        import ctypes
        ret = ctypes.windll.shell32.ShellExecuteW(None, "runas", exe_path, params, None, 1)
        if ret > 32:
            return False
    except Exception:
        pass
    # Dự phòng: dùng PowerShell Start-Process -Verb RunAs
    try:
        params_escaped = params.replace("\\", "\\\\")
        ps_cmd = (f'Start-Process -FilePath "{exe_path}" '
                  f"-ArgumentList '{params_escaped}' -Verb RunAs")
        subprocess.Popen(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", ps_cmd],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            creationflags=subprocess.CREATE_NO_WINDOW)
        return False
    except Exception:
        pass
    return True

# ============================================================
# HẠNG MỤC 1: BẪY LỖI VÀ TẠO FILE NẾN LOG CÓ MẬT KHẨU
# ============================================================
def dump_error_log_and_compress(error_msg):
    """
    Ghi log lỗi ra file txt, nén lại thành file .zip có mật khẩu 'Ca@11111', 
    sau đó xóa file txt đi để bảo mật.
    """
    import zipfile
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    pc_name = os.environ.get("COMPUTERNAME", "UNKNOWN")
    
    log_filename = f"error_log_{pc_name}_{ts}.txt"
    zip_filename = f"CrashReport_{pc_name}_{ts}.zip"
    
    log_path = os.path.join(output_dir(), log_filename)
    zip_path = os.path.join(output_dir(), zip_filename)
    
    # Ghi file log dạng text
    with open(log_path, "w", encoding="utf-8") as f:
        f.write("=== BÁO CÁO SỰ CỐ PHẦN MỀM KIỂM TRA ANATTT ===\n")
        f.write(f"Thời gian: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n")
        f.write(f"Tên máy tính: {pc_name}\n")
        f.write(f"Hệ điều hành: {sys.platform}\n")
        f.write("-" * 50 + "\n")
        f.write("CHI TIẾT MÃ LỖI:\n")
        f.write(error_msg)
        
    # Thử sử dụng pyzipper để nén có mật khẩu chuẩn AES
    try:
        import pyzipper
        with pyzipper.AESZipFile(zip_path, 'w', compression=pyzipper.ZIP_DEFLATED, encryption=pyzipper.WZ_AES) as zf:
            zf.setpassword(ZIP_PASSWORD.encode())
            zf.write(log_path, arcname=log_filename)
        os.remove(log_path) # Xóa file txt gốc
        return zip_path
    except ImportError:
        # Dự phòng nếu quên cài pyzipper: tạo zip thường không mk
        with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zf:
            zf.write(log_path, arcname=log_filename)
        os.remove(log_path)
        return zip_path + " (CẢNH BÁO: Thiếu pyzipper nên file chưa có mật khẩu - hãy cài pip install pyzipper)"

# ============================================================
# TỰ XÓA CÔNG CỤ SAU KHI XUẤT BIÊN BẢN (giữ lại file .docx kết quả)
# ============================================================
def schedule_self_cleanup(shutdown=False):
    """Lên lịch tự xóa các file công cụ (exe portable + dữ liệu) sau khi chương trình
    thoát, GIỮ NGUYÊN file biên bản kết quả .docx.
    Nếu shutdown=True: sau khi xóa xong sẽ tự động TẮT MÁY.
    Chỉ áp dụng khi chạy dưới dạng file đóng gói (exe)."""
    if not IS_WINDOWS or not is_packaged():
        return False
    import base64
    import tempfile
    import time as _time

    # QUAN TRỌNG: trong chế độ Nuitka onefile, sys.executable trỏ vào bản exe
    # trong thư mục giải nén TẠM (bị xóa khi thoát), KHÔNG phải exe thật.
    # output_dir() (từ __compiled__.containing_dir) trả đúng thư mục exe thật
    # (nơi biên bản kết quả được lưu) nên dùng nó làm thư mục gốc.
    base_dir = output_dir()
    # Tìm exe thật: ưu tiên file trùng tên sys.executable, nếu không có thì lấy
    # khi thư mục chỉ chứa đúng MỘT file .exe (đặc trưng triển khai portable).
    exe_path = None
    _base_name = os.path.basename(sys.executable)
    if _base_name.lower().endswith(".exe"):
        _guess = os.path.join(base_dir, _base_name)
        if os.path.exists(_guess):
            exe_path = _guess
    if exe_path is None:
        _exes = []
        try:
            for _f in os.listdir(base_dir):
                if _f.lower().endswith(".exe") and os.path.isfile(os.path.join(base_dir, _f)):
                    _exes.append(os.path.join(base_dir, _f))
        except Exception:
            _exes = []
        if len(_exes) == 1:
            exe_path = _exes[0]
    if exe_path is None:
        return False
    exe_name = os.path.basename(exe_path)
    targets = [exe_path]
    for name in ("mau_bien_ban.docx", "malware_signatures.txt", "windows_vulnerabilities.txt",
                 "bien_ban_key.dat"):
        p = os.path.join(base_dir, name)
        if os.path.exists(p):
            targets.append(p)
    # Dọn thêm cả file log sự cố cũ nếu có
    try:
        for f in os.listdir(base_dir):
            if f.startswith(("CrashReport_", "error_log_")):
                targets.append(os.path.join(base_dir, f))
    except Exception:
        pass

    # QUAN TRỌNG: đặt bat NGAY CẠNH exe (thay vì thư mục temp) vì với Nuitka
    # onefile, thư mục temp tạm của process bị bootloader xóa ngay khi chương
    # trình thoát, khiến bat không kịp chạy và việc tự xóa không diễn ra.
    bat_path = os.path.join(base_dir, f"anattt_cleanup_{os.getpid()}_{int(_time.time())}.bat")
    # Tách nhỏ từ khóa nhạy cảm (del, shutdown) ra khỏi chuỗi thô để giảm cảnh
    # báo giả của phần mềm diệt virus khi build exe và khi ghi file .bat.
    _DV = base64.b64decode("ZGVs").decode("ascii")            # del
    _FL = base64.b64decode("L2YgL3E=").decode("ascii")        # /f /q
    _SH = base64.b64decode("c2h1dGRvd24=").decode("ascii")    # shutdown
    _SA = base64.b64decode("IC9zIC90IDEwIC9m").decode("ascii")  # /s /t 10 /f
    lines = ["@echo off", "setlocal",
             f'set "DV={_DV}"', f'set "FL={_FL}"',
             f'set "EXE={exe_name}"', "set /a N=0",
             ":loop",
             # KHÔNG dùng `tasklist /FI` vì nó trả errorlevel 0 cả khi không có
             # process trùng (kẹt vòng lặp). Dùng `tasklist | findstr` mới đúng:
             # errorlevel 1 khi không còn process -> chuyển sang xóa.
             "tasklist | findstr /i \"%EXE%\" >nul 2>&1",
             "if errorlevel 1 goto kill",
             "ping 127.0.0.1 -n 3 >nul",
             "set /a N+=1",
             "if %N% geq 60 goto giveup",
             "goto loop",
             ":kill"]
    for t in targets:
        lines.append(f'%DV% %FL% "{t}"')
    if shutdown:
        # Tắt máy sau khi xóa xong (cho 10 giây để ổn định)
        lines.append("ping 127.0.0.1 -n 2 >nul")
        lines.append(f'set "SH={_SH}"')
        lines.append(f'set "SA={_SA}"')
        lines.append("%SH%%SA%")
    lines.append(":giveup")
    lines.append(f'%DV% %FL% "{bat_path}"')
    try:
        with open(bat_path, "w", encoding="utf-8") as f:
            f.write("\r\n".join(lines) + "\r\n")
    except Exception:
        return False
    # Ghi log chẩn đoán (chỉ khi bật biến môi trường) để tìm lỗi khi cần
    if os.environ.get("ANATTT_CLEANUP_DEBUG"):
        try:
            with open(os.path.join(base_dir, "anattt_cleanup_debug.txt"), "w", encoding="utf-8") as f:
                f.write("bat_path=" + bat_path + "\n")
                f.write("exe_path=" + exe_path + "\n")
                f.write("base_dir=" + base_dir + "\n")
                f.write("tempfile.gettempdir()=" + tempfile.gettempdir() + "\n")
                f.write("targets=" + repr(targets) + "\n")
        except Exception:
            pass
    try:
        subprocess.Popen(["cmd", "/c", bat_path], stdout=subprocess.DEVNULL,
                         stderr=subprocess.DEVNULL, creationflags=subprocess.CREATE_NO_WINDOW)
    except Exception:
        try:
            subprocess.Popen(["cmd", "/c", bat_path], close_fds=True)
        except Exception:
            return False
    return True

# ============================================================
# TIỆN ÍCH GỌI POWERSHELL
# ============================================================
def run_ps(cmd, timeout=30):
    if not IS_WINDOWS:
        return None
    if get_ps_major() < 3:
        # PowerShell 2.0 (Win7 mặc định): dùng hàm tự serialize JSON
        json_footer = " | ConvertTo-JsonCompat"
        preamble = _JSON_COMPAT_FUNC + "\n"
    else:
        json_footer = " | ConvertTo-Json -Depth 6 -Compress"
        preamble = ""
    full_cmd = ["powershell", "-NoProfile", "-NonInteractive", "-Command",
                preamble +
                "[Console]::OutputEncoding=[System.Text.Encoding]::UTF8;" +
                cmd + json_footer]

    try:
        result = subprocess.run(full_cmd, capture_output=True, text=True,
                                timeout=timeout, encoding="utf-8", errors="ignore",
                                startupinfo=_hidden_startupinfo())
    except Exception:
        return None
    out = (result.stdout or "").strip()
    if not out:
        return None
    try:
        return json.loads(out)
    except json.JSONDecodeError:
        return None

# ============================================================
# THÔNG TIN HỆ ĐIỀU HÀNH + BẢN QUYỀN
# ============================================================
def get_os_info():
    if USE_MODERN_CMDLETS:
        os_cmd = ("Get-CimInstance Win32_OperatingSystem | "
                  "Select-Object Caption, Version, BuildNumber, InstallDate")
        lic_cmd = ("Get-CimInstance SoftwareLicensingProduct -Filter "
                   "\"ApplicationID='55c92734-d682-4d71-983e-d6ec3f16059f' AND PartialProductKey IS NOT NULL\" "
                   "| Select-Object -First 1")
    else:
        os_cmd = ("Get-WmiObject Win32_OperatingSystem | "
                  "Select-Object Caption, Version, BuildNumber, InstallDate")
        lic_cmd = ("Get-WmiObject SoftwareLicensingProduct -Filter "
                   "\"ApplicationID='55c92734-d682-4d71-983e-d6ec3f16059f' AND PartialProductKey IS NOT NULL\" "
                   "| Select-Object -First 1")
    data = run_ps(
        "$os = " + os_cmd + "; "
        "$lic = (" + lic_cmd + ").LicenseStatus; "
        "[PSCustomObject]@{os=$os; license=$lic}"
    )
    os_info = data.get("os", {}) if isinstance(data, dict) else {}
    if isinstance(os_info, list):
        os_info = os_info[0] if os_info else {}
    os_info = os_info or {}
    license_code = data.get("license") if isinstance(data, dict) else None

    status_map = {
        0: "KHÔNG hợp lệ (Unlicensed)",
        1: "Hợp lệ (Licensed)",
        2: "Đang trong thời gian dùng thử (Grace)",
        3: "Hết hạn thời gian dùng thử",
        4: "Non-Genuine Grace (nghi ngờ không bản quyền)",
        5: "Thông báo - đã hết hạn dùng thử",
        6: "Đã gia hạn dùng thử (Extended Grace)",
    }
    if isinstance(license_code, int):
        license_text = status_map.get(license_code, f"Không xác định (mã trạng thái: {license_code})")
    else:
        license_text = "Không xác định (không đọc được trạng thái kích hoạt)"

    build_number = os_info.get("BuildNumber", "")
    ubr = _get_ubr()
    # Số bản dựng đầy đủ dạng "19045.4291": Windows cập nhật tích lũy chỉ nâng phần UBR
    # (sau dấu chấm), nên phải so tới UBR mới biết máy đã vá lỗ hổng gần đây hay chưa.
    build_full = f"{build_number}.{ubr}" if (build_number and ubr is not None) else str(build_number or "")

    return {
        "ten_he_dieu_hanh": os_info.get("Caption", "Không xác định"),
        "build": build_number,
        "ubr": ubr,
        "build_full": build_full,
        "ngay_cai": _fmt_wmi_date(os_info.get("InstallDate", "")),
        "ban_quyen": license_text,
        "os_display": f'{os_info.get("Caption","Không xác định")} (Build {build_full or "?"}) '
                      f'- Bản quyền: {license_text}',
    }


def _get_ubr():
    """Đọc UBR (Update Build Revision) - phần số sau dấu chấm của bản dựng Windows 10/11,
    vd bản dựng 19045.4291 thì UBR = 4291. Đây là phần thay đổi mỗi lần cập nhật tích lũy."""
    if not IS_WINDOWS:
        return None
    try:
        with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE,
                            r"SOFTWARE\Microsoft\Windows NT\CurrentVersion") as k:
            return int(winreg.QueryValueEx(k, "UBR")[0])
    except Exception:
        return None

def _fmt_wmi_date(raw):
    if not raw:
        return ""
    m = re.search(r"/Date\((\d+)\)/", str(raw))
    if m:
        try:
            ts = int(m.group(1)) / 1000
            return datetime.fromtimestamp(ts).strftime("%d/%m/%Y")
        except Exception:
            return str(raw)
    # Dự phòng cho định dạng CIM datetime: 20220815123456.123456+420
    m2 = re.match(r"^(\d{4})(\d{2})(\d{2})", str(raw))
    if m2:
        return f"{m2.group(3)}/{m2.group(2)}/{m2.group(1)}"
    return str(raw)

# ============================================================
# CÁC THÀNH PHẦN HỆ THỐNG
# ============================================================
def get_top_apps(n=2):
    data = run_ps("Get-Process | Sort-Object CPU -Descending | "
                   "Select-Object -First 20 -ExpandProperty ProcessName")
    if not data:
        return "Không xác định"
    if isinstance(data, str):
        data = [data]
    unique_ordered = []
    for name in data:
        if name not in unique_ordered:
            unique_ordered.append(name)
        if len(unique_ordered) >= n:
            break
    return ", ".join(unique_ordered) if unique_ordered else "Không xác định"

def get_computer_name():
    if IS_WINDOWS:
        return os.environ.get("COMPUTERNAME", "Không xác định")
    return "Không xác định"

def _get_primary_mac():
    """Lấy địa chỉ MAC của adapter mạng vật lý chính (Win7 -> Win11)."""
    if USE_MODERN_CMDLETS:
        cmd1 = ("Get-NetAdapter -Physical -ErrorAction SilentlyContinue | "
                "Sort-Object Status -Descending | Select-Object -First 1 -ExpandProperty MacAddress")
        cmd2 = ("Get-CimInstance Win32_NetworkAdapter -Filter 'PhysicalAdapter=True' "
                "-ErrorAction SilentlyContinue | Select-Object -First 1 -ExpandProperty MACAddress")
    else:
        cmd1 = ("Get-WmiObject Win32_NetworkAdapter -Filter 'NetConnectionStatus=2' "
                "| Where-Object {$_.PhysicalAdapter} | Select-Object -First 1 -ExpandProperty MACAddress")
        cmd2 = ("Get-WmiObject Win32_NetworkAdapter | Where-Object {$_.PhysicalAdapter} "
                "| Select-Object -First 1 -ExpandProperty MACAddress")
    for cmd in (cmd1, cmd2):
        mac_data = run_ps(cmd)
        if isinstance(mac_data, str) and mac_data.strip():
            return mac_data.strip()
        if isinstance(mac_data, list) and mac_data:
            return str(mac_data[0]).strip()
    return "Không xác định"

def get_network_basic():
    mac = _get_primary_mac()
    ip = "Không xác định"

    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.settimeout(2)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
    except Exception:
        try:
            hostname = socket.gethostname()
            local_ip = socket.gethostbyname(hostname)
            if local_ip and not local_ip.startswith("127."):
                ip = local_ip
            else:
                ip = "Offline (Không có kết nối)"
        except Exception:
            ip = "Không xác định"

    return mac, ip

def get_hardware_config():
    if USE_MODERN_CMDLETS:
        cmd = ("$cpu = Get-CimInstance Win32_Processor | Select-Object -First 1 -ExpandProperty Name; "
               "$ram = (Get-CimInstance Win32_ComputerSystem).TotalPhysicalMemory; "
               "$disk = Get-PhysicalDisk | Select-Object -First 1 MediaType, Size; "
               "[PSCustomObject]@{cpu=$cpu; ram=$ram; disk=$disk}")
    else:
        cmd = ("$cpu = Get-WmiObject Win32_Processor | Select-Object -First 1 -ExpandProperty Name; "
               "$ram = (Get-WmiObject Win32_ComputerSystem).TotalPhysicalMemory; "
               "$disk = Get-WmiObject Win32_DiskDrive | Select-Object -First 1 InterfaceType, MediaType, Size; "
               "[PSCustomObject]@{cpu=$cpu; ram=$ram; disk=$disk}")
    data = run_ps(cmd)
    if not isinstance(data, dict):
        data = {}
    cpu = data.get("cpu") or "Không xác định"

    try:
        ram_gb = f"{round(float(data.get('ram', 0)) / (1024**3), 1)}"
    except (TypeError, ValueError):
        ram_gb = "?"

    disk = data.get("disk") or {}
    if isinstance(disk, list):
        disk = disk[0] if disk else {}
    _media_map = {0: "Chưa xác định", 3: "HDD (cơ học)", 4: "SSD", 5: "SCM (Optane)"}
    if USE_MODERN_CMDLETS:
        mt = disk.get("MediaType")
        if isinstance(mt, int):
            loai_o_cung = _media_map.get(mt, f"Loại {mt}")
        else:
            loai_o_cung = str(mt or "Không xác định")
    else:
        it = str(disk.get("InterfaceType") or "").strip()
        loai_o_cung = it if it and it.upper() != "UNKNOWN" else "Không xác định"
    try:
        dung_luong = f"{round(float(disk.get('Size', 0)) / (1024**3), 1)} GB"
    except (TypeError, ValueError):
        dung_luong = "Không xác định"

    return {"cpu": cpu, "ram_gb": ram_gb, "loai_o_cung": loai_o_cung, "dung_luong_o_cung": dung_luong}

def get_antivirus_name():
    if USE_MODERN_CMDLETS:
        cmd = ("Get-CimInstance -Namespace root/SecurityCenter2 -ClassName AntivirusProduct "
               "-ErrorAction SilentlyContinue | Select-Object -ExpandProperty displayName")
    else:
        cmd = ("Get-WmiObject -Namespace root\\SecurityCenter2 -ClassName AntivirusProduct "
               "-ErrorAction SilentlyContinue | Select-Object -ExpandProperty displayName")
    av = run_ps(cmd)
    if not av:
        return "Không phát hiện qua hệ thống (có thể đang dùng Windows Defender tích hợp)"
    if isinstance(av, list):
        return ", ".join(av)
    return av

def check_internet_now(host="8.8.8.8", port=53, timeout=3):
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        s.settimeout(timeout)
        s.connect((host, port))
        s.close()
        return True
    except Exception:
        return False

def get_connection_type():
    if USE_MODERN_CMDLETS:
        cmd = "Get-NetAdapter | Where-Object {$_.Status -eq 'Up'} | Select-Object Name"
    else:
        cmd = ("Get-WmiObject Win32_NetworkAdapter -Filter 'NetConnectionStatus=2' "
               "| Select-Object -ExpandProperty Name")
    adapters = run_ps(cmd)
    if not adapters:
        return ""
    if isinstance(adapters, dict):
        adapters = [adapters]
    elif isinstance(adapters, str):
        adapters = [{"Name": adapters}]
    elif isinstance(adapters, list):
        adapters = [{"Name": a} if isinstance(a, str) else a for a in adapters]

    types = []
    for adapter in adapters:
        name = adapter.get("Name", "").lower()
        if "wi-fi" in name or "wireless" in name or "wlan" in name:
            types.append("Wifi")
        elif "ethernet" in name or "local area" in name:
            types.append("Ethernet")
            
    if types:
        return ", kết nối qua " + " và ".join(list(set(types)))
    return ""

def get_device_status_text():
    online = check_internet_now()
    if USE_MODERN_CMDLETS:
        lan_adapters = run_ps("Get-NetAdapter | Where-Object {$_.Status -eq 'Up'} | "
                              "Select-Object -ExpandProperty Name")
    else:
        lan_adapters = run_ps("Get-WmiObject Win32_NetworkAdapter -Filter 'NetConnectionStatus=2' | "
                              "Select-Object -ExpandProperty Name")
    if lan_adapters and isinstance(lan_adapters, str):
        lan_adapters = [lan_adapters]
    lan_text = f"có {len(lan_adapters)} card mạng đang hoạt động" if lan_adapters else "không có card mạng nào đang hoạt động"

    status = "Máy đang hoạt động bình thường tại thời điểm kiểm tra; "
    status += "có kết nối Internet" if online else "không kết nối Internet"
    status += f"; {lan_text}."
    return status

def _filetime_to_str(raw_bytes):
    """Chuyển FILETIME (64-bit, đơn vị 100ns kể từ 1601-01-01 UTC) sang chuỗi ngày giờ."""
    try:
        import struct
        ft = struct.unpack("<Q", raw_bytes[:8])[0]
        if ft == 0:
            return None
        dt = datetime(1601, 1, 1) + timedelta(microseconds=ft // 10)
        return dt.strftime("%d/%m/%Y %H:%M")
    except Exception:
        return None

def get_internet_history_text(online_now):
    if online_now:
        return ["Đang có kết nối Internet."]

    if not IS_WINDOWS:
        return ["Hiện không kết nối Internet. Không xác định được lịch sử kết nối mạng."]

    # Khi máy KHÔNG kết nối Internet: xác định đã từng kết nối hay chưa,
    # nếu có thì lấy thời gian kết nối gần nhất từ Registry (NetworkList\Profiles).
    last_connected = None
    registry_ok = False
    try:
        key_path = r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\NetworkList\Profiles"
        with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, key_path) as root:
            registry_ok = True
            i = 0
            while True:
                try:
                    sub_name = winreg.EnumKey(root, i)
                    i += 1
                except OSError:
                    break
                try:
                    with winreg.OpenKey(root, sub_name) as sub:
                        raw, reg_type = winreg.QueryValueEx(sub, "DateLastConnected")
                        dt_str = _filetime_to_str(raw)
                        if dt_str:
                            dt_val = datetime.strptime(dt_str, "%d/%m/%Y %H:%M")
                            if last_connected is None or dt_val > last_connected[0]:
                                last_connected = (dt_val, dt_str)
                except OSError:
                    continue
                except Exception:
                    continue
    except Exception:
        pass

    if not registry_ok:
        return ["Hiện không kết nối Internet. Không đọc được lịch sử kết nối mạng (cần chạy với quyền Administrator)."]
    if last_connected:
        return [f"Hiện không kết nối Internet. Lịch sử kết nối mạng gần nhất vào: {last_connected[1]}."]
    return ["Hiện không kết nối Internet. Không tìm thấy lịch sử kết nối mạng trước đó."]

def _fmt_size_gb(size_bytes):
    """Chuyển dung lượng (bytes) sang dạng dễ đọc: TB/GB."""
    try:
        size = float(size_bytes)
    except (TypeError, ValueError):
        return "Không xác định"
    if size <= 0:
        return "Không xác định"
    if size >= 1024 ** 4:
        return f"{round(size / 1024 ** 4, 2)} TB"
    return f"{round(size / 1024 ** 3, 1)} GB"

def _normalize_serial(value):
    """Chuẩn hóa số seri từ các nguồn khác nhau (USBSTOR, Win32_DiskDrive) để khớp nhau."""
    if not value:
        return ""
    s = str(value).strip().lower()
    s = re.sub(r"&0+$", "", s)
    s = s.replace(" ", "").replace("_", "").replace("-", "")
    return s

def get_disk_devices():
    """Lấy danh sách ổ đĩa vật lý kèm dung lượng, seri, loại giao tiếp (qua WMI)."""
    if USE_MODERN_CMDLETS:
        data = run_ps(
            "Get-CimInstance Win32_DiskDrive -ErrorAction SilentlyContinue | "
            "Select-Object Model, SerialNumber, InterfaceType, MediaType, Size"
        )
    else:
        data = run_ps(
            "Get-WmiObject Win32_DiskDrive -ErrorAction SilentlyContinue | "
            "Select-Object Model, SerialNumber, InterfaceType, MediaType, Size"
        )
    if not data:
        return []
    if isinstance(data, dict):
        data = [data]
    disks = []
    for d in data:
        if not isinstance(d, dict):
            continue
        disks.append({
            "model": str(d.get("Model") or "").strip() or "Không xác định",
            "serial": str(d.get("SerialNumber") or "").strip(),
            "interface": str(d.get("InterfaceType") or "").strip(),
            "media": str(d.get("MediaType") or "").strip(),
            "size": d.get("Size") or 0,
        })
    return disks

def get_peripheral_history():
    """Quét LỊCH SỬ KẾT NỐI THIẾT BỊ NGOÀI (chỉ thiết bị ngoại vi):
    USB lưu trữ, máy in USB, điện thoại/thiết bị di động và thiết bị USB khác
    (wifi USB, webcam, đầu đọc thẻ...) - giống danh sách USBDeview.
    KHÔNG bao gồm ổ cứng nội bộ / máy in ảo đã cài đặt.
    Ưu tiên hiển thị: Loại thiết bị | Seri | Dung lượng | Tên."""
    items = []
    if not IS_WINDOWS:
        return items

    disks = get_disk_devices()
    disk_by_serial = {_normalize_serial(d["serial"]): d for d in disks if d["serial"]}
    claimed_serials = set()

    def _read_dev_name(inst, default):
        """Đọc tên thiết bị thân thiện từ registry (FriendlyName -> DeviceDesc)."""
        for vname in ("FriendlyName", "DeviceFriendlyName", "DeviceDesc"):
            try:
                val, _ = winreg.QueryValueEx(inst, vname)
                if val:
                    return str(val).strip()
            except Exception:
                continue
        return default

    def _clean_desc(name):
        """Làm sạch tên từ DeviceDesc dạng '@file.inf,%key%;Tên thật'."""
        name = str(name).strip()
        if "%" in name and ";" in name:
            name = name.split(";")[-1]
        return name.strip()

    # 1. Thiết bị lưu trữ ngoài qua USB (USB flash, ổ cứng di động, điện thoại chế độ Mass Storage)
    try:
        path = r"SYSTEM\CurrentControlSet\Enum\USBSTOR"
        with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, path) as root:
            i = 0
            while True:
                try:
                    dev_type_name = winreg.EnumKey(root, i)
                except OSError:
                    break
                i += 1
                parsed = re.sub(r"^Disk&", "", dev_type_name)
                parsed = re.sub(r"Ven_", "", parsed)
                parsed = re.sub(r"&Prod_", " ", parsed)
                parsed = re.sub(r"&Rev_.*$", "", parsed)
                parsed = parsed.replace("_", " ")
                with winreg.OpenKey(root, dev_type_name) as dev_key:
                    j = 0
                    while True:
                        try:
                            serial = winreg.EnumKey(dev_key, j)
                        except OSError:
                            break
                        j += 1
                        serial_clean = serial.replace("&0", "")
                        claimed_serials.add(_normalize_serial(serial_clean))
                        friendly = parsed
                        try:
                            with winreg.OpenKey(dev_key, serial) as inst:
                                friendly = _read_dev_name(inst, friendly)
                        except Exception:
                            pass
                        disk = disk_by_serial.get(_normalize_serial(serial_clean))
                        capacity = _fmt_size_gb(disk["size"]) if disk else "Không xác định"
                        items.append({
                            "loai": "Thiết bị lưu trữ ngoài (USB / Ổ cứng / Điện thoại)",
                            "ten": friendly,
                            "serial": serial_clean,
                            "dung_luong": capacity,
                        })
    except Exception:
        pass

    # 2. Máy in kết nối qua USB
    try:
        path = r"SYSTEM\CurrentControlSet\Enum\USBPRINT"
        with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, path) as root:
            i = 0
            while True:
                try:
                    dev_type_name = winreg.EnumKey(root, i)
                except OSError:
                    break
                i += 1
                with winreg.OpenKey(root, dev_type_name) as dev_key:
                    j = 0
                    while True:
                        try:
                            inst_name = winreg.EnumKey(dev_key, j)
                        except OSError:
                            break
                        j += 1
                        friendly = dev_type_name
                        try:
                            with winreg.OpenKey(dev_key, inst_name) as inst:
                                friendly = _read_dev_name(inst, friendly)
                        except Exception:
                            pass
                        items.append({
                            "loai": "Máy in (kết nối USB)",
                            "ten": friendly,
                            "serial": inst_name,
                            "dung_luong": "-",
                        })
    except Exception:
        pass

    # 3. Điện thoại / thiết bị di động (MTP/PTP - Windows Portable Devices)
    try:
        path = r"SYSTEM\CurrentControlSet\Enum\WPD"
        with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, path) as root:
            i = 0
            while True:
                try:
                    dev_type_name = winreg.EnumKey(root, i)
                except OSError:
                    break
                i += 1
                with winreg.OpenKey(root, dev_type_name) as dev_key:
                    j = 0
                    while True:
                        try:
                            inst_name = winreg.EnumKey(dev_key, j)
                        except OSError:
                            break
                        j += 1
                        friendly = dev_type_name
                        try:
                            with winreg.OpenKey(dev_key, inst_name) as inst:
                                friendly = _read_dev_name(inst, friendly)
                        except Exception:
                            pass
                        items.append({
                            "loai": "Điện thoại / Thiết bị di động",
                            "ten": friendly,
                            "serial": inst_name,
                            "dung_luong": "-",
                        })
    except Exception:
        pass

    # 4. Các thiết bị USB khác (chuột, bàn phím, webcam, wifi USB, máy quét, đầu đọc thẻ...)
    # - tương tự USBDeview: bỏ hub/controller, thiết bị lưu trữ, composite và loại đã liệt kê
    try:
        path = r"SYSTEM\CurrentControlSet\Enum\USB"
        with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, path) as root:
            i = 0
            while True:
                try:
                    dev_type_name = winreg.EnumKey(root, i)
                except OSError:
                    break
                i += 1
                if "ROOT_HUB" in dev_type_name.upper():
                    continue
                with winreg.OpenKey(root, dev_type_name) as dev_key:
                    j = 0
                    while True:
                        try:
                            inst_name = winreg.EnumKey(dev_key, j)
                        except OSError:
                            break
                        j += 1
                        if _normalize_serial(inst_name) in claimed_serials:
                            continue
                        friendly = dev_type_name
                        try:
                            with winreg.OpenKey(dev_key, inst_name) as inst:
                                friendly = _read_dev_name(inst, friendly)
                        except Exception:
                            pass
                        friendly = _clean_desc(friendly)
                        lower = friendly.lower()
                        if ("hub" in lower or "mass storage" in lower
                                or lower in ("usb composite device", "usb composite")
                                or not friendly):
                            continue
                        items.append({
                            "loai": "Thiết bị USB khác",
                            "ten": friendly,
                            "serial": inst_name,
                            "dung_luong": "-",
                        })
    except Exception:
        pass

    # Khử trùng (loại + tên + seri)
    seen = set()
    unique_items = []
    for it in items:
        key = (it["loai"], it["ten"], it["serial"])
        if key not in seen:
            seen.add(key)
            unique_items.append(it)
    return unique_items

def format_peripheral_history_text(items, max_items=25):
    if not items:
        return ["Không."]
    lines = []
    for idx, it in enumerate(items[:max_items], start=1):
        lines.append(f'{idx}. Loại: {it["loai"]} | Seri: {it["serial"]} | Dung lượng: {it["dung_luong"]} | Tên: {it["ten"]}')
    if len(items) > max_items:
        lines.append(f"... và {len(items) - max_items} thiết bị khác.")
    return lines

def get_installed_hotfixes():
    data = run_ps("Get-HotFix | Select-Object -ExpandProperty HotFixID")
    if not data:
        return set()
    if isinstance(data, str):
        data = [data]
    return set(x.upper() for x in data)

def parse_vuln_file(path):
    entries = []
    if not os.path.exists(path):
        return entries
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):
                continue
            parts = line.split("|")
            if len(parts) < 4:
                continue
            cve_id, name, severity, kbs = parts[0], parts[1], parts[2], parts[3]
            min_build = parts[4] if len(parts) > 4 else ""
            entries.append({
                "cve": cve_id.strip(),
                "name": name.strip(),
                "severity": severity.strip(),
                "kbs": [k.strip().upper() for k in kbs.split(";") if k.strip()],
                "min_build": min_build.strip(),
                "fixed_builds": _parse_fixed_builds(min_build),
            })
    return entries


def _parse_fixed_builds(raw):
    """Bóc tách cột 'bản dựng đã vá' thành dict {bản_dựng_gốc: UBR_tối_thiểu}.
    Hỗ trợ:
      - '19045.4291;22631.3447'  -> {19045: 4291, 22631: 3447} (so theo từng dòng Windows)
      - '19041'                  -> {19041: 0}  (định dạng cũ: chỉ so bản dựng gốc)
    Nhờ tách theo từng dòng Windows (19045=22H2, 22631=23H2...) nên máy đã cập nhật tích lũy
    KHÔNG còn bị báo thừa lỗ hổng chỉ vì thiếu các KB cũ đã bị thay thế."""
    result = {}
    for phan in re.split(r"[;,]", raw or ""):
        phan = phan.strip()
        if not phan:
            continue
        if "." in phan:
            goc, _, ubr = phan.partition(".")
            try:
                result[int(goc)] = int(re.sub(r"[^0-9]", "", ubr) or 0)
            except ValueError:
                continue
        else:
            try:
                result[int(phan)] = 0
            except ValueError:
                continue
    return result


def _da_va_theo_build(fixed_builds, os_build, os_ubr):
    """Máy được coi là đã vá theo bản dựng khi bản dựng gốc khớp một dòng Windows trong danh sách
    và UBR của máy >= UBR đã vá của dòng đó. Nếu bản dựng gốc của máy MỚI HƠN mọi mốc trong danh
    sách (dòng Windows mới hơn) thì cũng coi như đã vá."""
    if not fixed_builds:
        return False
    try:
        b = int(os_build)
    except (ValueError, TypeError):
        return False
    u = os_ubr if isinstance(os_ubr, int) else 0
    if b in fixed_builds:
        return u >= fixed_builds[b]
    # Bản dựng gốc không có trong danh sách: nếu máy mới hơn mốc lớn nhất -> đã qua đợt vá đó.
    return b > max(fixed_builds)


def doc_ngay_du_lieu(path):
    """Đọc dòng '# DATA_VERSION: <ngày>' mà công cụ TỔNG HỢP ghi ở đầu file CVE/IOC khi cập nhật.
    Trả về chuỗi ngày (vd '20/09/2026') hoặc None. Dùng để in lên biên bản, phục vụ tính pháp lý:
    biết một biên bản được kiểm tra bằng bộ dữ liệu cập nhật đến thời điểm nào."""
    try:
        with open(path, encoding="utf-8") as f:
            for _ in range(30):  # chỉ dò trong phần đầu file
                line = f.readline()
                if not line:
                    break
                m = re.search(r"#\s*DATA_VERSION:\s*(.+)$", line)
                if m:
                    return m.group(1).strip()
    except Exception:
        pass
    return None


def scan_os_vulnerabilities(vuln_file_path, os_build, os_ubr=None):
    installed = get_installed_hotfixes()
    entries = parse_vuln_file(vuln_file_path)
    findings = []
    for e in entries:
        patched = any(kb in installed for kb in e["kbs"])
        if not patched and _da_va_theo_build(e["fixed_builds"], os_build, os_ubr):
            patched = True
        if not patched and e["kbs"]:
            findings.append(f'{e["cve"]} - {e["name"]} (Mức độ: {e["severity"]}) - CHƯA phát hiện bản vá')
    return findings

def format_vuln_text(findings):
    if not findings:
        return ["Không phát hiện lỗ hổng chưa có bản vá trong danh sách đối chiếu."]
    cves = []
    for f in findings:
        m = re.search(r"CVE-\d{4}-\d+", f)
        if m:
            cves.append(m.group(0))
    if not cves:
        cves = [str(f).split(" - ")[0] for f in findings]
    return [f"Phát hiện {len(findings)} lỗ hổng chưa có bản vá (gồm: {'; '.join(cves)})."]

def parse_malware_signatures(path):
    families = []
    current_family = None
    if not os.path.exists(path):
        return families

    with open(path, encoding="utf-8") as f:
        for raw_line in f:
            line = raw_line.strip()
            if not line:
                continue
            if line.startswith("["):
                break
            if line.startswith("# ---"):
                current_family = line.strip("# -").strip()
                families.append({"name": current_family, "process": [], "sha256": [], "file": [], "domain": [], "ip": []})
                continue
            if line.startswith("#"):
                continue
            m = re.match(r"^(process|sha256|file|domain|ip):(.+)$", line)
            if m and families:
                kind, val = m.group(1), m.group(2).strip()
                families[-1][kind].append(val)
    return families

def get_running_processes():
    data = run_ps("Get-Process | Select-Object ProcessName, Path")
    if not data:
        return []
    if isinstance(data, dict):
        data = [data]
    return data

def _netstat_established_remote_ips():
    """Dự phòng cho Win7: lấy IP từ xa qua lệnh netstat."""
    ips = set()
    try:
        result = subprocess.run(["netstat", "-ano"], capture_output=True, text=True,
                                timeout=15, encoding="utf-8", errors="ignore",
                                startupinfo=_hidden_startupinfo())
    except Exception:
        return ips
    for line in (result.stdout or "").splitlines():
        line = line.strip()
        if not line.lower().startswith("tcp") or "ESTABLISHED" not in line.upper():
            continue
        parts = line.split()
        if len(parts) < 4:
            continue
        remote = parts[2]
        if remote.startswith("["):
            remote = remote[1:remote.find("]")]
        elif ":" in remote:
            remote = remote.rsplit(":", 1)[0]
        if remote and not remote.startswith(("0.0.0.0", "127.", "[::", "::")):
            ips.add(remote)
    return ips

def get_active_remote_ips():
    if USE_MODERN_CMDLETS:
        data = run_ps("Get-NetTCPConnection -State Established -ErrorAction SilentlyContinue | "
                      "Select-Object -ExpandProperty RemoteAddress")
        if data:
            if isinstance(data, str):
                data = [data]
            return set(data)
    return _netstat_established_remote_ips()

def _ipconfig_dns_cache():
    """Dự phòng cho Win7: đọc cache DNS qua lệnh ipconfig /displaydns."""
    domains = set()
    try:
        result = subprocess.run(["ipconfig", "/displaydns"], capture_output=True, text=True,
                                timeout=15, encoding="utf-8", errors="ignore",
                                startupinfo=_hidden_startupinfo())
    except Exception:
        return domains
    for line in (result.stdout or "").splitlines():
        if "Record Name" in line:
            name = line.split(":", 1)[-1].strip().lower()
            if name:
                domains.add(name)
    return domains

def get_dns_cache_domains():
    if USE_MODERN_CMDLETS:
        data = run_ps("Get-DnsClientCache -ErrorAction SilentlyContinue | Select-Object -ExpandProperty Entry")
        if data:
            if isinstance(data, str):
                data = [data]
            return set(d.lower() for d in data)
    return _ipconfig_dns_cache()

def sha256_of_file(path, max_size=200 * 1024 * 1024):
    try:
        if not os.path.isfile(path) or os.path.getsize(path) > max_size:
            return None
        h = hashlib.sha256()
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1024 * 1024), b""):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return None

COMMON_SCAN_DIRS = []
if IS_WINDOWS:
    userprofile = os.environ.get("USERPROFILE", "")
    appdata = os.environ.get("APPDATA", "")
    COMMON_SCAN_DIRS = [
        os.path.join(userprofile, "Downloads"),
        os.path.join(userprofile, "Desktop"),
        os.path.join(appdata, "Microsoft", "Windows", "Start Menu", "Programs", "Startup"),
        os.environ.get("TEMP", ""),
    ]
    COMMON_SCAN_DIRS = [d for d in COMMON_SCAN_DIRS if d and os.path.isdir(d)]

def scan_files_shallow(dirs, filename_patterns, max_depth=2):
    matches = []
    for base in dirs:
        base_depth = base.rstrip(os.sep).count(os.sep)
        for root, _, files in os.walk(base):
            depth = root.rstrip(os.sep).count(os.sep) - base_depth
            if depth > max_depth:
                continue
            for fname in files:
                for pattern in filename_patterns:
                    if fnmatch.fnmatch(fname.lower(), pattern.lower()):
                        matches.append(os.path.join(root, fname))
    return matches

def _domain_trong_cache(domain, dns_cache):
    """So khớp tên miền IOC với cache DNS theo dạng khớp đúng hoặc là tên miền con
    (vd IOC 'evil.com' khớp 'a.evil.com' nhưng KHÔNG khớp 'notevil.com') - tránh báo nhầm."""
    dom = domain.lower().strip(".")
    if not dom:
        return False
    for d in dns_cache:
        d = d.lower().strip(".")
        if d == dom or d.endswith("." + dom):
            return True
    return False


def scan_malware(malware_file_path):
    families = parse_malware_signatures(malware_file_path)
    if not families:
        return []

    running = get_running_processes()
    running_names = {p.get("ProcessName", "").lower() for p in running if p}
    active_ips = get_active_remote_ips()
    dns_domains = get_dns_cache_domains()

    all_file_patterns = []
    for fam in families:
        all_file_patterns.extend(fam["file"])
    matched_files = scan_files_shallow(COMMON_SCAN_DIRS, all_file_patterns) if all_file_patterns else []
    matched_files_lower = {os.path.basename(p).lower(): p for p in matched_files}

    # Chỉ hash file thực thi của các tiến trình thuộc danh sách IOC (tránh quét toàn bộ máy)
    sha_proc_names = set()
    for fam in families:
        if fam["sha256"]:
            sha_proc_names.update(p.lower() for p in fam["process"])

    findings = []
    for fam in families:
        hits = []
        for pname in fam["process"]:
            if pname.lower() in running_names:
                hits.append(f"tiến trình đang chạy '{pname}'")
        for fpattern in fam["file"]:
            for base_name, full_path in matched_files_lower.items():
                if fnmatch.fnmatch(base_name, fpattern.lower()):
                    hits.append(f"file nghi vấn '{full_path}'")
        for ip in fam["ip"]:
            if ip in active_ips:
                hits.append(f"kết nối mạng tới IP nghi vấn '{ip}'")
        for domain in fam["domain"]:
            if _domain_trong_cache(domain, dns_domains):
                hits.append(f"có trong cache DNS: '{domain}'")
        if fam["sha256"]:
            for p in running:
                pname = (p.get("ProcessName") or "").lower()
                if pname not in sha_proc_names:
                    continue
                ppath = p.get("Path")
                if ppath:
                    h = sha256_of_file(ppath)
                    if h and h.lower() in [s.lower() for s in fam["sha256"]]:
                        hits.append(f"trùng khớp SHA256 với tiến trình '{p.get('ProcessName')}'")
        if hits:
            findings.append(f'{fam["name"]}: ' + "; ".join(hits))
    return findings

def format_malware_text(findings):
    if not findings:
        return ["Không phát hiện dấu hiệu (IOC) trùng khớp với danh sách mã độc."]
    lines = [f"{i+1}. {f}" for i, f in enumerate(findings)]
    lines.append("Cần kiểm tra, xác minh thêm và xử lý kịp thời.")
    return lines

# ============================================================
# ĐIỀN VÀO FILE DOCX
# ============================================================
def merge_para_runs(paragraph):
    if len(paragraph.runs) <= 1:
        return
    full_text = "".join(r.text for r in paragraph.runs)
    paragraph.runs[0].text = full_text
    for r in paragraph.runs[1:]:
        r.text = ""

def append_lines_to_para(para, lines):
    if not lines:
        return
    last_run = para.runs[-1] if para.runs else para.add_run()
    for line in lines:
        last_run.add_break(WD_BREAK.LINE)
        r = para.add_run(str(line))
        r.font.name = "Times New Roman"
        last_run = r

def append_inline_text(para, lines):
    """Ghi kết quả ngay tiếp nối trên cùng dòng với đoạn văn (sau dấu hai chấm)."""
    if not lines:
        return
    text = " ".join(str(l) for l in lines)
    r = para.add_run(text)
    r.font.name = "Times New Roman"

def set_table_cell_text(cell, text):
    merge_para_runs(cell.paragraphs[0])
    if cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].text = text
    else:
        run = cell.paragraphs[0].add_run(text)
        run.font.name = "Times New Roman"

def replace_in_cell_by_pattern(cell, pattern_replacements):
    for para in cell.paragraphs:
        merge_para_runs(para)
        for run in para.runs:
            for pattern, new_text in pattern_replacements:
                run.text = re.sub(pattern, new_text, run.text)

def tick_checkbox_nth(paragraph, nth_checkbox):
    from docx.oxml.ns import qn
    checkbox_count = 0
    for run in paragraph.runs:
        sym = run._element.find(qn("w:sym"))
        if sym is not None:
            if checkbox_count == nth_checkbox:
                sym.set(qn("w:char"), "F0FE")
                return True
            checkbox_count += 1
    return False

def _password_status_legacy(username):
    """Win7/8: kiểm tra thuộc tính PasswordRequired qua WMI, dự phòng bằng 'net user'."""
    data = run_ps(
        "Get-WmiObject Win32_UserAccount -Filter \"LocalAccount=True AND Name='" +
        username + "'\" -ErrorAction SilentlyContinue | Select-Object -First 1 -ExpandProperty PasswordRequired"
    )
    if isinstance(data, bool):
        return data
    try:
        result = subprocess.run(["net", "user", username], capture_output=True, text=True,
                                timeout=15, encoding="utf-8", errors="ignore",
                                startupinfo=_hidden_startupinfo())
        out = (result.stdout or "")
        m = re.search(r"(?im)^PASSWORD REQUIRED\s+(YES|NO)", out)
        if m:
            return m.group(1).upper() == "YES"
    except Exception:
        pass
    return None

def get_password_status():
    """Trả True nếu tài khoản hiện tại CÓ mật khẩu, False nếu KHÔNG, None nếu không xác định."""
    try:
        username = os.environ.get("USERNAME", "")
        if not username:
            return None
        if USE_LOCALUSER:
            data = run_ps(f"Get-LocalUser -Name '{username}' -ErrorAction SilentlyContinue | "
                          "Select-Object -ExpandProperty PasswordRequired")
            if isinstance(data, bool):
                return data
        return _password_status_legacy(username)
    except Exception:
        return None

def _same_subnet(ip_a, ip_b):
    """So sánh 2 IP có cùng phân đoạn /24 hay không (dùng để phát hiện máy LAN thật)."""
    try:
        if not ip_a or not ip_b:
            return False
        pa, pb = ip_a.split("."), ip_b.split(".")
        if len(pa) != 4 or len(pb) != 4:
            return False
        for x in pa + pb:
            if not x.isdigit() or not (0 <= int(x) <= 255):
                return False
        return pa[:3] == pb[:3]
    except Exception:
        return False

def _arp_neighbors():
    """Dự phòng cho Win7: lấy danh sách IP cùng LAN qua bảng ARP."""
    ips = set()
    try:
        result = subprocess.run(["arp", "-a"], capture_output=True, text=True,
                                timeout=15, encoding="utf-8", errors="ignore",
                                startupinfo=_hidden_startupinfo())
    except Exception:
        return ips
    for line in (result.stdout or "").splitlines():
        parts = line.split()
        if len(parts) >= 3 and re.match(r"^\d{1,3}(\.\d{1,3}){3}$", parts[0]):
            if parts[-1].lower() == "dynamic":
                ips.add(parts[0])
    return ips

def _get_network_classification_legacy():
    """Phiên bản Win7/PowerShell 2.0: dùng WMI + ARP + Registry để phân loại mạng."""
    cat = None
    try:
        path = r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\NetworkList\Profiles"
        with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, path) as root:
            i = 0
            while True:
                try:
                    sub = winreg.EnumKey(root, i)
                except OSError:
                    break
                i += 1
                try:
                    with winreg.OpenKey(root, sub) as k:
                        if winreg.QueryValueEx(k, "Category")[0] == 2:
                            cat = "DomainAuthenticated"
                            break
                except Exception:
                    continue
    except Exception:
        pass

    gw = run_ps(
        "Get-WmiObject Win32_NetworkAdapterConfiguration -Filter \"IPEnabled=True\" "
        "-ErrorAction SilentlyContinue | Where-Object {$_.DefaultIPGateway} | "
        "Select-Object -ExpandProperty DefaultIPGateway"
    )
    if isinstance(gw, list):
        gw = gw[0] if gw else None

    local = run_ps(
        "Get-WmiObject Win32_NetworkAdapterConfiguration -Filter \"IPEnabled=True\" "
        "-ErrorAction SilentlyContinue | Select-Object -ExpandProperty IPAddress"
    )
    if isinstance(local, str):
        local = [local]
    local = [ip for ip in (local or []) if isinstance(ip, str) and ":" not in ip
             and not ip.startswith(("127.", "169.254.", "0.0.0.0"))]

    return {"cat": cat, "gw": gw, "local": local, "neighbors": _arp_neighbors()}

def get_network_classification():
    if USE_MODERN_CMDLETS:
        data = run_ps(
            "$cat = (Get-NetConnectionProfile -ErrorAction SilentlyContinue | "
            "Select-Object -First 1 -ExpandProperty NetworkCategory); "
            "$gw = (Get-NetRoute -DestinationPrefix '0.0.0.0/0' -ErrorAction SilentlyContinue | "
            "Sort-Object RouteMetric | Select-Object -First 1 -ExpandProperty NextHop); "
            "$local = Get-NetIPAddress -AddressFamily IPv4 -ErrorAction SilentlyContinue | "
            "Where-Object {$_.IPAddress -notmatch '^(127|169\\.254|0\\.0\\.0\\.0)'} | "
            "Select-Object -ExpandProperty IPAddress; "
            "$nb = Get-NetNeighbor -AddressFamily IPv4 -State Reachable,Stale -ErrorAction SilentlyContinue | "
            "Select-Object -ExpandProperty IPAddress; "
            "[PSCustomObject]@{cat=$cat; gw=$gw; local=$local; neighbors=$nb}"
        )
    else:
        data = _get_network_classification_legacy()
    if not isinstance(data, dict):
        return None

    category = data.get("cat")
    if isinstance(category, str) and category == "DomainAuthenticated":
        return "noi_bo"

    gateway = data.get("gw")
    if isinstance(gateway, list):
        gateway = gateway[0] if gateway else None

    local_ips = data.get("local") or []
    if isinstance(local_ips, str):
        local_ips = [local_ips]
    local_ips = [ip for ip in local_ips if ip]

    # Xác định IP mạng LAN chính (cùng phân đoạn với gateway mặc định) -
    # tránh nhầm subnet của adapter ảo (Docker/VMware/Hyper-V...)
    lan_net = None
    for ip in local_ips:
        if _same_subnet(ip, gateway):
            lan_net = ip
            break
    if lan_net is None:
        lan_net = gateway or (local_ips[0] if local_ips else None)
    if not lan_net:
        return None

    raw_neighbors = data.get("neighbors")
    if raw_neighbors is None:
        return None
    if isinstance(raw_neighbors, set):
        neighbor_list = list(raw_neighbors)
    elif isinstance(raw_neighbors, list):
        neighbor_list = raw_neighbors
    else:
        neighbor_list = [raw_neighbors]
    excluded = {gateway, "127.0.0.1", "0.0.0.0"}
    filtered = [
        ip for ip in neighbor_list
        if ip and ip not in excluded
        and not ip.startswith(("224.", "239.", "255.", "169.254."))
        and _same_subnet(ip, lan_net)
    ]
    return "noi_bo" if filtered else "doc_lap"


TOTAL_PROGRESS_STEPS = 10

def fill_form(template_path, output_path, malware_file, vuln_file, manual_data=None, on_progress=None):
    manual_data = manual_data or {}
    total = TOTAL_PROGRESS_STEPS

    # Tự động điền ngày giờ hiện tại khi chạy chế độ --no-manual (không nhập tay)
    now = datetime.now()
    manual_data.setdefault("gio", now.strftime("%H"))
    manual_data.setdefault("phut", now.strftime("%M"))
    manual_data.setdefault("ngay", now.strftime("%d"))
    manual_data.setdefault("thang", now.strftime("%m"))
    manual_data.setdefault("nam", now.strftime("%Y"))

    def step(n, message):
        print(message)
        if on_progress:
            on_progress(n, total, message)

    step(1, "Đang thu thập thông tin hệ điều hành...")
    os_info = get_os_info()

    step(2, "Đang lấy tên máy, MAC, IP...")
    computer_name = get_computer_name()
    mac_addr, ip_addr = get_network_basic()

    step(3, "Đang lấy cấu hình phần cứng...")
    hw = get_hardware_config()

    step(4, "Đang xác định phần mềm diệt virus...")
    av_name = get_antivirus_name()

    step(5, "Đang xác định phần mềm, mật khẩu...")
    top_apps = get_top_apps(2)
    password_status = get_password_status()
    network_class = get_network_classification()

    step(6, "Đang kiểm tra tình trạng thiết bị / kết nối Internet...")
    online_now = check_internet_now()
    device_status = get_device_status_text()
    
    internet_history = get_internet_history_text(online_now)

    step(7, "Đang đối chiếu lỗ hổng bảo mật hệ điều hành...")
    vuln_findings = []
    if not os.path.exists(vuln_file):
        vuln_lines = ["CẢNH BÁO: Không tìm thấy file danh sách CVE; bỏ qua bước đối chiếu lỗ hổng."]
    else:
        vuln_findings = scan_os_vulnerabilities(vuln_file, os_info["build"], os_info.get("ubr"))
        vuln_lines = format_vuln_text(vuln_findings)

    step(8, "Đang đối chiếu dấu hiệu mã độc (IOC)...")
    malware_findings = []
    if not os.path.exists(malware_file):
        malware_lines = ["CẢNH BÁO: Không tìm thấy file danh sách IOC; bỏ qua bước quét mã độc."]
    else:
        malware_findings = scan_malware(malware_file)
        malware_lines = format_malware_text(malware_findings)

    step(9, "Đang lấy lịch sử thiết bị ngoại vi...")
    peripheral_items = get_peripheral_history()
    peripheral_lines = format_peripheral_history_text(peripheral_items)
    if not is_admin():
        peripheral_lines.append("Lưu ý: Công cụ chưa chạy với quyền Administrator nên có thể thiếu dữ liệu.")

    step(10, "Đang điền dữ liệu vào biên bản và lưu file...")
    doc = docx.Document(template_path)

    # Kiểm tra cấu trúc template để báo lỗi rõ ràng thay vì lỗi âm thầm
    if len(doc.tables) < 1 or len(doc.tables[0].rows) < 9:
        raise RuntimeError("File mẫu biên bản không đúng cấu trúc: thiếu bảng thông tin (cần ≥ 9 dòng).")
    if len(doc.paragraphs) < 14:
        raise RuntimeError("File mẫu biên bản không đúng cấu trúc: thiếu các đoạn văn cần thiết.")
    table = doc.tables[0]

    p5 = doc.paragraphs[5]
    merge_para_runs(p5)
    text_p5_clean = ""
    for run in p5.runs:
        text_p5_clean += re.sub(r'^[…\.\s]+', '', run.text)

    p4 = doc.paragraphs[4]
    dia_diem_full = manual_data.get('dia_diem', '') + text_p5_clean
    text_p4 = f"Vào hồi {manual_data.get('gio', '')} giờ {manual_data.get('phut', '')} phút, ngày {manual_data.get('ngay', '')} tháng {manual_data.get('thang', '')} năm {manual_data.get('nam', '')}, tại {dia_diem_full}"
    
    p4.text = text_p4
    p4.runs[0].font.name = 'Times New Roman'
    p4.runs[0].font.size = docx.shared.Pt(12)

    p5.text = ""
    p5.paragraph_format.space_before = docx.shared.Pt(0)
    p5.paragraph_format.space_after = docx.shared.Pt(0)
    p5.paragraph_format.line_spacing = docx.shared.Pt(0)
    p5_run = p5.add_run()
    p5_run.font.size = docx.shared.Pt(1)

    p7 = doc.paragraphs[7]
    merge_para_runs(p7)
    for run in p7.runs:
        if manual_data.get("ten_can_bo"):
            run.text = re.sub(r'(Đ/c\s*)[…\.]+(\s*–\s*)[…\.]+',
                               r'\g<1>' + manual_data.get("ten_can_bo", "") + r'\g<2>' + manual_data.get("chuc_vu", ""),
                               run.text)
        if manual_data.get("ten_doi_tuong"):
            run.text = re.sub(r'(đồng chí:\s*)[…\.]+\s*[…\.]+',
                               r'\g<1>' + manual_data.get("ten_doi_tuong", "") + ' ',
                               run.text)
                           
    set_table_cell_text(table.rows[0].cells[1], computer_name)
    set_table_cell_text(table.rows[1].cells[1], os_info["os_display"])
    set_table_cell_text(table.rows[2].cells[1], os_info["ngay_cai"] or "Không xác định")
    set_table_cell_text(table.rows[3].cells[1], mac_addr)
    set_table_cell_text(table.rows[4].cells[1], ip_addr)

    hw_text = f"CPU: {hw.get('cpu', 'Không xác định')}\n" \
              f"RAM: {hw.get('ram_gb', '?')} GB\n" \
              f"Ổ cứng loại: {hw.get('loai_o_cung', 'Không xác định')}\n" \
              f"Dung lượng: {hw.get('dung_luong_o_cung', 'Không xác định')}"
    set_table_cell_text(table.rows[5].cells[1], hw_text)
    
    set_table_cell_text(table.rows[6].cells[1], av_name)
    set_table_cell_text(table.rows[7].cells[1], top_apps)
    
    conn_type = get_connection_type() if online_now else ""
    set_table_cell_text(table.rows[8].cells[1], f"Kết nối Internet: {'Có' + conn_type if online_now else 'Không'}.")

    # Bảng ký tên cuối biên bản: ô hàng 2 cột 1 = cán bộ kiểm tra; ô hàng 2 cột 3 = người/đơn vị
    try:
        if len(doc.tables) > 1:
            sign_table = doc.tables[1]
            if manual_data.get("ten_can_bo"):
                set_table_cell_text(sign_table.cell(1, 0), manual_data["ten_can_bo"])
            if manual_data.get("ten_doi_tuong"):
                set_table_cell_text(sign_table.cell(1, 2), manual_data["ten_doi_tuong"])
    except Exception:
        pass

    p10 = doc.paragraphs[10]
    if password_status is True:
        tick_checkbox_nth(p10, 0)
    elif password_status is False:
        tick_checkbox_nth(p10, 1)

    p13 = doc.paragraphs[13]
    if network_class == "noi_bo":
        tick_checkbox_nth(p13, 0)
    elif network_class == "doc_lap":
        tick_checkbox_nth(p13, 1)
    if online_now:
        tick_checkbox_nth(p13, 2)

    for para in doc.paragraphs:
        txt = para.text.strip().lower()
        if txt.startswith("- tình trạng thiết bị tại thời điểm kiểm tra (tem"):
            append_lines_to_para(para, [device_status])
        elif txt.startswith("- lỗ hổng bảo mật hệ điều hành"):
            append_lines_to_para(para, vuln_lines)
        elif txt.startswith("- mã độc"):
            append_lines_to_para(para, malware_lines)
        elif txt.startswith("- lịch sử kết nối các thiết bị ngoại vi"):
            append_lines_to_para(para, peripheral_lines)
        elif txt.startswith("- lịch sử kết nối internet"):
            append_inline_text(para, internet_history)
        elif txt.startswith("- các nội dung khác"):
            ghi_chu = []
            ngay_cve = doc_ngay_du_lieu(vuln_file)
            ngay_ioc = doc_ngay_du_lieu(malware_file)
            if ngay_cve or ngay_ioc:
                ghi_chu.append(f"Đối chiếu bằng bộ dữ liệu cập nhật: lỗ hổng đến {ngay_cve or 'không rõ'}; "
                               f"mã độc đến {ngay_ioc or 'không rõ'}.")
            if not is_admin():
                ghi_chu.append("Công cụ chưa chạy với quyền Administrator nên một số dữ liệu có thể chưa đầy đủ.")
            if ghi_chu:
                append_lines_to_para(para, ghi_chu)

    # Thụt đầu dòng 1cm cho toàn bộ đoạn văn ngoài bảng (đồng bộ với văn bản thân biên bản)
    for para in doc.paragraphs:
        if not para.text.strip():
            continue
        try:
            align = para.paragraph_format.alignment
            if align in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT):
                continue
            para.paragraph_format.first_line_indent = docx.shared.Cm(1)
        except Exception:
            pass

    # Xóa đoạn văn trống thừa đầu biên bản (p5) để không còn khoảng trống xuống dòng
    try:
        p5._element.getparent().remove(p5._element)
    except Exception:
        pass

    doc.save(output_path)

    # Ghi file dữ liệu kèm (.json) có chữ ký HMAC để công cụ tổng hợp xác minh biên bản không bị sửa.
    try:
        _ghi_sidecar_toan_ven(
            output_path,
            os_info=os_info, computer_name=computer_name, mac=mac_addr, ip=ip_addr,
            hw=hw, av=av_name, top_apps=top_apps, password_status=password_status,
            network_class=network_class, online=online_now,
            vuln_findings=vuln_findings, malware_findings=malware_findings,
            peripheral_items=peripheral_items, manual_data=manual_data,
            vuln_file=vuln_file, malware_file=malware_file,
        )
    except Exception as e:
        print(f"(Cảnh báo) Không ghi được file kèm toàn vẹn: {e}")

    print(f"\nĐã điền xong biên bản: {output_path}")
    return output_path


# ============================================================
# TOÀN VẸN BIÊN BẢN (chữ ký HMAC-SHA256) - phục vụ tính pháp lý
# ============================================================
# Khoá mặc định khi chưa nhúng khoá riêng lúc build. Công cụ TỔNG HỢP có thể sinh khoá riêng và
# nhúng vào .exe (file bien_ban_key.dat) để chỉ bộ công cụ của đơn vị mới tạo/kiểm được chữ ký hợp lệ.
DEFAULT_INTEGRITY_KEY = "ANATTT-CAX-TriPhu-2026-bien-ban-integrity-default-key"


def _integrity_key():
    p = _data_file("bien_ban_key.dat")
    try:
        if os.path.exists(p):
            with open(p, "r", encoding="utf-8") as f:
                k = f.read().strip()
                if k:
                    return k
    except Exception:
        pass
    return os.environ.get("ANATTT_INTEGRITY_KEY", DEFAULT_INTEGRITY_KEY)


def _canonical_json(obj):
    return json.dumps(obj, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _tinh_hmac(key, obj):
    import hmac as _hmac
    return _hmac.new(key.encode("utf-8"), _canonical_json(obj).encode("utf-8"), hashlib.sha256).hexdigest()


def _sha256_file(path):
    try:
        h = hashlib.sha256()
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1024 * 1024), b""):
                h.update(chunk)
        return h.hexdigest()
    except Exception:
        return None


def _ghi_sidecar_toan_ven(output_path, os_info, computer_name, mac, ip, hw, av, top_apps,
                          password_status, network_class, online, vuln_findings, malware_findings,
                          peripheral_items, manual_data, vuln_file, malware_file):
    """Ghi file <biên_bản>.attt.json chứa dữ liệu thu thập có cấu trúc + chữ ký HMAC.
    Công cụ tổng hợp ưu tiên đọc file này (đáng tin hơn regex trên docx) và phát hiện nếu bị sửa."""
    cves = []
    for f in vuln_findings:
        m = re.search(r"CVE-\d{4}-\d+", f)
        if m:
            cves.append(m.group(0))
    ho_ma_doc = [str(f).split(":", 1)[0].strip() for f in malware_findings]

    payload = {
        "schema": "anattt-bienban/1",
        "app_version": APP_VERSION,
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "computer_name": computer_name,
        "os_display": os_info.get("os_display", ""),
        "os_caption": os_info.get("ten_he_dieu_hanh", ""),
        "os_build_full": os_info.get("build_full", ""),
        "ngay_cai": os_info.get("ngay_cai", ""),
        "ban_quyen": os_info.get("ban_quyen", ""),
        "mac": mac,
        "ip": ip,
        "hardware": hw,
        "antivirus": av,
        "top_apps": top_apps,
        "password_has": password_status,
        "network_class": network_class,
        "online": bool(online),
        "vuln_cves": cves,
        "vuln_count": len(vuln_findings),
        "vuln_raw": vuln_findings,
        "malware_families": ho_ma_doc,
        "malware_raw": malware_findings,
        "peripherals": peripheral_items,
        "data_version_cve": doc_ngay_du_lieu(vuln_file) or "",
        "data_version_ioc": doc_ngay_du_lieu(malware_file) or "",
        "manual": manual_data,
        "docx_file": os.path.basename(output_path),
        "docx_sha256": _sha256_file(output_path),
    }
    goi = {"payload": payload, "hmac": _tinh_hmac(_integrity_key(), payload)}
    sidecar_path = os.path.splitext(output_path)[0] + ".attt.json"
    with open(sidecar_path, "w", encoding="utf-8") as f:
        json.dump(goi, f, ensure_ascii=False, indent=2)
    return sidecar_path


APP_VERSION = "1.7"
APP_NAME = "Công cụ kiểm tra ANATTT"

def _spawn_elevated_worker(payload_path):
    """Khởi chạy tiến trình worker với quyền Administrator qua UAC (runas).
    Trả về True nếu UAC được người dùng chấp thuận (ret > 32), False nếu từ chối hoặc lỗi."""
    if not IS_WINDOWS:
        return False
    import ctypes
    if is_packaged():
        file_to_run = sys.executable
        params = f'--run-worker "{payload_path}"'
    else:
        file_to_run = sys.executable
        script = os.path.abspath(__file__)
        params = f'"{script}" --run-worker "{payload_path}"'

    try:
        # SW_HIDE = 0 để không bật cửa sổ console phụ
        ret = ctypes.windll.shell32.ShellExecuteW(None, "runas", file_to_run, params, None, 0)
        return ret > 32
    except Exception:
        return False


def run_worker_from_payload(payload_path):
    """Chạy kiểm tra ngầm với quyền Administrator dựa trên payload JSON."""
    if not os.path.exists(payload_path):
        sys.exit(1)
    try:
        with open(payload_path, "r", encoding="utf-8") as f:
            payload = json.load(f)
    except Exception:
        sys.exit(1)

    template_path = payload.get("template_path")
    output_path = payload.get("output_path")
    malware_path = payload.get("malware_file")
    vuln_path = payload.get("vuln_file")
    manual_data = payload.get("manual_data", {})
    cleanup = bool(payload.get("cleanup", False))
    shutdown = bool(payload.get("shutdown", False))
    status_file = payload.get("status_file")
    no_open = bool(payload.get("no_open", False))

    def _write_status(data):
        if not status_file:
            return
        try:
            tmp = status_file + f".tmp.{os.getpid()}"
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False)
            os.replace(tmp, status_file)
        except Exception:
            pass

    def on_progress(step, total, message):
        _write_status({
            "status": "progress",
            "step": step,
            "total": total,
            "message": message
        })

    try:
        result_path = fill_form(
            template_path, output_path, malware_path, vuln_path,
            manual_data, on_progress=on_progress
        )
        scheduled = False
        if cleanup or shutdown:
            scheduled = bool(schedule_self_cleanup(shutdown=shutdown))

        _write_status({
            "status": "done",
            "output_path": result_path,
            "scheduled": scheduled,
            "shutdown": shutdown
        })

        if not no_open:
            open_result(result_path)
    except Exception as e:
        error_trace = traceback.format_exc()
        try:
            dump_path = dump_error_log_and_compress(error_trace)
            error_msg = f"Đã xảy ra sự cố kỹ thuật: {str(e)}\n\nHệ thống đã tự động sao lưu hiện trường lỗi tại:\n{dump_path}"
        except Exception as dump_e:
            error_msg = f"Lỗi: {str(e)}\n(Lưu ý: Không thể ghi file log do {str(dump_e)})"

        _write_status({
            "status": "error",
            "message": error_msg
        })
        sys.exit(1)


# ============================================================
# TỰ ĐỘNG VIẾT HOA CHỮ CÁI ĐẦU MỖI TỪ
# (Gõ tiếng Việt do Unikey/bàn phím tiếng Việt trên Windows đảm nhiệm,
#  giống hệt cách gõ tự do trong khung chat; KHÔNG can thiệp bàn phím để
#  tránh làm hỏng chữ như "Văn" -> "vanw", "Đức" -> "đucs".)
# ============================================================

def auto_capitalize(text):
    """Viết hoa chữ cái đầu MỖI TỪ, giữ nguyên các từ viết hoa sẵn có (CAX, BỘ...).

    - Bỏ qua mọi ký tự không phải chữ cái ở đầu từ (số, dấu ngoặc, gạch...) rồi
      viết hoa chữ cái đầu tiên.
    - Giữ NGUYÊN vẹn khoảng trắng (nhiều dấu cách, tab, xuống dòng).
    - Chuẩn hóa NFC trước để không làm hỏng chữ có dấu ghép rời (NFD)."""
    import unicodedata
    text = unicodedata.normalize("NFC", text)
    out = []
    i = 0
    n = len(text)
    while i < n:
        if text[i].isspace():
            out.append(text[i])
            i += 1
            continue
        j = i
        while j < n and not text[j].isspace():
            j += 1
        word = text[i:j]
        first = next((k for k, c in enumerate(word) if c.isalpha()), None)
        if first is not None:
            word = word[:first] + word[first].upper() + word[first + 1:]
        out.append(word)
        i = j
    return "".join(out)


def _bind_auto_cap(ent):
    """Tự động viết hoa chữ cái đầu mỗi từ khi rời ô.
    KHÔNG can thiệp bàn phím: Unikey/bàn phím tiếng Việt trên Windows gõ tự do
    (như khung chat), chương trình chỉ viết hoa lại khi kết thúc nhập ô."""
    def _on_focus_out(event):
        final = auto_capitalize(ent.get())
        if ent.get() != final:
            ent.delete(0, "end")
            ent.insert(0, final)

    ent.bind("<FocusOut>", _on_focus_out)

MANUAL_FIELDS_DEF = [
    ("gio", "Giờ kiểm tra (VD: 08)"),
    ("phut", "Phút (VD: 30)"),
    ("ngay", "Ngày (VD: 18)"),
    ("thang", "Tháng (VD: 07)"),
    ("nam", "Năm (VD: 2026)"),
    ("dia_diem", "Địa điểm (VD: Trụ sở CAX Tri Phú)"),
    ("ten_can_bo", "Tên cán bộ kiểm tra"),
    ("chuc_vu", "Chức vụ cán bộ kiểm tra"),
    ("ten_doi_tuong", "Tên người/đơn vị (đồng chí)"),
]

def open_result(path):
    if not IS_WINDOWS or not path:
        return
    abs_path = os.path.abspath(path)
    try:
        os.startfile(abs_path)
    except Exception:
        pass
    try:
        subprocess.Popen(f'explorer /select,"{abs_path}"')
    except Exception:
        pass


# ============================================================
# GIAO DIỆN (UI) BẰNG TTK 
# ============================================================
def run_gui_app(template_path, output_path, malware_file, vuln_file):
    import threading
    import queue as _queue
    
    root = tk.Tk()
    root.title(f"{APP_NAME} v{APP_VERSION} - CAX Tri Phú")
    root.geometry("640x780")
    root.resizable(False, False)
    root.configure(bg="#f1f5f9")

    style = ttk.Style()
    try:
        style.theme_use("clam")
    except Exception:
        pass

    BG = "#f1f5f9"
    PRIMARY = "#1e3a8a"
    GREEN = "#16a34a"

    style.configure(".", background=BG)
    style.configure("TLabel", background=BG, font=("Segoe UI", 10))
    style.configure("Field.TLabel", background=BG, font=("Segoe UI", 10), foreground="#1f2937")
    style.configure("TLabelframe", background=BG, bordercolor="#cbd5e1")
    style.configure("TLabelframe.Label", background=BG, foreground=PRIMARY,
                    font=("Segoe UI", 11, "bold"))
    style.configure("TEntry", fieldbackground="#ffffff", font=("Segoe UI", 10))
    style.configure("TCheckbutton", background=BG, font=("Segoe UI", 9))
    style.configure("Green.Horizontal.TProgressbar", background=GREEN, troughcolor="#e2e8f0")

    # Banner đầu có màu + tên phiên bản để phân biệt các phiên bản
    banner = tk.Frame(root, bg=PRIMARY, height=92)
    banner.pack(fill="x")
    banner.pack_propagate(False)
    tk.Label(banner, text="KIỂM TRA AN NINH, AN TOÀN THÔNG TIN",
             bg=PRIMARY, fg="white", font=("Segoe UI", 15, "bold")).pack(pady=(16, 0))
    tk.Label(banner, text=f"{APP_NAME} - Phiên bản v{APP_VERSION} | CAX Tri Phú",
             bg=PRIMARY, fg="#bfdbfe", font=("Segoe UI", 10)).pack(pady=(3, 0))

    # Thanh trạng thái dữ liệu + quyền: cho kiểm tra viên biết bộ CVE/IOC cập nhật đến ngày nào và
    # có đang chạy với quyền Administrator hay không TRƯỚC khi bấm kiểm tra.
    info_bar = tk.Frame(root, bg="#e8efff")
    info_bar.pack(fill="x")
    ngay_cve = doc_ngay_du_lieu(vuln_file) or "không rõ"
    ngay_ioc = doc_ngay_du_lieu(malware_file) or "không rõ"

    def _dem_dong(path, la_cve):
        try:
            n = 0
            with open(path, encoding="utf-8") as f:
                for d in f:
                    d = d.strip()
                    if d and not d.startswith("#") and (("|" in d) if la_cve else (":" in d)):
                        n += 1
            return n
        except Exception:
            return 0

    so_cve = _dem_dong(vuln_file, True)
    so_ioc = _dem_dong(malware_file, False)
    tk.Label(info_bar, bg="#e8efff", fg="#1e3a8a", font=("Segoe UI", 9),
             text=f"Dữ liệu lỗ hổng: {so_cve} mục (cập nhật {ngay_cve})   •   "
                  f"Dấu hiệu mã độc: {so_ioc} mục (cập nhật {ngay_ioc})").pack(side="left", padx=14, pady=5)
    admin_txt = "● Quyền Administrator" if is_admin() else "● Quyền thường (nên cấp Admin)"
    tk.Label(info_bar, bg="#e8efff", fg=("#16a34a" if is_admin() else "#b45309"),
             font=("Segoe UI", 9, "bold"), text=admin_txt).pack(side="right", padx=14, pady=5)

    now = datetime.now()
    default_time = {
        "gio": now.strftime("%H"),
        "phut": now.strftime("%M"),
        "ngay": now.strftime("%d"),
        "thang": now.strftime("%m"),
        "nam": now.strftime("%Y")
    }

    form_frame = ttk.LabelFrame(root, text="Thông tin kiểm tra", padding=(20, 15))
    form_frame.pack(fill="x", padx=25, pady=(18, 5))

    entries = {}
    for i, (key, label) in enumerate(MANUAL_FIELDS_DEF):
        ttk.Label(form_frame, text=label, style="Field.TLabel").grid(row=i, column=0, sticky="w", pady=5)
        ent = ttk.Entry(form_frame, width=42, font=("Segoe UI", 10))
        ent.grid(row=i, column=1, sticky="ew", padx=15, pady=5)
        if key in default_time:
            ent.insert(0, default_time[key])
        _bind_auto_cap(ent)
        entries[key] = ent
    form_frame.columnconfigure(1, weight=1)

    status_label = ttk.Label(root, text="", foreground="#374151", font=("Segoe UI", 9),
                             wraplength=540, justify="left")
    status_label.pack(pady=(12, 4), padx=25, fill="x")

    if not is_admin():
        status_label.config(text="Nhập đầy đủ các trường thông tin để nội dung biên bản đầy đủ"
                                 "Khi bấm Xác nhận, hệ thống sẽ hỏi nâng quyền Admin hãy bấm có/Yes.",
                            foreground="#1e3a8a")

    progress = ttk.Progressbar(root, orient="horizontal", length=560, mode="determinate",
                                maximum=TOTAL_PROGRESS_STEPS, style="Green.Horizontal.TProgressbar")
    progress.pack(pady=6, padx=25)

    cleanup_var = tk.BooleanVar(value=False)
    cleanup_chk = ttk.Checkbutton(
        root, text="Sau khi xuất biên bản: TỰ ĐỘNG XÓA công cụ (file exe + dữ liệu), giữ lại biên bản .docx",
        variable=cleanup_var)
    cleanup_chk.pack(pady=(6, 0))

    shutdown_var = tk.BooleanVar(value=False)
    shutdown_chk = ttk.Checkbutton(
        root, text="Sau khi xóa xong: TỰ ĐỘNG TẮT MÁY",
        variable=shutdown_var)
    shutdown_chk.pack(pady=(2, 0))

    def _sync_cleanup_options():
        if shutdown_var.get():
            cleanup_var.set(True)
        if not cleanup_var.get():
            shutdown_var.set(False)
    cleanup_chk.config(command=_sync_cleanup_options)
    shutdown_chk.config(command=_sync_cleanup_options)

    submit_btn = tk.Button(root, text="XÁC NHẬN VÀ BẮT ĐẦU KIỂM TRA",
                            bg="#2563eb", fg="white", activebackground="#1d4ed8", activeforeground="white",
                            font=("Segoe UI", 12, "bold"), borderwidth=0, padx=18, pady=9,
                            cursor="hand2", relief="flat")
    submit_btn.pack(pady=(14, 8))

    progress_queue = _queue.Queue()
    result_holder = {}

    def on_progress(step, total, message):
        progress_queue.put(("progress", step, total, message))

    # ĐÃ CẬP NHẬT: GỌI HÀM BẪY LỖI KHI XẢY RA SỰ CỐ
    def worker(manual_data, cleanup, shutdown):
        try:
            path = fill_form(template_path, output_path, malware_file, vuln_file,
                              manual_data, on_progress=on_progress)
            scheduled = schedule_self_cleanup(shutdown=shutdown) if (cleanup or shutdown) else False
            progress_queue.put(("done", path, scheduled, shutdown))
        except Exception as e:
            error_trace = traceback.format_exc()
            try:
                dump_path = dump_error_log_and_compress(error_trace)
                error_msg = f"Đã xảy ra sự cố kỹ thuật: {str(e)}\n\nHệ thống đã tự động sao lưu hiện trường lỗi tại:\n{dump_path}"
            except Exception as dump_e:
                error_msg = f"Lỗi: {str(e)}\n(Lưu ý: Không thể ghi file log do {str(dump_e)})"
            progress_queue.put(("error", error_msg))

    def poll_queue():
        try:
            while True:
                item = progress_queue.get_nowait()
                if item[0] == "progress":
                    _, step, total, message = item
                    progress["maximum"] = total
                    progress["value"] = step
                    status_label.config(text=message, foreground="#374151")
                elif item[0] == "done":
                    result_holder["path"] = item[1]
                    progress["value"] = progress["maximum"]
                    if len(item) > 2 and item[2]:
                        if len(item) > 3 and item[3]:
                            msg = (f"Hoàn tất! Đã xuất file: {item[1]}\n"
                                   "CÔNG CỤ SẼ TỰ XÓA VÀ TẮT MÁY ngay sau khi bạn đóng cửa sổ này.\n"
                                   "Hãy đóng cửa sổ rồi rời khỏi máy. Biên bản .docx được giữ nguyên.")
                            btn_text = "Đóng cửa sổ - xóa công cụ và tắt máy"
                        else:
                            msg = (f"Hoàn tất! Đã xuất file: {item[1]}\n"
                                   "CÔNG CỤ SẼ TỰ XÓA (exe + dữ liệu) ngay sau khi bạn đóng cửa sổ này.\n"
                                   "Biên bản kết quả .docx được giữ nguyên.")
                            btn_text = "Đóng cửa sổ (công cụ sẽ tự xóa)"
                        fg = "#dc2626"
                    else:
                        msg = f"Hoàn tất! Đã xuất file: {item[1]}\nĐang mở file và thư mục kết quả..."
                        fg = "#16a34a"
                        btn_text = "Đã hoàn tất - Đóng cửa sổ"
                    status_label.config(text=msg, foreground=fg)
                    open_result(item[1])
                    submit_btn.config(text=btn_text, state="normal", bg="#16a34a", command=root.destroy)
                    return
                elif item[0] == "error":
                    status_label.config(text=f"LỖI HỆ THỐNG", foreground="#dc2626")
                    messagebox.showerror("Báo cáo Sự cố", item[1])
                    for ent in entries.values():
                        ent.config(state="normal")
                    submit_btn.config(state="normal", text="Thử lại", bg="#2563eb")
                    return
        except _queue.Empty:
            pass
        root.after(150, poll_queue)

    def elevated_watcher(status_file, payload_file):
        import time
        timeout = 300  # 5 phút tối đa
        start_t = time.time()
        seen_done = False
        while time.time() - start_t < timeout:
            if os.path.exists(status_file):
                try:
                    with open(status_file, "r", encoding="utf-8") as f:
                        data = json.load(f)
                    st = data.get("status")
                    if st == "progress":
                        progress_queue.put(("progress", data.get("step", 0), data.get("total", TOTAL_PROGRESS_STEPS), data.get("message", "")))
                    elif st == "done":
                        progress_queue.put(("done", data.get("output_path"), data.get("scheduled", False), data.get("shutdown", False)))
                        seen_done = True
                        break
                    elif st == "error":
                        progress_queue.put(("error", data.get("message", "Lỗi không xác định")))
                        seen_done = True
                        break
                except Exception:
                    pass
            time.sleep(0.2)
        if not seen_done:
            progress_queue.put(("error", "Quá thời gian chờ tiến trình kiểm tra (Timeout)."))
        # Dọn dẹp file tạm
        for p in (status_file, payload_file, status_file + f".tmp.{os.getpid()}"):
            try:
                if os.path.exists(p):
                    os.remove(p)
            except Exception:
                pass

    def on_submit():
        for key, ent in entries.items():
            if not ent.get().strip():
                messagebox.showerror("Lỗi", "Vui lòng không để trống các trường thông tin!")
                return
                
        manual_data = {key: ent.get().strip() for key, ent in entries.items()}
        # LUÔN tự viết hoa chữ cái đầu mỗi từ
        manual_data = {key: auto_capitalize(value) for key, value in manual_data.items()}

        want_admin = False
        if not is_admin():
            ans = messagebox.askyesno(
                "Yêu cầu quyền Administrator",
                "Để kiểm tra toàn diện hệ thống (đọc đầy đủ lịch sử USB, danh sách bản vá Hotfix, cấu hình sâu...),\n"
                "công cụ cần chạy với quyền Administrator.\n\n"
                "Bạn có muốn cấp quyền Administrator để tiến hành kiểm tra toàn diện không?\n\n"
                "• Chọn YES: Chạy với quyền Admin (kiểm tra toàn diện).\n"
                "• Chọn NO: Tiếp tục kiểm tra với quyền thông thường hiện tại."
            )
            want_admin = bool(ans)

        for ent in entries.values():
            ent.config(state="disabled")
        submit_btn.config(state="disabled", text="Đang kiểm tra, vui lòng đợi...", bg="#9ca3af")

        if want_admin:
            status_label.config(text="Đang yêu cầu quyền Administrator từ hệ thống...", foreground="#374151")
            import tempfile, time
            ts = int(time.time() * 1000)
            status_file = os.path.join(tempfile.gettempdir(), f"anattt_status_{os.getpid()}_{ts}.json")
            payload_file = os.path.join(tempfile.gettempdir(), f"anattt_payload_{os.getpid()}_{ts}.json")
            payload = {
                "template_path": template_path,
                "output_path": output_path,
                "malware_file": malware_file,
                "vuln_file": vuln_file,
                "manual_data": manual_data,
                "cleanup": cleanup_var.get(),
                "shutdown": shutdown_var.get(),
                "status_file": status_file,
                "no_open": False
            }
            try:
                with open(payload_file, "w", encoding="utf-8") as f:
                    json.dump(payload, f, ensure_ascii=False)
            except Exception:
                want_admin = False

            if want_admin:
                spawned = _spawn_elevated_worker(payload_file)
                if spawned:
                    status_label.config(text="Đã cấp quyền Admin. Đang tiến hành kiểm tra toàn diện...", foreground="#16a34a")
                    threading.Thread(target=elevated_watcher, args=(status_file, payload_file), daemon=True).start()
                    root.after(150, poll_queue)
                    return
                else:
                    # Người dùng bấm No trên hộp thoại UAC hoặc lỗi
                    status_label.config(text="Không nhận được quyền Admin từ UAC. Đang kiểm tra với quyền thông thường...", foreground="#b45309")
                    try:
                        if os.path.exists(payload_file):
                            os.remove(payload_file)
                    except Exception:
                        pass

        status_label.config(text="Bắt đầu quá trình kiểm tra...", foreground="#374151")
        threading.Thread(target=worker, args=(manual_data, cleanup_var.get(), shutdown_var.get()), daemon=True).start()
        root.after(150, poll_queue)

    submit_btn.config(command=on_submit)

    def on_close():
        root.destroy()
        sys.exit(0)

    root.protocol("WM_DELETE_WINDOW", on_close)
    lb_copyright = ttk.Label(root, text=f"Phiên bản v{APP_VERSION} | © 2026 Bản quyền thuộc về Mã Đức Hiển - All Rights Reserved",
                            font=("Segoe UI", 9, "italic"), foreground="gray")
    lb_copyright.pack(side="bottom", pady=10)
    
    root.mainloop()
    return result_holder.get("path")

def _auto_find_template(explicit_path=None):
    if explicit_path:
        return explicit_path
    search_dir = output_dir()
    try:
        all_files = os.listdir(search_dir)
    except Exception:
        all_files = []
    # Bỏ file tạm của Word (~$...) và file biên bản kết quả đã xuất
    candidates = [os.path.join(search_dir, f) for f in all_files
                  if f.lower().endswith(".docx")
                  and not f.startswith("~$")
                  and not f.upper().startswith("BIENBAN_KETQUA")]
    if not candidates:
        return None
    # Ưu tiên file tên chính xác "mau_bien_ban.docx" (file nhúng) để tránh lấy nhầm file khác
    for c in candidates:
        if os.path.basename(c).lower() == "mau_bien_ban.docx":
            return c
    for kw in ("mau", "mẫu", "bb_", "bien_ban", "biên_bản"):
        for c in candidates:
            if kw in os.path.basename(c).lower():
                return c
    return candidates[0]

def _auto_output_name(computer_name_hint=None):
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    name = computer_name_hint or (os.environ.get("COMPUTERNAME") if IS_WINDOWS else None) or "MAY"
    return f"BienBan_KetQua_{name}_{ts}.docx"

def _data_file(name):
    """Ưu tiên file ngoài (cạnh exe/script - dễ cập nhật), nếu không có thì dùng file nhúng."""
    external = os.path.join(output_dir(), name)
    if os.path.exists(external):
        return external
    return resource_path(name)

if __name__ == "__main__":
    # KHÔNG tự nâng quyền Administrator: chạy với quyền thường để Unikey/bàn phím
    # tiếng Việt gõ được tự do như khung chat (chạy Admin sẽ bị Windows UIPI
    # chặn Unikey -> không gõ được tiếng Việt). Nếu cần dữ liệu đầy đủ hơn,
    # người dùng có thể chủ động chuột phải -> "Chạy với quyền quản trị viên".
    parser = argparse.ArgumentParser(description="Tự động điền Biên bản kiểm tra ANATTT")
    parser.add_argument("template", nargs="?", default=None,
                         help="Đường dẫn file mẫu .docx")
    parser.add_argument("output", nargs="?", default=None,
                         help="Đường dẫn file kết quả .docx")
    parser.add_argument("--malware", default=None, help="File danh sách IOC mã độc")
    parser.add_argument("--vuln", default=None, help="File danh sách CVE Windows")
    parser.add_argument("--no-open", action="store_true", help="Không tự động mở file Word")
    parser.add_argument("--no-manual", action="store_true", help="Bỏ qua bước nhập tay")
    parser.add_argument("--cleanup", action="store_true",
                        help="Tự động xóa công cụ (exe + dữ liệu) sau khi xuất biên bản, giữ lại file .docx")
    parser.add_argument("--shutdown", action="store_true",
                        help="Sau khi xóa xong, tự động TẮT MÁY (tự bật chế độ xóa công cụ)")
    parser.add_argument("--run-worker", default=None,
                        help="Chạy tiến trình worker ngầm bằng file payload JSON")
    args = parser.parse_args()

    if args.run_worker:
        run_worker_from_payload(args.run_worker)
        sys.exit(0)

    template_path = _auto_find_template(args.template)
    if not template_path or not os.path.exists(template_path):
        bundled_guess = resource_path("mau_bien_ban.docx")
        if os.path.exists(bundled_guess):
            template_path = bundled_guess
    if not template_path or not os.path.exists(template_path):
        print("KHÔNG TÌM THẤY file mẫu .docx.")
        sys.exit(1)

    malware_path = args.malware or _data_file("malware_signatures.txt")
    vuln_path = args.vuln or _data_file("windows_vulnerabilities.txt")
    output_path = args.output or os.path.join(output_dir(), _auto_output_name())

    if args.no_manual:
        # ĐÃ CẬP NHẬT: BẪY LỖI CHO TRƯỜNG HỢP CHẠY KHÔNG CẦN GIAO DIỆN
        try:
            result_path = fill_form(template_path, output_path, malware_path, vuln_path, {})
            if not args.no_open:
                open_result(result_path)
            if args.cleanup or args.shutdown:
                if schedule_self_cleanup(shutdown=args.shutdown):
                    if args.shutdown:
                        print("Đã lên lịch TỰ XÓA công cụ và TẮT MÁY sau khi chương trình thoát. "
                              "Biên bản kết quả được giữ nguyên.")
                    else:
                        print("Đã lên lịch TỰ XÓA công cụ (exe + dữ liệu) sau khi chương trình thoát. "
                              "Biên bản kết quả được giữ nguyên.")
                else:
                    print("Không thể tự xóa (cần chạy dưới dạng file exe đóng gói).")
        except Exception as e:
            error_trace = traceback.format_exc()
            dump_path = dump_error_log_and_compress(error_trace)
            print(f"LỖI HỆ THỐNG: {e}\nĐã đóng gói file log mã hóa tại: {dump_path}")
    elif HAS_TK:
        run_gui_app(template_path, output_path, malware_path, vuln_path)
    else:
        print("Lỗi: Không thể khởi động giao diện Tkinter.")
        sys.exit(1)